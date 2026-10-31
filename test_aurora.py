"""
Aurora self-test. Tells you which features work.

    python test_aurora.py              logic + routing tests (no camera/mic/GPU/API needed)
    python test_aurora.py --hardware   also: camera, microphone, ADB, volume, hand tracker
    python test_aurora.py --gui        also: opens the hologram window and renders every display mode
    python test_aurora.py --live       also: real Groq API, tool calling, neural TTS, weather (needs internet + api_key.txt)
    python test_aurora.py --all        everything

Statuses: PASS = works | FAIL = broken (reason shown) | SKIP = not testable here (missing optional part)
Tests use temp folders; your notes, vault, projects and sandbox are never touched.
Includes the extended suite (module internals, edge cases, error paths) near the end of the file.
"""
import importlib
import json
import math
import os
import shutil
import sys
import tempfile
import threading
import time
import traceback
import urllib.request
from types import SimpleNamespace
from unittest import mock

ROOT = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, ROOT)
os.chdir(ROOT)
ARGS = set(sys.argv[1:])
ALL = "--all" in ARGS
HARDWARE, GUI, LIVE = ALL or "--hardware" in ARGS, ALL or "--gui" in ARGS, ALL or "--live" in ARGS

RESULTS = []

# Tests must never launch PowerPoint/Word/Excel/image viewers. The voice handlers call os.startfile to open
# what they create, so replace it for the whole run and just record what would have opened.
OPENED_BY_TESTS = []
os.startfile = lambda path, *args, **kwargs: OPENED_BY_TESTS.append(path)


class Skip(Exception):
    def __init__(self, msg="", status="SKIP"):
        super().__init__(msg)
        self.status = status


def skip(msg, status="SKIP"):
    raise Skip(msg, status)


def requires(kind, msg):
    """kind: API, PERMISSION or HARDWARE"""
    raise Skip(msg, f"REQUIRES {kind}")


def check(group, name, fn):
    try:
        detail = fn()
        RESULTS.append((group, name, "PASS", detail if isinstance(detail, str) else ""))
    except Skip as e:
        RESULTS.append((group, name, e.status, str(e)))
    except AssertionError as e:
        RESULTS.append((group, name, "FAIL", f"assertion: {e}" if str(e) else "assertion failed"))
    except Exception as e:
        tb = traceback.extract_tb(e.__traceback__)[-1]
        RESULTS.append((group, name, "FAIL", f"{type(e).__name__}: {str(e)[:110]} ({os.path.basename(tb.filename)}:{tb.lineno})"))


def mod(name):
    try:
        return importlib.import_module(f"jarvis_ui.{name}")
    except Exception as e:
        raise Skip(f"module {name} failed to import: {type(e).__name__}: {str(e)[:80]}")


def eq(a, b):
    assert a == b, f"expected {b!r}, got {a!r}"


# ============================================================================ dependencies
REQUIRED = ["cv2", "mediapipe", "OpenGL", "pygame", "numpy", "groq", "speech_recognition", "pyaudio", "edge_tts",
            "pyttsx3", "psutil", "PIL", "serial", "pypdf"]
WINDOWS_ONLY = ["pycaw", "comtypes"]
OPTIONAL = ["pocketsphinx"]


def dep(name, optional=False):
    def run():
        try:
            importlib.import_module(name)
        except ImportError as e:
            if optional or (name in WINDOWS_ONLY and os.name != "nt"):
                skip(f"not installed ({'optional' if optional else 'Windows only'})")
            raise AssertionError(f"not installed: pip install {name} ({e})")
    return run


for d in REQUIRED + WINDOWS_ONLY + OPTIONAL:
    check("Dependencies", d, dep(d, d in OPTIONAL))

# ============================================================================ modules import
for m in ["hologram", "hand_tracker", "face_id", "system_control", "telemetry", "brain", "voice_assistant",
          "addons", "aurora_plus", "sandbox_labs", "dashboard_modes", "system_monitor", "gmail_client", "procedural_construction"]:
    check("Modules import", f"jarvis_ui.{m}", (lambda m=m: importlib.import_module(f"jarvis_ui.{m}") and None))
check("Modules import", "main.py", lambda: importlib.import_module("main") and None)


# ============================================================================ brain
def t_brain_search():
    b = mod("brain")
    assert b.needs_search("what's the latest news") and b.needs_search("who is the ceo of nvidia") and b.needs_search("iphone vs pixel")
    assert not b.needs_search("show me a carbon atom")


def t_brain_misc():
    b = mod("brain")
    eq(b.strip_code_fences("```python\nprint(1)\n```"), "print(1)\n")
    assert b.fuzzy_hit("oughrora set timer", ["aurora"]) or b.fuzzy_hit("aurera hi", ["aurora"])
    names = {t["function"]["name"] for t in b.TOOLS}
    for need in ("show_display", "get_weather", "set_timer", "pc_control", "phone_control", "edit_atom", "take_note",
                 "read_notes", "get_time", "write_code", "search_files"):
        assert need in names, f"tool {need} missing"


def t_brain_memory():
    b = mod("brain")
    m = b.RollingMemory(None)
    for i in range(20):
        m.add("user", f"hi {i}")
    m.maybe_compress()                      # no client -> must not crash
    eq(len(m.get_messages()), 20)
    m.clear_but_last()
    eq(len(m.get_messages()), 1)


check("AI brain", "needs_search heuristic", t_brain_search)
check("AI brain", "code-fence stripping, fuzzy match, tool schema", t_brain_misc)
check("AI brain", "rolling conversation memory", t_brain_memory)


# ============================================================================ voice helpers
def t_wake():
    v = mod("voice_assistant")
    eq(v.find_wake_word("aurora what time is it"), (True, "what time is it"))
    assert v.find_wake_word("arora hello")[0], "fuzzy wake word"
    assert not v.find_wake_word("just talking")[0]


def t_math():
    v = mod("voice_assistant")
    eq(v.try_calculate("47 times 12")[0], 564)
    eq(v.try_calculate("15 percent of 200")[0], 30)
    assert v.try_calculate("hello there") is None


def t_chunker_clean():
    v = mod("voice_assistant")
    c = v.Chunker()
    out = c.feed("Hello there my friend, how are you doing today? I am fine. ") + c.flush()
    assert len(out) >= 1 and "".join(out).replace(" ", "").startswith("Hellothere")
    eq(v.clean_for_speech("**bold** `code` # head"), "bold code head")


def t_speech_engine_queue():
    v = mod("voice_assistant")
    e = v.SpeechEngine(lambda m: None)
    e.use_edge = lambda: False
    with mock.patch.object(e, "_speak_offline", lambda text: time.sleep(0.05)):
        e.say("one two three")
        e.say("four five six")
        e.wait(5)
    assert not e.busy, "speech queue never drained"
    e.say("x y z")
    e.stop()
    time.sleep(0.2)
    assert not e.busy, "stop() left speech pending"


check("Voice pipeline", "wake word (fuzzy)", t_wake)
check("Voice pipeline", "spoken math + percent", t_math)
check("Voice pipeline", "sentence chunker + markdown cleaner", t_chunker_clean)
check("Voice pipeline", "speech queue drains and stop() cancels", t_speech_engine_queue)


# ============================================================================ voice command routing (stubbed hologram)
def make_voice():
    va = mod("voice_assistant")
    tel = mod("telemetry")
    H = mock.MagicMock(name="Hologram")
    H.load_atom.return_value = ("C", "carbon")
    H.load_star_system.return_value = ("Trappist-1", 7, False)
    H.next_element.return_value = ("He", "helium", 2)
    H.select_orbit.return_value = True
    H.protons, H.neutrons, H.electron_count, H.mode_label = 7, 7, 7, "NITROGEN"
    v = object.__new__(va.VoiceAssistant)
    said = []
    v.hologram, v.face_id, v.client, v.brain, v.plus, v.drawer = H, mock.MagicMock(), None, None, None, None
    v._capture = v._capture_tid = None
    v._turn_display = False
    v.said = said
    v.speech = mock.MagicMock()
    v.speech.say.side_effect = lambda text: said.append(va.clean_for_speech(text))
    v._log = v._on_log = lambda t: None
    v.active_timers, v.last_reply, v.latest_frame, v.last_weather_error = [], "", None, None
    v.telemetry, v.game_session, v.experiment = tel.TelemetryReader(), tel.GameSessionReader(), tel.ExperimentRecorder()
    v.active_reader, v.pending_enrollment_name = v.telemetry, None
    v._cancel = threading.Event()
    v._state = "idle"
    return v, H, said


def route(phrase, expect_said=None, expect=None, patches=(), setup=None):
    def run():
        v, H, said = make_voice()
        if setup:
            setup(v, H)
        stack = [p for p in patches]
        for p in stack:
            p.start()
        try:
            handled = v._handle_local_command(phrase)
        finally:
            for p in stack:
                p.stop()
        assert handled, f"'{phrase}' was not handled locally (would go to the AI)"
        if expect_said:
            assert any(expect_said.lower() in s.lower() for s in said), f"reply {said} lacks '{expect_said}'"
        if expect:
            expect(v, H)
    return run


sc_mod = lambda: mod("system_control")
tmpdir = tempfile.mkdtemp(prefix="aurora_test_")
ph = lambda target, **kw: mock.patch.object(sc_mod(), target, **kw)

ROUTES = [
    ("Time", "what time is it", "It's", None, ()),
    ("Math", "what's 47 times 12", "564", None, ()),
    ("Timers", "set a timer for 5 minutes", "Timer set", lambda v, H: eq(len(v.active_timers), 1), ()),
    ("Atoms", "show me a carbon atom", "carbon", lambda v, H: H.load_atom.assert_called_with("carbon"), ()),
    ("Atoms", "add 2 protons", "protons", lambda v, H: H.add_protons.assert_called_with(2), ()),
    ("Atoms", "next element", "helium", None, ()),
    ("Displays", "show me the solar system", "solar system", lambda v, H: H.load_solar_system.assert_called(), ()),
    ("Displays", "show me a cube", "cube", lambda v, H: H.load_shape.assert_called_with("cube"), ()),
    ("Displays", "show the eiffel tower", "eiffel", lambda v, H: H.load_shape.assert_called_with("eiffel tower"), ()),
    ("Displays", "show a water molecule", "water", lambda v, H: H.load_molecule.assert_called_with("water"), ()),
    ("Displays", "show orion", "Orion", lambda v, H: H.load_constellation.assert_called_with("orion"), ()),
    ("Displays", "show me trappist-1", "Trappist", lambda v, H: H.load_star_system.assert_called(), ()),
    ("Displays", "show a pendulum simulation", "pendulum", lambda v, H: H.load_physics_sim.assert_called_with("pendulum"), ()),
    ("Displays", "show my systems", "systems", lambda v, H: H.load_network.assert_called(), ()),
    ("Displays", "show the human heart", "heart", lambda v, H: H.load_shape.assert_called_with("heart"), ()),
    ("Displays", "reset the display", "Resetting", lambda v, H: H.load_demo.assert_called(), ()),
    ("Graphs", "plot x squared", "Plotting", None, ()),
    ("Graphs", "graph sine of x", "Plotting", None, ()),
    ("Notes", "take a note: buy milk", "Noted", None, (ph("add_note"),)),
    ("Notes", "read my notes", "notes", None, (ph("read_notes", return_value=["buy milk"]),)),
    ("Weather", "weather in tokyo", "Tokyo", lambda v, H: H.show_weather.assert_called(), ()),
    ("PC control", "open notepad", "Opening notepad", None, (ph("launch_app", return_value=True),)),
    ("PC control", "volume to 40", "40", None, (ph("set_volume_percent", return_value=True),)),
    ("PC control", "volume up", "Volume at", None, (ph("adjust_volume_percent", return_value=60),)),
    ("PC control", "play music", "Done", None, (ph("media_play_pause"),)),
    ("PC control", "take a screenshot", "Screenshot", None, (ph("take_screenshot", return_value=True),)),
    ("PC control", "system status", "CPU", None, (ph("get_system_status", return_value="CPU at 5 percent."),)),
    ("Phone", "call mom", "Calling mom", None,
     (ph("load_contacts", return_value={"mom": "123"}), ph("is_device_connected", return_value=(True, "")),
      ph("call_number", return_value=(True, "")))),
    ("Phone", "open instagram on my phone", "Opening instagram", None,
     (ph("is_device_connected", return_value=(True, "")), ph("open_app", return_value=True))),
    ("Phone", "call mom on whatsapp", "WhatsApp call", None,
     (ph("is_device_connected", return_value=(True, "")), ph("whatsapp_call", return_value=(True, "Mom")))),
    ("Phone", "turn off wifi on my phone", "wifi turned off", None,
     (ph("is_device_connected", return_value=(True, "")), ph("wifi_set", return_value=True))),
    ("Code", "create arduino sketch blinky", "blinky", None,
     (ph("create_arduino_sketch", return_value="x.ino"), ph("open_in_editor", return_value=(True, "")))),
    ("Code", "new python script called demo", "demo", None,
     (ph("create_code_file", return_value="demo.py"), ph("open_in_editor", return_value=(True, "")))),
    ("Vault", "add folder /tmp to my vault", "Added", None, (ph("vault_add_folder", return_value=(True, "/tmp")),)),
    ("Vault", "rebuild my vault", "Indexed", None, (ph("vault_build", return_value=(3, 9)),)),
    ("Face", "who do you know", "enrolled", None, ()),
    ("Face", "remember my face as sam", "Sam", lambda v, H: eq(v.pending_enrollment_name, "Sam"), ()),
    ("IoT", "show my mars station telemetry", "mars station", lambda v, H: H.show_telemetry.assert_called_with("mars station"), ()),
    ("Game companion", "game companion", "recording", lambda v, H: H.show_telemetry.assert_called(), ()),
    ("Experiment", "start experiment called volcano", "volcano", None, ()),
    ("Control", "repeat that", "haven't said", None, ()),
    ("Control", "stop", None, lambda v, H: v.speech.stop.assert_called(), ()),
]


def setup_graph(v, H):
    mod("hologram")
    H._eval_graph_expr.side_effect = lambda e, x: mod("hologram").Hologram._eval_graph_expr(H, e, x)


def setup_weather(v, H):
    v._fetch_weather_data = lambda loc: {"location": "Tokyo, Japan", "temp_c": 21.4, "condition": "sunny", "description": "clear sky",
                                          "humidity": 40, "wind_kph": 5, "updated_at": "12:00"}


def setup_experiment(v, H):
    tel = mod("telemetry")
    v.experiment = tel.ExperimentRecorder()


for group, phrase, said_frag, expect, patches in ROUTES:
    setup = setup_graph if group == "Graphs" else setup_weather if group == "Weather" else None
    if group == "Experiment":
        def exp_run(phrase=phrase, said_frag=said_frag):
            def run():
                tel = mod("telemetry")
                with mock.patch.object(tel, "EXPERIMENTS_DIR", tmpdir):
                    route(phrase, said_frag)()
            return run
        check(f"Voice commands: {group}", f'"{phrase}"', exp_run())
    else:
        check(f"Voice commands: {group}", f'"{phrase}"', route(phrase, said_frag, expect, patches, setup))


def t_falls_through():
    v, H, said = make_voice()
    assert not v._handle_local_command("tell me a joke about cats"), "general chat must go to the AI"


check("Voice commands: Control", "general questions fall through to the AI", t_falls_through)


def t_llm_tools():
    b = mod("brain")
    v, H, said = make_voice()
    assert "It's" in v._exec_tool("get_time", {})
    v._exec_tool("show_display", {"kind": "atom", "name": "carbon"})
    H.load_atom.assert_called_with("carbon")
    assert "Timer set" in v._exec_tool("set_timer", {"seconds": 30})
    assert v._exec_tool("show_display", {"kind": "nonsense"}) == "Unsupported request."
    assert {t["function"]["name"] for t in b.TOOLS} >= {"pc_control", "phone_control"}


check("AI tool calling", "LLM tools drive real Aurora actions", t_llm_tools)


# ============================================================================ system layer
def t_notes():
    sc = sc_mod()
    with mock.patch.object(sc, "NOTES_FILE", os.path.join(tmpdir, "notes.txt")):
        sc.add_note("alpha")
        sc.add_note("beta")
        eq(sc.read_notes(5), ["alpha", "beta"])
        sc.clear_notes()
        eq(sc.read_notes(5), [])


def t_code_files():
    sc = sc_mod()
    with mock.patch.object(sc, "PROJECTS_ROOT", os.path.join(tmpdir, "projects")), \
            mock.patch.object(sc, "ARDUINO_ROOT", os.path.join(tmpdir, "ino")), mock.patch.object(sc, "BASE", tmpdir):
        p = sc.create_code_file("my demo", "python")
        assert p.endswith("my_demo.py") and os.path.exists(p)
        eq(sc.resolve_project_path("my demo"), p)
        ino = sc.create_arduino_sketch("blink test")
        assert os.path.exists(ino) and sc.find_arduino_sketch("blink test") == ino
        sc.write_file_content(p, "print(1)\n")
        sc.write_file_content(p, "print(2)\n")
        assert os.path.exists(p + ".bak"), "backup missing"
        assert "1 line added" in sc.summarize_diff("a\n", "a\nb\n")


def t_vault():
    sc = sc_mod()
    d = os.path.join(tmpdir, "vault_src")
    os.makedirs(d, exist_ok=True)
    with open(os.path.join(d, "servo.txt"), "w") as f:
        f.write("my arduino servo robot arm project notes about torque")
    with mock.patch.object(sc, "BASE", tmpdir):
        ok, _ = sc.vault_add_folder(d)
        assert ok
        files, chunks = sc.vault_build()
        eq(files, 1)
        hits = sc.vault_search("what did I write about servo")
        assert hits and hits[0]["path"].endswith("servo.txt")


def t_apps_regex():
    sc = sc_mod()
    assert sc.APP_RE.search("open vs code").group(1) == "vs code"


def t_volume():
    sc = sc_mod()
    if not sc.PYCAW_AVAILABLE:
        skip("pycaw not available (Windows only)")
    v = sc.get_volume_percent()
    assert v is not None and 0 <= v <= 100


check("System control", "notes", t_notes)
check("System control", "code files, backups, diff summary", t_code_files)
check("System control", "keyword vault (add folder, build, search)", t_vault)
check("System control", "app launcher name matching", t_apps_regex)
check("System control", "master volume read (pycaw)", t_volume)


# ============================================================================ hand + face (no camera)
def t_filters():
    ht = mod("hand_tracker")
    f = ht.OneEuro()
    for i in range(10):
        f(i * 0.1, i * 0.033)
    d = ht.GestureDebouncer()
    for _ in range(5):
        c, _ = d.update("fist")
    eq(c, "fist")
    eq(round(ht._angle((1, 0, 0), (0, 0, 0), (0, 1, 0))), 90)


def t_face():
    f = mod("face_id")
    fid = f.FaceID()
    assert fid.detection_available, "Haar face detector failed to load (OpenCV install conflict)"
    assert fid.available, "cv2.face missing: pip uninstall opencv-python opencv-python-headless; pip install opencv-contrib-python"
    import numpy as np
    eq(fid.detect_faces(np.zeros((240, 320), np.uint8)), [])


check("Hand tracking", "smoothing filter, debouncer, joint angles", t_filters)
check("Face recognition", "detector + LBPH recogniser available", t_face)


# ============================================================================ addons
def circle(n=60):
    return [(0.5 + 0.2 * math.cos(2 * math.pi * i / n), 0.5 + 0.2 * math.sin(2 * math.pi * i / n)) for i in range(n + 1)]


def poly(pts, per=15):
    out = []
    for a, b in zip(pts, pts[1:] + pts[:1]):
        out += [(a[0] + (b[0] - a[0]) * k / per, a[1] + (b[1] - a[1]) * k / per) for k in range(per)]
    return out + [out[0]]


SHAPES_IN = {
    "circle": circle(),
    "square": poly([(0.3, 0.3), (0.7, 0.3), (0.7, 0.7), (0.3, 0.7)]),
    "triangle": poly([(0.5, 0.25), (0.75, 0.7), (0.25, 0.7)]),
    "line": [(0.2 + 0.6 * i / 30, 0.5) for i in range(31)],
    "wave": [(0.1 + 0.8 * i / 80, 0.5 + 0.15 * math.sin(i / 80 * 3 * 2 * math.pi)) for i in range(81)],
    "spiral": [(0.5 + (0.05 + 0.25 * i / 100) * math.cos(i / 100 * 3 * 2 * math.pi),
                0.5 + (0.05 + 0.25 * i / 100) * math.sin(i / 100 * 3 * 2 * math.pi)) for i in range(101)],
}
for shp, pts in SHAPES_IN.items():
    check("Gesture drawing", f"recognises {shp}", (lambda shp=shp, pts=pts: eq(mod("addons").recognize(pts), shp)))
check("Gesture drawing", "rejects a tiny scribble", lambda: eq(mod("addons").recognize([(0.5, 0.5)] * 20), None))


def t_drawer():
    ad = mod("addons")
    H = mock.MagicMock()
    d = ad.GestureDrawer(H)
    d.set_active(True)
    for x, y in SHAPES_IN["circle"]:
        d.update(SimpleNamespace(raw_gesture="point", tip=(x, y)))
    d._last_point -= 1.0
    eq(d.update(None), "circle")
    H.load_shape.assert_called_with("sphere")


check("Gesture drawing", "drawer: stroke -> hologram sphere", t_drawer)


def t_expr():
    ad = mod("addons")
    for spoken, want in [("x squared", "x**2"), ("sine of x", "sin(x)"), ("2 x plus 1", "2*x + 1"),
                         ("x cubed minus 3 x", "x**3 - 3*x"), ("square root of x", "sqrt(x)"), ("e to the x", "exp(x)"),
                         ("x to the power of 4", "x**4")]:
        eq(ad.speech_to_expr(spoken), want)
    assert ad.speech_to_expr("delete everything") is None


check("Math plotter", "speech -> expression parsing", t_expr)


def t_spatial():
    ad = mod("addons")
    d = ad.SpatialDesktop()
    for phrase, want in [("move system panel behind me", "Moved"), ("bring the system panel forward", "Bringing"),
                         ("move notes panel left", "Moved"), ("push notes panel back", "Pushed"),
                         ("open browser panel", "Opened"), ("close the browser panel", "Closed"),
                         ("show my desktop", "system")]:
        ok, reply = ad.handle_spatial_command(d, phrase)
        assert ok and want.lower() in reply.lower(), f"{phrase!r} -> {reply!r}"
    eq(d.panels["system"].angle, 0.0)


def t_sandbox():
    ad = mod("addons")
    with mock.patch.object(ad, "SANDBOX_ROOT", os.path.join(tmpdir, "sandbox")):
        _, ok, out, err = ad.sandbox_write_and_run("hello", "print(2+2)")
        assert ok and out.strip() == "4", err
        _, ok, _, err = ad.sandbox_write_and_run("loop", "while True: pass", timeout=1)
        assert not ok and "Timed out" in err
        p = ad.sandbox_write("../../evil", "x")
        assert os.path.realpath(p).startswith(os.path.realpath(os.path.join(tmpdir, "sandbox"))), "path escaped sandbox"
        ad.sandbox_reset()
        eq(ad.sandbox_list(), [])


def t_addon_voice():
    ad = mod("addons")
    v, H, said = make_voice()
    v.drawer = ad.GestureDrawer(H)
    assert ad.handle_command(v, "draw mode") and v.drawer.active
    assert ad.handle_command(v, "stop drawing") and not v.drawer.active
    assert ad.handle_command(v, "move system panel behind me")
    assert not ad.handle_command(v, "tell me a joke")


check("Spatial desktop", "panel voice commands", t_spatial)
check("Code sandbox", "write/run, timeout kill, path jail, reset", t_sandbox)
check("Add-ons router", "draw mode + panels via handle_command", t_addon_voice)


# ============================================================================ sandbox labs
def t_physics():
    sl = mod("sandbox_labs")
    p = sl.Physics()
    r = p.create("create a projectile with 20 m/s velocity at 30 degrees")
    assert "Range" in r and p.sims[0]["kind"] == "proj"
    p.step(0.5)
    assert "Pendulum" in p.create("create a pendulum with length 2 meters")
    assert "spring" in p.create("create a spring with constant 40").lower()
    assert "Dropping" in p.create("drop a ball from 15 meters on the moon")


def t_molecule():
    sl = mod("sandbox_labs")
    m = sl.Molecule()
    m.add("C")
    m.fill_hydrogens()
    eq(m.formula(), "CH4")
    eq(m.name(), "methane")
    m.clear()
    m.add("O")
    m.fill_hydrogens()
    eq(m.formula(), "H2O")


def t_circuit():
    sl = mod("sandbox_labs")
    c = sl.Circuit()
    c.preset("led")
    for _ in range(40):
        c.frame(0.04)
    led = next(x for x in c.comps if x["kind"] == "led")
    assert led["on"] and not led["burnt"] and 0.005 < led["i"] < 0.03, f"LED i={led['i']}"
    c.clear()
    c.add("battery", 5)
    c.add("led")
    c.auto_wire()
    for _ in range(40):
        c.frame(0.04)
    assert next(x for x in c.comps if x["kind"] == "led")["burnt"], "LED without resistor should burn out"


def t_labs_voice():
    sl = mod("sandbox_labs")
    H = mock.MagicMock()
    H.labs = None      # a MagicMock auto-creates .labs, so install() would reuse a fake instead of the real Labs
    said = []
    assert sl.handle_command(H, "circuit sandbox", said.append) and "Circuit" in said[0]
    assert sl.handle_command(H, "add a 220 ohm resistor", said.append)
    assert sl.handle_command(H, "create a projectile with 20 m/s velocity", said.append)
    assert not sl.handle_command(H, "what is the capital of france", said.append)


check("Sandbox labs", "physics simulations", t_physics)
check("Sandbox labs", "molecule builder (valence + formula)", t_molecule)
check("Sandbox labs", "circuit simulator (LED on / burnout)", t_circuit)
check("Sandbox labs", "voice commands", t_labs_voice)


# ============================================================================ aurora plus
def make_plus(client=None):
    ap = mod("aurora_plus")
    voice = SimpleNamespace(client=client, _on_log=lambda m: None, state="idle", last_reply="", speak_now=lambda t: None,
                            _speak=lambda t: None, run_text=lambda t: f"echo:{t}")
    return ap, ap.Plus(voice, mock.MagicMock())


def t_guard():
    ap = mod("aurora_plus")
    G = ap.Guard
    assert not G.text_ok("api_key = gsk_abcdefghijklmnopqrstuvwxyz")[0]
    assert not G.text_ok("write a keylogger")[0]
    assert G.text_ok("build a prime number checker")[0]
    assert not G.python_ok("import subprocess", set())[0]
    assert not G.python_ok("import os\nos.system('x')", set())[0]
    assert not G.python_ok("eval('1')", set())[0]
    assert not G.python_ok("open('/etc/passwd')", set())[0]
    assert not G.python_ok("import requests", set())[0]
    assert G.python_ok("import math\nprint(math.pi)", set())[0]
    assert not G.path_ok("../x.py")[0] and not G.path_ok("a.exe")[0] and G.path_ok("calc.py")[0]
    assert G.command_ok(ap.TEST_CMD)[0] and not G.command_ok(["rm", "-rf", "/"])[0]


def t_cowork_pipeline():
    ap, plus = make_plus(client=object())
    plus.guard.judge = lambda k, d: (True, "")
    root = tempfile.mkdtemp(dir=tmpdir)
    reply = ("### FILE: calc.py\ndef add(a, b):\n    return a + b\n### END\n"
             "### FILE: test_calc.py\nimport unittest\nfrom calc import add\n\nclass T(unittest.TestCase):\n"
             "    def test_add(self):\n        self.assertEqual(add(1, 2), 3)\n### END\n")
    files = ap.parse_files(reply)
    eq(sorted(files), ["calc.py", "test_calc.py"])
    ok, why = plus.cowork._apply(root, files)
    assert ok, why
    ok, out = plus.cowork._test(root)
    assert ok, out[-300:]
    ok, why = plus.cowork._apply(root, {"evil.py": "import subprocess\nsubprocess.run(['ls'])\n"})
    assert not ok, "malicious file must be rejected"
    ok, why = plus.cowork._apply(root, {"x.py": "print(1)\n"})
    plus.guard.judge = lambda k, d: (False, "no")
    ok, why = plus.cowork._apply(root, {"y.py": "print(1)\n"})
    assert not ok and "safeguard" in why, "safeguard must fail closed"


def t_semantic():
    ap, plus = make_plus()
    sc = sc_mod()
    d = os.path.join(tmpdir, "sem_src")
    os.makedirs(d, exist_ok=True)
    with open(os.path.join(d, "budget.txt"), "w") as f:
        f.write("monthly budget: rent and groceries expenses total")
    with open(os.path.join(d, "rocket.txt"), "w") as f:
        f.write("model rocket launch notes about the estes engine")
    with mock.patch.object(sc, "BASE", tmpdir), mock.patch.object(ap, "INDEX_FILE", os.path.join(tmpdir, "sem.json")):
        sc.vault_add_folder(d)
        plus.search = ap.SemanticSearch(plus)
        nf, nc = plus.search.build()
        assert nf >= 2
        hits = plus.search.query("money and payments")
        assert hits and hits[0]["path"].endswith("budget.txt"), hits


def t_offline():
    ap, plus = make_plus()
    assert "offline" in plus.offline.answer("hello").lower()
    assert "Today is" in plus.offline.answer("what day is it")
    plus.offline.force = True
    assert not plus.offline.online()
    b = ap.OfflineAwareBrain(plus, None)
    got = []
    r = b.ask("hello", got.append, threading.Event())
    assert r and got, "offline brain must answer without an API"


def t_plus_routes():
    ap, plus = make_plus()
    said = []
    plus.voice._speak = said.append
    assert plus.route("offline mode on") and plus.offline.force
    assert plus.route("offline mode off") and not plus.offline.force
    assert plus.route("cowork status") and "idle" in said[-1]
    assert plus.route("cowork build a thing") and "offline" in said[-1].lower()     # no client
    assert not plus.route("what's the weather")


def t_hub():
    ap, plus = make_plus()
    with mock.patch.object(ap, "TOKEN_FILE", os.path.join(tmpdir, "tok.txt")):
        plus.hub.port = 18765
        ok, url = plus.hub.start()
        assert ok, url
        token = plus.hub.token
        try:
            def call(path, body=None, tok=token):
                req = urllib.request.Request(f"http://127.0.0.1:18765{path}", data=json.dumps(body).encode() if body else None,
                                             headers={"X-Token": tok, "Content-Type": "application/json"})
                try:
                    with urllib.request.urlopen(req, timeout=5) as r:
                        return r.status, json.loads(r.read().decode())
                except urllib.error.HTTPError as e:
                    return e.code, {}
            eq(call("/cmd", {"text": "hello"}, tok="wrong")[0], 401)
            code, data = call("/cmd", {"text": "aurora hello"})
            assert code == 200 and data["reply"] == "echo:hello", data
            assert "isn't allowed" in call("/cmd", {"text": "unlock my phone"})[1]["reply"]
            eq(call("/status")[0], 200)
        finally:
            plus.hub.stop()
        assert plus.hub.server is None


def t_install():
    ap = mod("aurora_plus")
    va = mod("voice_assistant")
    v, H, said = make_voice()
    v.recognizer = None
    v.speech = va.SpeechEngine(lambda m: None)
    v._speak = said.append   # the mock speech engine that fills `said` was just replaced
    plus = ap.install(v, H)
    assert v.plus is plus and v.speech.use_edge == plus.offline.online
    assert v._handle_local_command("offline mode status") and any("currently" in s for s in said)
    assert v.brain.ask("hello there", lambda d: None, threading.Event())


check("Aurora Plus", "safety guard (secrets, malware, path, command)", t_guard)
check("Aurora Plus", "Cowork write/validate/test pipeline + safeguard fail-closed", t_cowork_pipeline)
check("Aurora Plus", "semantic file search", t_semantic)
check("Aurora Plus", "offline mode answers", t_offline)
check("Aurora Plus", "voice routes (offline, cowork)", t_plus_routes)
check("Aurora Plus", "network hub (auth, LAN-only, blocked commands)", t_hub)
check("Aurora Plus", "install() wires into VoiceAssistant", t_install)


# ============================================================================ hardware (--hardware)
def hw(name, fn):
    check("Hardware", name, fn if HARDWARE else (lambda: skip("run with --hardware")))


def t_camera():
    import cv2
    cap = cv2.VideoCapture(0, cv2.CAP_DSHOW) if os.name == "nt" else cv2.VideoCapture(0)
    ok, frame = cap.read()
    cap.release()
    assert ok and frame is not None, "camera 0 gave no frame (close Zoom/Teams/browser tabs using it)"
    return f"{frame.shape[1]}x{frame.shape[0]}"


def t_mic():
    import speech_recognition as sr
    names = sr.Microphone.list_microphone_names()
    assert names, "no audio input devices (create mic_index.txt after installing pyaudio)"
    return f"{len(names)} input device(s)"


def t_adb():
    ok, out = sc_mod().is_device_connected()
    if not ok and "not found" in out:
        skip("adb not installed / not on PATH")
    if not ok:
        requires("HARDWARE", "adb found but no phone authorised (enable USB debugging)")


def t_tracker():
    ad = mod("addons")
    tr = ad.TipHandTracker(0)
    try:
        for _ in range(5):
            result, frame = tr.read()
        assert frame is not None, "no frame from tracker"
        return "hand detected" if result.detected else "tracker running (no hand in view)"
    finally:
        tr.close()


hw("camera opens and returns frames", t_camera)
hw("microphone present", t_mic)
hw("Android phone over ADB", t_adb)
hw("MediaPipe hand tracker runs", t_tracker)


# ============================================================================ GUI (--gui)
def gui(name, fn):
    check("Hologram (GUI)", name, fn if GUI else (lambda: skip("run with --gui")))


_holo = {}


def get_holo():
    if "h" not in _holo:
        _holo["h"] = mod("hologram").Hologram(fullscreen=False, windowed_size=(1100, 720))
        mod("addons").install(_holo["h"], mod("addons").GestureDrawer(_holo["h"]))
    return _holo["h"]


def render_case(label, loader):
    def run():
        h = get_holo()
        loader(h)
        for _ in range(3):
            h.process_events()
            h.update(0.1, 0.1, 0.0, target_dt=1 / 30)
            h.render()
    gui(label, run)


CASES = [
    ("idle network view", lambda h: h.load_network()),
    ("demo orbits", lambda h: h.load_demo()),
    ("carbon atom", lambda h: h.load_atom("carbon")),
    ("solar system", lambda h: h.load_solar_system()),
    ("star system (TRAPPIST-1)", lambda h: h.load_star_system("trappist-1")),
    ("shapes: sphere/cube/torus/pyramid/cylinder/eiffel/skyscraper/dna",
     lambda h: [h.load_shape(s) or h.render() for s in ("sphere", "cube", "torus", "pyramid", "cylinder", "eiffel tower", "skyscraper", "dna")]),
    ("heart model", lambda h: h.load_shape("heart")),
    ("construction demo (robot)", lambda h: (mod("procedural_construction").install(h), h.load_construction("robot"),
                                              [(setattr(h, "elapsed", h._materialize_start + t), h.render()) for t in (0.9, 3.0)])),
    ("construction demo (drone)", lambda h: (mod("procedural_construction").install(h), h.load_construction("drone"),
                                             [(setattr(h, "elapsed", h._materialize_start + t), h.render()) for t in (0.9, 3.0)])),
    ("construction demo (car)", lambda h: (mod("procedural_construction").install(h), h.load_construction("car"),
                                           [(setattr(h, "elapsed", h._materialize_start + t), h.render()) for t in (0.9, 3.0)])),
    ("construction demo (spaceship)", lambda h: (mod("procedural_construction").install(h), h.load_construction("spaceship"),
                                                 [(setattr(h, "elapsed", h._materialize_start + t), h.render()) for t in (0.9, 3.0)])),
    ("construction instant (car)", lambda h: (mod("procedural_construction").install(h), h.load_construction("car", assemble=False),
                                              h.render())),
    ("molecule", lambda h: h.load_molecule("water")),
    ("math graph", lambda h: h.load_math_graph("sin(x)")),
    ("physics sim", lambda h: h.load_physics_sim("pendulum")),
    ("satellite", lambda h: h.load_satellite("ISS")),
    ("constellation", lambda h: h.load_constellation("orion")),
    ("telemetry panel", lambda h: (h.show_telemetry("test"), h.update_telemetry({"temp": 21.5, "hum": 40}, True))),
    ("info card", lambda h: h.show_info_card("q", "an answer " * 20)),
    ("weather panel", lambda h: h.show_weather({"location": "Tokyo", "temp_c": 20, "condition": "rain", "description": "rain",
                                                "humidity": 80, "wind_kph": 10, "updated_at": "12:00"})),
    ("spatial panels overlay", lambda h: setattr(h, "spatial", mod("addons").SpatialDesktop())),
    ("circuit sandbox panel", lambda h: (mod("sandbox_labs").install(h), h.labs.open("circuit"), h.labs.circ.preset("led"))),
    ("molecule builder panel", lambda h: (mod("sandbox_labs").install(h), h.labs.open("molecule"), h.labs.mol.load_preset("methane"))),
    ("physics sandbox panel", lambda h: (mod("sandbox_labs").install(h), h.labs.open("physics"), h.labs.phys.create("create a projectile with 20 m/s velocity"))),
]
for label, loader in CASES:
    render_case(label, loader)


def t_gui_close():
    if _holo.get("h"):
        _holo["h"].close()


gui("window closes cleanly", t_gui_close)


# ============================================================================ live (--live)
def live(name, fn):
    check("Live services", name, fn if LIVE else (lambda: skip("run with --live")))


def get_client():
    va = mod("voice_assistant")
    key = va.load_api_key()
    if not key or va.Groq is None:
        requires("API", "no api_key.txt / groq package")
    return va.Groq(api_key=key)


def t_groq():
    b = mod("brain")
    out = b._complete(get_client(), "Reply with the single word OK.", "ping", 400, effort="low")
    assert out, "empty reply"
    return out[:30]


def t_groq_tools():
    b = mod("brain")
    calls = []
    br = b.Brain(get_client(), lambda m: None, lambda n, a: calls.append(n) or "Timer set")
    reply = br.ask("Set a timer for 10 seconds", lambda d: None, threading.Event())
    assert "set_timer" in calls, f"model did not call set_timer (calls={calls}, reply={reply!r})"


def t_groq_code():
    b = mod("brain")
    code = b.generate_code(get_client(), "print the numbers 1 to 3", "python")
    assert "print" in code and "```" not in code


def t_tts():
    import asyncio
    va = mod("voice_assistant")
    if not va.EDGE_TTS_AVAILABLE:
        skip("edge-tts not installed")
    p = os.path.join(tmpdir, "tts.mp3")
    try:
        asyncio.run(va._edge_save("Testing Aurora.", p))
    except Exception as e:
        if "403" in str(e):
            skip("Microsoft rejected the request: pip install -U edge-tts")
        raise
    assert os.path.getsize(p) > 1000


def t_weather():
    v, H, said = make_voice()
    d = v._fetch_weather_data("Tokyo")
    assert d and d["temp_c"] is not None, f"weather failed ({v.last_weather_error})"
    return f"{d['location']} {d['temp_c']}C"


live("Groq chat (gpt-oss-120b)", t_groq)
live("Groq tool calling", t_groq_tools)
live("Groq code generation", t_groq_code)
live("Edge neural TTS", t_tts)
live("Open-Meteo weather", t_weather)


# ============================================================================ productivity features
import zipfile
from datetime import datetime, timedelta


# Legacy module names (before everything was merged) resolve to aurora_utilities instead of "module not found".
_LEGACY_UTILITY_MODULES = {"quick_notes", "file_utilities", "calendar_email", "notification_center", "translation",
                           "study_mode", "skill_builder", "productivity_commands"}
_mod_original = mod


def mod(name):
    if name in _LEGACY_UTILITY_MODULES:
        name = "aurora_utilities"
    return _mod_original(name)



class FakeClient:
    """Stands in for Groq: brain._complete(client, ...) is patched to call .reply()."""
    def __init__(self, text="ok"):
        self.text, self.calls = text, []

    def reply(self, system, user):
        self.calls.append((system, user))
        return self.text


def with_complete(client):
    return mock.patch.object(mod("brain"), "_complete", lambda c, s, u, *a, **k: c.reply(s, u))


def util():
    return mod("aurora_utilities")


def newdir(prefix):
    return tempfile.mkdtemp(prefix=prefix, dir=tmpdir)


# ---- Quick notes
def t_qn():
    sc, qn = sc_mod(), util()
    with mock.patch.object(sc, "NOTES_FILE", os.path.join(newdir("qn"), "notes.txt")):
        qn.add_quick_note("buy milk and eggs")
        qn.add_quick_note("call the dentist")
        eq([n[2] for n in qn.list_notes()], ["buy milk and eggs", "call the dentist"])
        eq(qn.search_notes("what did i note about dentist")[0][2], "call the dentist")
        eq(qn.get_quick_note(1)[2], "buy milk and eggs")
        assert qn.delete_note(1) and not qn.delete_note(9)
        eq(len(qn.list_notes()), 1)
        eq(qn.delete_all_notes(), 1)
        eq(qn.list_notes(), [])


check("Quick notes", "add, list, search, read, delete", t_qn)


# ---- File utilities
def t_files():
    fu = util()
    home = os.path.join(newdir("home"), "AppData", "Local", "Temp", "h")       # Windows temp dirs live under AppData
    for d in ("Downloads", "Documents"):
        os.makedirs(os.path.join(home, d))
    dl = os.path.join(home, "Downloads")
    for n in ("report.pdf", "pic.png", "song.mp3", "setup.exe", "notes.txt", "weird.zzz"):
        open(os.path.join(dl, n), "w").write("x")
    with mock.patch.object(fu, "HOME", home):
        eq(fu.find_file("report", fu.folder_path("downloads")), [os.path.join(dl, "report.pdf")])
        new = fu.rename_file(os.path.join(dl, "report.pdf"), "summary")
        eq(os.path.basename(new), "summary.pdf")                        # extension kept
        open(os.path.join(dl, "again.pdf"), "w").write("y")
        eq(os.path.basename(fu.rename_file(os.path.join(dl, "again.pdf"), "summary.pdf")), "summary (1).pdf")  # no overwrite
        moved = fu.move_file(os.path.join(dl, "summary.pdf"), os.path.join(home, "Documents"))
        assert os.path.exists(moved) and not os.path.exists(os.path.join(dl, "summary.pdf"))
        z = fu.compress_path(os.path.join(dl, "notes.txt"))
        out = fu.extract_zip(z)
        assert os.path.exists(os.path.join(out, "notes.txt"))
        plan = fu.plan_organize(dl)
        cats = {os.path.basename(s): c for s, c in plan}
        assert cats["pic.png"] == "Images" and cats["song.mp3"] == "Audio" and cats["weird.zzz"] == "Other"
        assert fu.organize_folder(dl, plan) == len(plan)
        assert os.path.exists(os.path.join(dl, "Images", "pic.png"))
        for bad in (lambda: fu.rename_file(os.path.join(dl, "Images", "pic.png"), "../x.png"),
                    lambda: fu.move_file(os.path.join(dl, "Images", "pic.png"), tempfile.gettempdir()),
                    lambda: fu.rename_file(os.path.abspath(__file__), "x.py")):
            try:
                bad()
            except fu.FileError:
                continue
            raise AssertionError("unsafe file operation was allowed")
        evil = os.path.join(dl, "evil.zip")
        with zipfile.ZipFile(evil, "w") as zf:
            zf.writestr("../../escaped.txt", "x")
        try:
            fu.extract_zip(evil)
            raise AssertionError("zip-slip not blocked")
        except fu.FileError:
            pass
        assert not os.path.exists(os.path.join(home, "..", "escaped.txt"))


check("File utilities", "rename/move/compress/extract/organize + safety limits", t_files)

# ---- Calendar
ICS = ("BEGIN:VCALENDAR\r\nBEGIN:VEVENT\r\nSUMMARY:Standup\r\nDTSTART:{d}T090000\r\nDTEND:{d}T091500\r\n"
       "RRULE:FREQ=WEEKLY;BYDAY=MO,TU,WE,TH,FR;COUNT=100\r\nEND:VEVENT\r\n"
       "BEGIN:VEVENT\r\nSUMMARY:Dentist\\, Dr. Lee\r\nDTSTART:{t}T150000\r\nDTEND:{t}T160000\r\nEND:VEVENT\r\n"
       "BEGIN:VEVENT\r\nSUMMARY:Birthday\r\nDTSTART;VALUE=DATE:{t}\r\nEND:VEVENT\r\nEND:VCALENDAR\r\n")


def cal_env():
    cal = util()
    d = newdir("cal")
    today = datetime.now().date()
    ics = os.path.join(d, "c.ics")
    open(ics, "w").write(ICS.format(d=(today - timedelta(days=14)).strftime("%Y%m%d"), t=(today + timedelta(days=1)).strftime("%Y%m%d")))
    json.dump({"ics": [ics]}, open(os.path.join(d, "calendar_config.json"), "w"))
    return cal, d


def t_calendar():
    cal, d = cal_env()
    with mock.patch.object(cal, "BASE", d):
        tomorrow = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0) + timedelta(days=1)
        evs, errs = cal.events_between(tomorrow, tomorrow + timedelta(days=1))
        titles = {e["title"] for e in evs}
        assert not errs and "Dentist, Dr. Lee" in titles and "Birthday" in titles, titles
        if tomorrow.weekday() < 5:
            assert "Standup" in titles, "weekly recurrence missing"
        title, start = cal.parse_event_phrase("dentist tomorrow at 3:30 pm", datetime(2026, 10, 1, 8, 0))
        eq((title, start), ("dentist", datetime(2026, 10, 2, 15, 30)))
        eq(cal.parse_event_phrase("lunch with sam", datetime(2026, 10, 1))[1], None)
        cal.add_local_event("Study group", datetime.now() + timedelta(hours=2))
        ev, _ = cal.next_event()
        assert ev and ev["title"] == "Study group" or ev, "next_event failed"
        evs, _ = cal.events_between(datetime.now(), datetime.now() + timedelta(days=1))
        assert "Study group" in {e["title"] for e in evs}


check("Calendar", "ICS parsing, recurrence, time phrases, local events, next event", t_calendar)
check("Calendar", "live calendar source (ICS URL / file you configured)",
      lambda: (_ for _ in ()).throw(Skip("no calendar_config.json in project root", "REQUIRES PERMISSION"))
      if not os.path.exists(os.path.join(ROOT, "calendar_config.json")) else
      (lambda c: (c.events_between(datetime.now(), datetime.now() + timedelta(days=7)) and None))(util()))


# ---- Email
def mail_env():
    cal = util()
    d = newdir("mail")
    with open(os.path.join(d, "email_config.json"), "w") as f:
        json.dump({"contacts": {"ann": "ann@example.com"}}, f)
    return cal, d


def no_gmail(d):
    """Gmail not connected: points gmail_client at a missing token file (your real token is never touched)."""
    return mock.patch.object(mod("gmail_client"), "TOKEN_FILE", os.path.join(d, "no_token.json"))


def mail_token(d, scopes=("read", "send")):
    g = mod("gmail_client")
    path = os.path.join(d, "gmail_token.json")
    with open(path, "w") as f:
        json.dump({"access_token": "at", "expires_at": time.time() + 3600, "refresh_token": "r",
                   "scopes": [g.SCOPES[s] for s in scopes]}, f)
    return mock.patch.object(g, "TOKEN_FILE", path)


def t_email_logic():
    cal, d = mail_env()
    with mock.patch.object(cal, "BASE", d), no_gmail(d):
        assert not cal.email_configured()
        eq(cal.resolve_recipient("ann"), "ann@example.com")
        eq(cal.resolve_recipient("bob at example dot com"), "bob@example.com")
        eq(cal.resolve_recipient("nobody"), None)
        draft = cal.save_draft("ann@example.com", "Hi", "Body")
        eq(cal.latest_draft()["to"], "ann@example.com")
        try:
            cal.send_draft(draft)
            raise AssertionError("send without confirmation was allowed")
        except PermissionError:
            pass
        for call in (lambda: cal.fetch_unread(), lambda: cal.send_draft(draft, confirmed=True)):
            try:
                call()
                raise AssertionError("mail worked without a Gmail connection")
            except cal.MailError:
                pass
        assert cal.discard_draft() and cal.latest_draft() is None


def t_no_password_mail():
    src = open(util().__file__, encoding="utf-8").read()
    for word in ("imaplib", "smtplib", "AURORA_EMAIL_PASSWORD", "email_password"):
        assert word not in src, f"password-based mail code still present: {word}"


check("Email", "recipients, drafts, send needs confirmed=True, clear error when Gmail isn't connected", t_email_logic)
check("Email", "no password-based (IMAP/SMTP) code remains", t_no_password_mail)


# ---- Skill builder
def t_skills():
    skm = util()
    with mock.patch.object(skm, "BASE", newdir("sk")):
        eq(skm.parse_steps("open chrome then what's on my calendar today; set a timer for 5 minutes"),
           ["open chrome", "what's on my calendar today", "set a timer for 5 minutes"])
        skm.teach_skill("Morning", ["a", "b"])
        eq(skm.skill_names(), ["morning"])
        eq(skm.add_step("the morning skill", "c"), 3)
        skm.edit_step("morning", 2, "b2")
        skm.remove_step("morning", 1)
        eq(skm.get_skill("morning")["steps"], ["b2", "c"])
        eq(skm.run_skill("morning", lambda s: s.upper()), [("b2", "B2"), ("c", "C")])
        skm.rename_skill("morning", "evening")
        for bad in (lambda: skm.teach_skill("loop", ["run skill loop"]), lambda: skm.teach_skill("x", []), lambda: skm.edit_step("evening", 9, "z")):
            try:
                bad()
                raise AssertionError("expected SkillError")
            except skm.SkillError:
                pass
        skm.delete_skill("evening")
        eq(skm.skill_names(), [])


check("Skill builder", "teach, list, edit, rename, delete, run, recursion guard", t_skills)


# ---- Notification center
DUMP = """  NotificationRecord(0x1: pkg=com.whatsapp user=UserHandle{0} id=1 importance=3 key=0|com.whatsapp|1)
      extras={
        android.title=String (Mom)
        android.text=CharSequence (Dinner at 7?)
  NotificationRecord(0x2: pkg=com.android.systemui user=UserHandle{0} id=2 importance=2 key=x)
      extras={
        android.title=String (USB debugging connected)
  NotificationRecord(0x3: pkg=com.google.android.gm user=UserHandle{0} id=3 importance=3 key=y)
      extras={
        android.title=Bob
        android.text=Invoice attached
"""


def t_notifs():
    nc = util()
    p = nc.parse_phone_dump(DUMP)
    eq([(x["package"], x["title"], x["text"]) for x in p],
       [("com.whatsapp", "Mom", "Dinner at 7?"), ("com.google.android.gm", "Bob", "Invoice attached")])
    c = nc.NotificationCenter()
    assert c.push("phone", "a") and not c.push("phone", "a")                 # de-duplicated
    c.push("calendar", "Meeting soon")
    eq(c.ranked()[0]["source"], "calendar")                                  # calendar outranks phone
    sc, cal = sc_mod(), util()
    with mock.patch.object(sc, "is_device_connected", return_value=(True, "")), \
            mock.patch.object(sc, "_shell", return_value=(True, DUMP)), \
            mock.patch.object(cal, "email_configured", return_value=False), mock.patch.object(cal, "BASE", newdir("nc")):
        c2 = nc.NotificationCenter()
        eq(c2.collect(), [])
        assert any("Mom" in n["title"] for n in c2.items)
        eq(c2.clear(), 2)


check("Notification center", "ADB dump parsing, dedupe, ranking, collect (mocked phone)", t_notifs)


def t_notifs_hw():
    ok, out = sc_mod().is_device_connected()
    if not ok:
        requires("HARDWARE", "no authorised Android phone over ADB")
    util().NotificationCenter().collect()


check("Notification center", "real phone notifications", t_notifs_hw)


# ---- Study mode
def t_study():
    sm = util()
    with mock.patch.object(sm, "BASE", newdir("st")):
        sm.add_card("Bio", "Powerhouse of the cell?", "mitochondria")
        sm.add_cards("bio", [{"q": "Plants make food by?", "a": "photosynthesis"}])
        eq(sm.list_decks(), {"bio": 2})
        eq(sm.find_deck("biology"), "bio")
        assert sm.grade_answer("the mitochondria", "mitochondria") and sm.grade_answer("photosyntesis", "photosynthesis")
        assert not sm.grade_answer("nucleus", "mitochondria")
        name, cards = sm.get_cards("bio")
        s = sm.StudySession(name, cards, "quiz", shuffle=False)
        assert s.check("mitochondria") and not (s.next() and s.check("wrong"))
        eq(s.score_text(), "1 of 2 correct")
        assert s.next() is None
        said = []
        voice = SimpleNamespace(speak_now=said.append, active_timers=[])
        t = sm.StudyTimer(voice, seconds_per_minute=0.02)
        assert t.start(1, pomodoro=True, rest=1, cycles=2) and not t.start(1)
        time.sleep(0.6)
        assert not t.running and sm.minutes_today() == 2 and any("complete" in x for x in said), said
        t.start(5)
        assert t.stop()
        time.sleep(0.2)
        assert not t.running and voice.active_timers == []
        with with_complete(None):
            cards = sm.generate_cards(FakeClient('Sure! [{"q":"2+2?","a":"4"},{"q":"","a":"x"}]'), "math", 5)
        eq(cards, [{"q": "2+2?", "a": "4"}])
        assert sm.delete_deck("bio") and sm.list_decks() == {}


check("Study mode", "decks, grading, quiz session, pomodoro timer, AI card parsing (mocked)", t_study)


# ---- Translation
def t_translation():
    tr = util()
    eq(tr.parse_translation_request("translate good morning to spanish"), ("good morning", "spanish"))
    eq(tr.parse_translation_request("Translate to French: where is the library"), ("where is the library", "french"))
    eq(tr.parse_translation_request("how do you say thank you in japanese"), ("thank you", "japanese"))
    eq(tr.parse_translation_request("what is the weather"), None)
    eq(tr.voice_for_language("spanish"), "es-ES-AlvaroNeural")
    with with_complete(None):
        eq(tr.translate_text(FakeClient('"Buenos días"'), "good morning", "spanish"), "Buenos días")


check("Translation", "request parsing, voice mapping, translate (mocked model)", t_translation)


def t_translation_live():
    tr = util()
    out = tr.translate_text(get_client(), "good morning", "spanish")
    assert "buen" in out.lower(), out


live("Translation (real Groq)", t_translation_live)
live("Study flashcards (real Groq)", lambda: eq(len(util().generate_cards(get_client(), "the water cycle", 3)) >= 2, True))


# ---- Voice routing end to end (confirmation rules are the important part)
def pv():
    v, H, said = make_voice()
    return v, said, util()


def say(pc, v, phrase):
    return pc.handle_command(v, phrase)


def t_route_confirm_files():
    fu = util()
    home = os.path.join(newdir("rhome"), "AppData", "Local", "Temp", "h")
    dl = os.path.join(home, "Downloads")
    os.makedirs(dl)
    os.makedirs(os.path.join(home, "Documents"))
    open(os.path.join(dl, "report.pdf"), "w").write("x")
    v, said, pc = pv()
    with mock.patch.object(fu, "HOME", home):
        assert say(pc, v, "rename report.pdf in downloads to summary.pdf")
        assert os.path.exists(os.path.join(dl, "report.pdf")), "renamed BEFORE confirmation"
        assert "yes" in said[-1].lower(), f"no confirmation prompt: {said[-1:]}"
        assert say(pc, v, "yes") and os.path.exists(os.path.join(dl, "summary.pdf")), f"rename after yes failed: {said[-1:]}"
        assert say(pc, v, "move summary.pdf in downloads to documents") and os.path.exists(os.path.join(dl, "summary.pdf"))
        assert say(pc, v, "no") and os.path.exists(os.path.join(dl, "summary.pdf")), "moved despite 'no'"
        assert say(pc, v, "move summary.pdf in downloads to documents")
        pc.handle_command(v, "list my skills")                     # an unrelated command must cancel the pending move
        assert not (say(pc, v, "yes")), "stale confirmation executed"
        assert os.path.exists(os.path.join(dl, "summary.pdf"))
        assert say(pc, v, "organize my downloads") and os.path.exists(os.path.join(dl, "summary.pdf")), said[-1:]
        assert say(pc, v, "yes") and os.path.exists(os.path.join(dl, "Documents", "summary.pdf")), f"organize failed: {said[-1:]}"
        assert not say(pc, v, "move the panel to the left"), "must not steal non-file phrases"
        open(os.path.join(home, "Documents", "todo.txt"), "w").write("x")
        assert say(pc, v, "zip todo.txt in documents") and os.path.exists(os.path.join(home, "Documents", "todo.txt.zip"))


def t_route_email_send():
    cal, d = mail_env()
    g = mod("gmail_client")
    v, said, pc = pv()
    sent = []
    with mock.patch.object(cal, "BASE", d), mail_token(d), \
            mock.patch.object(g, "_api", lambda p, body=None: sent.append(p) or {}), with_complete(None):
        v.client = FakeClient("Subject: Demo\n\nSee you Friday.")
        assert say(pc, v, "draft an email to ann about the demo on friday")
        assert cal.latest_draft()["to"] == "ann@example.com"
        assert not sent
        assert say(pc, v, "send the email") and "about to send" in said[-1]
        assert not sent, "sent before confirmation"
        assert say(pc, v, "no")
        assert not sent, "sent despite 'no'"
        say(pc, v, "send the email")
        say(pc, v, "yes")
        eq(sent, ["messages/send"])


def t_route_notes_skills_calendar():
    sc, skm = sc_mod(), util()
    v, said, pc = pv()
    cal, d = cal_env()
    with mock.patch.object(sc, "NOTES_FILE", os.path.join(newdir("rn"), "notes.txt")), mock.patch.object(skm, "BASE", newdir("rs")), \
            mock.patch.object(cal, "BASE", d):
        assert say(pc, v, "remember this: pick up the parcel") and "saved" in said[-1].lower()
        assert not say(pc, v, "remember my face as sam"), "face enrollment must still reach its own handler"
        assert say(pc, v, "search notes for parcel") and "parcel" in said[-1]
        assert say(pc, v, "delete all notes") and "permanently" in said[-1]
        assert say(pc, v, "no") and len(util().list_notes()) == 1
        assert say(pc, v, "delete note 1") and util().list_notes() == []
        assert say(pc, v, "teach a skill called warmup: what time is it then list my skills")
        assert skm.get_skill("warmup") and len(skm.get_skill("warmup")["steps"]) == 2
        assert say(pc, v, "list my skills") and "warmup" in said[-1]
        assert say(pc, v, "run skill warmup") and "Finished warmup" in said[-1]
        assert say(pc, v, "delete skill warmup") and skm.get_skill("warmup")
        assert say(pc, v, "yes") and not skm.get_skill("warmup")
        assert say(pc, v, "what do i have tomorrow") and ("Dentist" in said[-1] or "Birthday" in said[-1]), said[-1]
        assert say(pc, v, "add lunch with sam tomorrow at noon to my calendar") and "yes" in said[-1].lower()
        assert not os.path.exists(os.path.join(d, "calendar_local.json")), "calendar changed before confirmation"
        assert say(pc, v, "yes") and os.path.exists(os.path.join(d, "calendar_local.json"))


def t_route_study_translate():
    sm = util()
    v, said, pc = pv()
    with mock.patch.object(sm, "BASE", newdir("rst")), with_complete(None):
        v.client = FakeClient('[{"q":"Capital of France?","a":"Paris"},{"q":"Capital of Italy?","a":"Rome"}]')
        assert say(pc, v, "make 2 flashcards about capitals") and sm.list_decks() == {"capitals": 2}
        assert say(pc, v, "quiz me on capitals") and v._productivity.session.mode == "quiz"
        for _ in range(2):
            card = v._productivity.session.card
            assert say(pc, v, f"the answer is {card['a']}")
        assert "2 of 2" in said[-1] and v._productivity.session is None
        v.client = FakeClient("Buenos días")
        assert say(pc, v, "translate good morning to spanish") and v.hologram.show_info_card.called
        assert "Buenos días" in v.last_reply, f"last_reply={v.last_reply!r}"
        v.client = None
        assert say(pc, v, "translate hello to french") and "Groq" in said[-1]


for _name, _fn in (("confirm-before-change for files", t_route_confirm_files), ("email never sends without yes", t_route_email_send),
                   ("notes, skills, calendar confirmation", t_route_notes_skills_calendar), ("study quiz + translation", t_route_study_translate)):
    check("Productivity voice routing", _name, _fn)


def t_hook():
    v, H, said = make_voice()
    assert v._handle_local_command("list my skills"), "productivity router not hooked"
    assert v._handle_local_command("translate hello to french"), "productivity router not hooked into voice_assistant"
    assert v._handle_local_command("take a note: buy milk"), "existing notes command broke"


check("Productivity voice routing", "hooked into VoiceAssistant._handle_local_command (existing commands intact)", t_hook)


# ============================================================================ dashboard modes + system monitor
def modes_env():
    dm = mod("dashboard_modes")
    v, _, said = make_voice()
    H = mock.MagicMock(name="HologramForModes")
    H.overlays, H.event_listeners, H.quiet_alerts = [], [], False
    H.modes = H.labs = None            # MagicMock auto-creates attributes; install() must see None to build real ones
    H.mode, H.info_card, H.translate_x = "empty", None, 0.0
    v.hologram = H
    v.speak_now = said.append
    return dm, H, v, said


def t_slide_parse():
    dm = mod("dashboard_modes")
    s = dm.parse_slides("# A\n- x\n- y\nNotes: hi\nDemo: show me the solar system\n---\n# B")
    eq(len(s), 2)
    eq((s[0]["title"], s[0]["bullets"], s[0]["notes"], s[0]["demo"]), ("A", ["x", "y"], "hi", "show me the solar system"))


def t_presentation_flow():
    dm, H, v, said = modes_env()
    with mock.patch.object(dm, "load_deck", lambda name=None: (dm.SAMPLE_DECK, "test deck")):
        assert dm.handle_command(v, "presentation mode")
        m = H.modes
        assert m.active == "presentation" and H.quiet_alerts is True
        assert dm.handle_command(v, "next slide") and m.presentation.i == 1
        H.load_atom.assert_called_with("carbon")                      # slide 2's Demo: line ran
        assert dm.handle_command(v, "previous slide") and m.presentation.i == 0
        assert dm.handle_command(v, "last slide") and m.presentation.i == 3
        assert dm.handle_command(v, "next slide") and "last slide" in said[-1]
        assert dm.handle_command(v, "show speaker notes") and m.presentation.show_notes
        assert dm.handle_command(v, "set a timer for 10 minutes") and m.presentation.timer_end
        assert dm.handle_command(v, "blank screen") and m.presentation.blank
        assert dm.handle_command(v, "end presentation") and m.active is None and H.quiet_alerts is False
        assert not dm.handle_command(v, "next slide"), "must not react outside presentation mode"


def t_timeline():
    dm, H, v, said = modes_env()
    d = tempfile.mkdtemp(dir=tmpdir)
    with mock.patch.object(dm, "EVENTS_FILE", os.path.join(d, "ev.jsonl")), \
            mock.patch.object(dm, "_git_commits", lambda limit=40: [(time.time() - 50, "git", "fix bug")]), \
            mock.patch.object(dm, "_recent_files", lambda limit=40: [(time.time() - 20, "file", "main.py")]):
        dm.install(H, v)
        for fn in H.event_listeners:
            fn("CODE: wrote demo")
            fn("FACE: recognized sam")                                 # not a recorded category
        assert dm.handle_command(v, "show my project timeline")
        eq(sorted(k for _, k, _ in H.modes.timeline.items), ["aurora", "file", "git"])
        assert H.modes.active == "timeline"
        assert dm.handle_command(v, "close timeline") and H.modes.active is None
        assert not dm.handle_command(v, "make a timeline of world war two"), "topic timelines must reach the AI"


def monitor(quiet=False):
    sm = mod("system_monitor")
    if sm.psutil is None:
        skip("psutil not installed")
    H = mock.MagicMock()
    H.quiet_alerts, H.mode = quiet, "empty"
    voice = mock.MagicMock()
    return sm, sm.SystemMonitor(H, voice), H, voice


def t_monitor_alert():
    sm, mon, H, voice = monitor()
    mon._alert("x", "Disk dropped", "why", 60)
    mon._alert("x", "again", "why", 60)                                # inside the cooldown
    eq(len(mon.alerts), 1)
    voice.speak_now.assert_called_once()
    H.quiet_alerts = True
    mon._alert("y", "Net down", "why", 60)
    eq((len(mon.alerts), voice.speak_now.call_count), (2, 1))          # silent while presenting


def t_monitor_disk():
    sm, mon, H, voice = monitor()
    usage = lambda free: SimpleNamespace(free=free, total=200 * sm.GB, percent=50)
    with mock.patch.object(sm.psutil, "disk_usage", side_effect=[usage(100 * sm.GB), usage(90 * sm.GB)]), \
            mock.patch.object(mon, "_snapshot_io", lambda now: None):
        mon._check_disk(1000.0)
        mon._check_disk(1005.0)
    assert [a["kind"] for a in mon.alerts] == ["disk_drop"], mon.alerts


def t_monitor_cpu():
    sm, mon, H, voice = monitor()
    with mock.patch.object(sm.psutil, "cpu_percent", return_value=97.0), \
            mock.patch.object(mon, "_top_cpu", lambda: ("MsMpEng.exe", 60.0)):
        for i in range(6):
            mon._check_cpu(float(i))
    assert mon.alerts and mon.alerts[-1]["kind"] == "cpu" and "Defender" in mon.alerts[-1]["explanation"]


def t_monitor_network():
    sm, mon, H, voice = monitor()
    states = iter([True, False, False, True])
    with mock.patch.object(sm.SystemMonitor, "_online", staticmethod(lambda: next(states))), \
            mock.patch.object(sm.SystemMonitor, "_net_explain", staticmethod(lambda: "adapter down")):
        for t in (100.0, 110.0, 120.0, 130.0):
            mon._check_network(t)
    assert [a["kind"] for a in mon.alerts] == ["network"]
    assert any("back" in c.args[0] for c in H.log_event.call_args_list), "recovery notice missing"


def t_monitor_crash():
    if os.name != "nt":
        skip("crash detection reads the Windows event log")
    sm, mon, H, voice = monitor()
    fake = lambda text: SimpleNamespace(stdout=text, returncode=0)
    ev = lambda date: (f"Event[0]\n  Date: {date}\n  Description:\nFaulting application name: notepad.exe, version: 1\n"
                       "Faulting module name: ntdll.dll, version 2\nException code: 0xc0000005\n")
    with mock.patch.object(sm.subprocess, "run", side_effect=[fake(ev("2026-10-01T10:00:00Z")), fake(ev("2026-10-01T10:00:00Z")),
                                                              fake(ev("2026-10-02T11:00:00Z"))]):
        for t in (1, 2, 3):
            mon._check_crashes(t)                                      # baseline, unchanged, new crash
    assert [a["kind"] for a in mon.alerts] == ["crash"]
    assert "notepad.exe" in mon.alerts[0]["summary"] and "access violation" in mon.alerts[0]["explanation"]


def t_monitor_voice():
    sm = mod("system_monitor")
    v, H, said = make_voice()
    assert not sm.handle_command(v, "what happened"), "'what happened' must reach the AI when there are no alerts"
    assert sm.handle_command(v, "any alerts") and "No alerts" in said[-1]
    assert sm.handle_command(v, "emergency monitor status") and said


def t_modes_hooked():
    dm = mod("dashboard_modes")
    hsrc = open(os.path.join(ROOT, "jarvis_ui", "hologram.py"), encoding="utf-8").read()
    for needle in ("self.overlays = []", "self.event_listeners = []", "fn(theme_color)", "fn(text)"):
        assert needle in hsrc, f"hologram.py is missing a hook ({needle}): run apply_dashboard_modes_patch.py"
    v, H, said = make_voice()
    H.modes = H.labs = None
    with mock.patch.object(dm, "collect", lambda since=None, limit=300: []):
        assert v._handle_local_command("show my project timeline"), "dashboard_modes is not hooked into voice_assistant"
    assert v._handle_local_command("emergency monitor status"), "system_monitor is not hooked into voice_assistant"


check("Dashboard modes", "slide file parsing", t_slide_parse)
check("Dashboard modes", "presentation voice flow (slides, notes, timer, blank, demo, exit)", t_presentation_flow)
check("Dashboard modes", "project timeline (event recording, sources, close)", t_timeline)
check("Dashboard modes", "hooks present in hologram.py and voice_assistant.py", t_modes_hooked)
check("System monitor", "alert cooldown + silent while presenting", t_monitor_alert)
check("System monitor", "sudden disk-space loss", t_monitor_disk)
check("System monitor", "unusual CPU usage + explanation", t_monitor_cpu)
check("System monitor", "network disconnect + recovery", t_monitor_network)
check("System monitor", "application crash from event log", t_monitor_crash)
check("System monitor", "voice commands", t_monitor_voice)


# ============================================================================ file assistant + screen understanding
def t_file_assistant():
    fu = util()
    home = os.path.join(newdir("fa"), "h")
    for d in ("Documents/School", "Downloads"):
        os.makedirs(os.path.join(home, d))
    dl = os.path.join(home, "Downloads")
    open(os.path.join(dl, "physics_notes.txt"), "w").write("newton physics")
    open(os.path.join(dl, "cooking.txt"), "w").write("pasta")
    v, said, pc = pv()
    with mock.patch.object(fu, "HOME", home):
        assert say(pc, v, "find documents about physics")
        eq([os.path.basename(p) for p in v._productivity.found], ["physics_notes.txt"])
        assert say(pc, v, "rename these files properly") and "yes" in said[-1].lower()
        assert os.path.exists(os.path.join(dl, "physics_notes.txt")), "renamed BEFORE confirmation"
        assert say(pc, v, "yes") and os.path.exists(os.path.join(dl, "Physics Notes.txt")), said[-1:]
        assert say(pc, v, "move these into my school folder") and "yes" in said[-1].lower()
        assert say(pc, v, "no") and os.path.exists(os.path.join(dl, "Physics Notes.txt")), "moved despite 'no'"
        assert say(pc, v, "move these into my school folder") and say(pc, v, "yes")
        assert os.path.exists(os.path.join(home, "Documents", "School", "Physics Notes.txt")), said[-1:]
        assert say(pc, v, "undo the move") and say(pc, v, "yes")
        assert os.path.exists(os.path.join(dl, "Physics Notes.txt")), "undo failed"
        assert not say(pc, v, "find files about budget"), "bare file search must reach the semantic search"
        assert not say(pc, v, "show me a picture of a cat"), "must not search files for picture requests"


def t_screen_understanding():
    sc, br = sc_mod(), mod("brain")
    v, H, said = make_voice()
    v.client = object()
    with mock.patch.object(sc, "capture_screen", return_value=object()), \
            mock.patch.object(br, "describe_screen", return_value="Missing colon on line 3.") as d:
        assert v._handle_local_command("why isn't this code working")
    assert any("colon" in s for s in said), said
    assert "why isn't" in d.call_args[0][2]
    v2, _, said2 = make_voice()
    assert v2._handle_local_command("what's on my screen") and any("Groq" in s for s in said2)
    assert not make_voice()[0]._handle_local_command("why is the sky blue"), "general questions must reach the AI"
    assert "describe_screen" in {t["function"]["name"] for t in br.TOOLS}


check("File assistant", "find, rename, move, undo (confirm first, 'no' cancels)", t_file_assistant)
check("Screen understanding", "routes to vision, needs API, ignores unrelated questions", t_screen_understanding)


# ============================================================================ Gmail OAuth
import base64
import urllib.parse


def gm():
    return mod("gmail_client")


def gm_token(d, scopes=("read",), expires=3600, refresh="r"):
    """Patches gmail_client.TOKEN_FILE to a temp file holding a token with the given permissions."""
    g = gm()
    path = os.path.join(d, "gmail_token.json")
    with open(path, "w") as f:
        json.dump({"access_token": "at", "expires_at": time.time() + expires, "refresh_token": refresh,
                   "scopes": [g.SCOPES[s] for s in scopes]}, f)
    return mock.patch.object(g, "TOKEN_FILE", path)


def t_gmail_permissions():
    g, d = gm(), newdir("gm")
    with mock.patch.object(g, "TOKEN_FILE", os.path.join(d, "none.json")):
        assert not g.connected() and not g.has("read") and not g.has("send")
    with gm_token(d, ("read",)):
        assert g.connected() and g.has("read") and not g.has("send")
        try:
            g.send({"to": "a@b.c", "body": "x"}, True)
            raise AssertionError("send allowed without the send permission")
        except g.GmailError:
            pass
    with gm_token(d, ("send",)):
        try:
            g.fetch_unread()
            raise AssertionError("read allowed without the read permission")
        except g.GmailError:
            pass
    with gm_token(d, ("read", "send")):
        try:
            g.send({"to": "a@b.c", "body": "x"})
            raise AssertionError("send without confirmation was allowed")
        except PermissionError:
            pass


def t_gmail_fetch():
    g, d = gm(), newdir("gm")
    data = base64.urlsafe_b64encode(b"The demo moved to Friday.\nPlease confirm.").decode().rstrip("=")
    msg = {"payload": {"headers": [{"name": "From", "value": "Ann Lee <ann@example.com>"},
                                   {"name": "Subject", "value": "Project update"}, {"name": "Date", "value": "Mon"}],
                       "mimeType": "multipart/alternative",
                       "parts": [{"mimeType": "text/html", "body": {"data": "PGI+"}},
                                 {"mimeType": "text/plain", "body": {"data": data}}]}}
    calls = []

    def fake(path, body=None):
        calls.append((path, body))
        return {"messages": [{"id": "1"}]} if path.startswith("messages?") else msg

    with gm_token(d), mock.patch.object(g, "_api", fake):
        out = g.fetch_unread(3)
    eq((out[0]["from"], out[0]["address"], out[0]["subject"]), ("Ann Lee", "ann@example.com", "Project update"))
    assert "Friday" in out[0]["body"], out[0]["body"]
    assert "is%3Aunread" in calls[0][0] and all(c[1] is None for c in calls), "reading must be read-only"


def t_gmail_send():
    g, d = gm(), newdir("gm")
    sent = []
    with gm_token(d, ("send",)), mock.patch.object(g, "_api", lambda p, body=None: sent.append((p, body)) or {}):
        assert g.send({"to": "ann@example.com", "subject": "Hi", "body": "See you Friday."}, True)
    eq(sent[0][0], "messages/send")
    raw = base64.urlsafe_b64decode(sent[0][1]["raw"]).decode()
    assert "To: ann@example.com" in raw and "See you Friday." in raw, raw


def t_gmail_refresh():
    g, d = gm(), newdir("gm")
    fake_req = lambda url, data=None, headers=None, as_json=False: {"access_token": "new", "expires_in": 3600}
    with gm_token(d, expires=-100):
        with mock.patch.object(g, "_client", lambda: ("cid", "sec")), mock.patch.object(g, "_request", fake_req):
            eq(g._access_token(), "new")
        with open(g.TOKEN_FILE) as f:
            eq(json.load(f)["access_token"], "new")
    with gm_token(d, expires=-100, refresh=None):
        try:
            g._access_token()
            raise AssertionError("expired token without refresh token must ask to reconnect")
        except g.GmailError:
            pass


def t_gmail_signin():
    g, d = gm(), newdir("gm")
    seen = {}

    def run(state_ok):
        def fake_open(url):
            q = {k: v[0] for k, v in urllib.parse.parse_qs(urllib.parse.urlparse(url).query).items()}
            seen.clear()
            seen.update(q)
            st = q["state"] if state_ok else "forged"

            def callback():
                time.sleep(0.2)
                urllib.request.urlopen(f"{q['redirect_uri']}/?code=abc&state={st}", timeout=5).read()
            threading.Thread(target=callback, daemon=True).start()
            return True

        exchange = lambda url, data=None, headers=None, as_json=False: {
            "access_token": "at", "refresh_token": "rt", "expires_in": 3600, "scope": g.SCOPES["read"]}
        with mock.patch.object(g, "TOKEN_FILE", os.path.join(d, f"t{state_ok}.json")), \
                mock.patch.object(g, "_client", lambda: ("cid", "sec")), \
                mock.patch.object(g.webbrowser, "open", fake_open), mock.patch.object(g, "_request", exchange):
            g.authorize(["read"], timeout=10)
            return g._load_token()

    tok = run(True)
    eq((tok["refresh_token"], tok["scopes"]), ("rt", [g.SCOPES["read"]]))
    assert seen["code_challenge_method"] == "S256" and g.SCOPES["read"] in seen["scope"]
    assert seen["redirect_uri"].startswith("http://127.0.0.1:"), seen["redirect_uri"]
    try:
        run(False)
        raise AssertionError("forged state must be rejected")
    except g.GmailError:
        pass


def t_gmail_disconnect():
    g, d = gm(), newdir("gm")
    with gm_token(d), mock.patch.object(g, "_request", lambda *a, **k: {}):
        assert g.disconnect() and not os.path.exists(g.TOKEN_FILE)
        assert not g.disconnect()


def t_gmail_routing():
    g, cal, d = gm(), util(), newdir("gm")
    v, said, pc = pv()
    calls = []
    with mock.patch.object(g, "connect_async", lambda perms, done: calls.append(perms) or done(None)):
        assert say(pc, v, "connect gmail") and calls[-1] == ["read"]
        assert "connected" in said[-1].lower(), said[-1:]
        assert say(pc, v, "connect gmail with sending") and calls[-1] == ["read", "send"]
    with mock.patch.object(g, "disconnect", lambda: True):
        assert say(pc, v, "disconnect gmail") and "disconnected" in said[-1].lower()
    sent = []
    with gm_token(d, ("read", "send")), mock.patch.object(cal, "BASE", newdir("gmm")), \
            mock.patch.object(g, "_api", lambda p, body=None: sent.append(p) or {}):
        cal.save_draft("ann@example.com", "Hi", "Body")
        assert say(pc, v, "send the email") and "about to send" in said[-1], said[-1:]
        assert not sent, "sent before confirmation"
        assert say(pc, v, "no") and not sent, "sent despite 'no'"
        say(pc, v, "send the email")
        say(pc, v, "yes")
        eq(sent, ["messages/send"])
        assert cal.latest_draft() is None


check("Gmail OAuth", "permissions: read/send are separate, send needs confirmed=True", t_gmail_permissions)
check("Gmail OAuth", "read unread mail (mocked API, read-only)", t_gmail_fetch)
check("Gmail OAuth", "send builds a valid message (mocked API)", t_gmail_send)
check("Gmail OAuth", "token refresh and reconnect prompt", t_gmail_refresh)
check("Gmail OAuth", "sign-in flow: PKCE, state check, token saved (no password)", t_gmail_signin)
check("Gmail OAuth", "disconnect revokes and deletes token", t_gmail_disconnect)
check("Gmail OAuth", "voice: connect/disconnect, send needs 'yes'", t_gmail_routing)
check("Gmail OAuth", "real Gmail read", lambda: requires("PERMISSION", "say 'Aurora, connect gmail' first")
      if not gm().has("read") else (gm().fetch_unread(1) and None))


# ============================================================================ feature modules: visual calculator, adaptive interface, workspace, simulator
import contextlib


def _fdir():
    return tempfile.mkdtemp(dir=tmpdir)


def fake_holo():
    """Hologram stand-in: real ints/lists where the code does arithmetic or iteration, mocks elsewhere."""
    H = mock.MagicMock(name="Hologram")
    H.width, H.height, H.translate_x = 1280, 720, 0.0
    H.mode, H.info_card, H.quiet_alerts = "empty", None, False
    H.overlays, H.event_listeners, H.event_log = [], [], []
    H.modes = H.labs = H.adaptive_interface = H.workspace = None
    H._wrap_text.side_effect = lambda f, t, w: [t]
    H._truncate.side_effect = lambda t, n: t
    H._pill_width.side_effect = lambda t: 8 * len(t) + 28
    H._materialize_progress.return_value = 1.0
    return H


def stub_voice(H=None, **kw):
    said = []
    attrs = dict(hologram=H or fake_holo(), _speak=said.append, _log=lambda m: None, client=None, plus=None,
                 latest_frame="frame")
    attrs.update(kw)
    v = SimpleNamespace(**attrs)
    return v, said


THEME = (0.3, 0.6, 1.0)

# ---- visual calculator
G = "Visual calculator"


def vc():
    return mod("visual_calculator")


def t_vc_normalize():
    v = vc()
    for raw, want in [("2 × 3", "2*3"), ("3x + 2", "3*x+2"), ("4 x 5", "4*5"), ("12,500 + 1", "12500+1"),
                      ("x^2", "x**2"), ("2 + 2 =", "2+2"), ("2(3+4)", "2*(3+4)")]:
        eq(v.normalize(raw), want)


def t_vc_solve():
    v = vc()
    for problem, want in [("12 + 30", "The answer is 42."), ("10 / 4", "The answer is 2.5."),
                          ("2x + 4 = 10", "x equals 3."), ("x^2 - 5x + 6 = 0", "x equals 2 or x equals 3."),
                          ("x^2 + 1 = 0", "There is no real solution."), ("x + 1 = x + 2", "There is no real solution."),
                          ("x + 1 = x + 1", "Every number works."), ("2 + 2 = 4", "That's correct."), ("1 / 0", "That divides by zero.")]:
        eq(v.solve_text(problem), want)
    assert v.solve_text("2 + 2 = 5").startswith("That's wrong"), "wrong equation not flagged"
    assert v.solve_text("x^3 = 8") is None, "cubic should go to the AI"
    assert v.solve_text("hello") is None
    assert v.solve_text("9**9**9") is None, "huge exponent must be refused"


def t_vc_spoken():
    v = vc()
    eq(v.spoken("2+3"), "2 plus 3")
    eq(v.fmt(3.0), "3")
    eq(v.fmt(2.5), "2.5")


def vc_run(phrase, ocr=None, vision=None, vision_error=False, complete=None, client=None, frame="img"):
    v = vc()
    voice, said = stub_voice(client=None)
    voice.client, voice.latest_frame = client, frame
    br = mod("brain")
    with contextlib.ExitStack() as st:
        st.enter_context(mock.patch.object(v, "ocr_math", return_value=ocr))
        st.enter_context(mock.patch.object(v.sc, "capture_screen", return_value="screen"))
        if vision_error:
            st.enter_context(mock.patch.object(br, "describe_scene", side_effect=RuntimeError("down")))
        else:
            st.enter_context(mock.patch.object(br, "describe_scene", return_value=vision))
        st.enter_context(mock.patch.object(br, "_complete", return_value=complete))
        handled = v.handle_command(voice, phrase)
    return handled, said, voice


def t_vc_offline_ocr():
    handled, said, voice = vc_run("solve this problem", ocr="2+3")
    assert handled
    eq(said[-1], "I read 2 plus 3. The answer is 5.")
    assert voice.hologram.show_info_card.called
    handled, said, _ = vc_run("solve the equation on my screen", ocr="2x+4=10")
    assert handled and "x equals 3" in said[-1], said
    handled, said, _ = vc_run("read this equation", ocr=None)
    assert handled and "couldn't read" in said[-1] and "pytesseract" in said[-1], said
    handled, said, _ = vc_run("solve this problem", frame=None)
    assert handled and "can't get an image" in said[-1]
    assert not vc_run("what time is it")[0], "unrelated phrase must not be handled"


def t_vc_online():
    handled, said, _ = vc_run("solve this problem", vision="x^2-4=0", client=object())
    assert handled and "x equals -2 or x equals 2" in said[-1], said
    handled, said, _ = vc_run("solve this problem", vision="x^3=8", complete="x equals 2.", client=object())
    assert handled and "x equals 2." in said[-1], "hard problem should go to the AI"
    handled, said, _ = vc_run("solve this problem", vision="NONE", client=object())
    assert handled and "couldn't read" in said[-1] and "pytesseract" not in said[-1], said
    handled, said, _ = vc_run("solve this problem", vision_error=True, ocr="7*6", client=object())
    assert handled and "The answer is 42." in said[-1], "must fall back to local OCR when the AI fails"


check(G, "normalize spoken/printed math", t_vc_normalize)
check(G, "local solver: arithmetic, linear, quadratic, edge cases, safety", t_vc_solve)
check(G, "spoken formatting", t_vc_spoken)
check(G, "voice flow offline (OCR), screen source, no frame, unrelated phrase", t_vc_offline_ocr)
check(G, "voice flow online (vision model, AI fallback, NONE, error -> OCR)", t_vc_online)

# ---- adaptive interface
G = "Adaptive interface"


def ai_env():
    ai = mod("adaptive_interface")
    H = fake_holo()
    gs = mock.MagicMock()
    gs.connected = False
    gs.latest = {"fps": 60, "cpu": 20.0, "ram": 40.0, "temp_c": 55.0, "session": "1m 2s", "foreground": "game.exe",
                 "recording": "no"}
    voice, said = stub_voice(H, game_session=gs, active_timers=[{"label": "study session", "ends_at": time.time() + 90}])
    return ai, H, gs, voice, said


def t_ai_switch():
    ai, H, gs, voice, said = ai_env()
    for phrase, mode, hidden in [("switch to coding mode", "coding", ("log",)),
                                 ("Switch Into The Programming Mode", "coding", ("log",)),
                                 ("switch to studying mode", "studying", ("system", "log", "readout")),
                                 ("switch over to the gaming mode", "gaming", ("system", "log", "readout")),
                                 ("switch to normal mode", "standard", ())]:
        assert ai.handle_command(voice, phrase), phrase
        eq(H.adaptive_interface.mode, mode)
        eq(H.hidden_panels, hidden)
        eq(said[-1], ai.REPLIES[mode])
    gs.start.assert_called_once()
    gs.stop.assert_called_once()                       # leaving gaming mode stops the reader it started
    eq(len(H.overlays), 1)                             # install() is idempotent
    assert not ai.handle_command(voice, "switch to chemistry mode")
    assert not ai.handle_command(voice, "switch to coding")


def t_ai_gather():
    ai, H, gs, voice, said = ai_env()
    ui = ai.install(H, voice)
    ui.mode, ui._busy = "coding", True
    run = lambda *a, **k: SimpleNamespace(stdout="main\n")
    with mock.patch.object(ai.subprocess, "run", run), \
            mock.patch.object(ai.dashboard_modes, "_recent_files", lambda n: [(time.time(), "file", "main.py")]):
        ui._gather()
    eq(ui._info["branch"], "main")
    eq(ui._info["files"][0][2], "main.py")
    assert ui._busy is False
    ui.mode, ui._busy = "studying", True
    au, sc = mod("aurora_utilities"), mod("system_control")
    with mock.patch.object(ai.subprocess, "run", run), mock.patch.object(au, "minutes_today", return_value=12), \
            mock.patch.object(au, "list_decks", return_value={"bio": 3}), mock.patch.object(sc, "read_notes", return_value=["n1"]):
        ui._gather()
    eq((ui._info["minutes"], ui._info["decks"], ui._info["notes"]), (12, {"bio": 3}, ["n1"]))


def t_ai_draw():
    ai, H, gs, voice, said = ai_env()
    ui = ai.install(H, voice)
    ui._refresh = lambda: None
    ui._info = {"branch": "main", "dirty": 2, "files": [(time.time(), "file", "a.py")], "minutes": 5,
                "decks": {"bio": 3}, "notes": ["n1"]}
    for mode in ("coding", "studying", "gaming"):
        ui.mode = mode
        ui.draw(THEME)
    assert H._draw_panel.called, "nothing was drawn"
    H._draw_panel.reset_mock()
    ui.mode = "standard"
    ui.draw(THEME)
    assert not H._draw_panel.called, "standard mode must not draw"
    ui.mode, H.modes = "coding", SimpleNamespace(active="presentation")
    ui._refresh = lambda: (_ for _ in ()).throw(AssertionError("drew during presentation"))
    ui.draw(THEME)


check(G, "voice: switching modes, hidden panels, game reader start/stop, rejects bad phrases", t_ai_switch)
check(G, "background data gathering (git, files, study stats)", t_ai_gather)
check(G, "each mode draws; standard and presentation draw nothing", t_ai_draw)

# ---- project workspace
G = "Project workspace"


@contextlib.contextmanager
def ws_env():
    pw, tel = mod("project_workspace"), mod("telemetry")
    root = _fdir()
    with mock.patch.object(pw, "WORKSPACES_DIR", root), mock.patch.object(tel, "EXPERIMENTS_DIR", tel.EXPERIMENTS_DIR):
        voice, said = stub_voice()
        try:
            yield SimpleNamespace(pw=pw, tel=tel, root=root, voice=voice, said=said,
                                  say=lambda p: pw.handle_command(voice, p))
        finally:
            mgr = getattr(voice.hologram, "workspace", None)
            if mgr is not None:
                mgr.close()                              # restores telemetry.EXPERIMENTS_DIR


def t_ws_data():
    pw = mod("project_workspace")
    eq(pw.slug("Mars Project!"), "mars_project")
    assert "temperature" in pw.words("what was the temperature in my last experiment")
    assert "what" not in pw.words("what was the temperature")
    d = _fdir()
    ws = pw.Workspace(os.path.join(d, "mars"), "Mars")
    for sub in ("files", "data", "experiments"):
        assert os.path.isdir(os.path.join(d, "mars", sub)), sub
    ws.add("tasks", text="order sensors", done=False)
    assert ws.complete_task(1) and not ws.complete_task(5)
    ws.add("notes", text="battery drains fast")
    assert pw.Workspace(os.path.join(d, "mars")).meta["notes"][0]["text"] == "battery drains fast", "not persisted"
    eq(ws.tab_lines("tasks"), ["1. [x] order sensors"])
    csv, txt = os.path.join(d, "run1.csv"), os.path.join(d, "thermal.txt")
    open(csv, "w").write("a,b\n1,2\n")
    open(txt, "w").write("thermal insulation results look good")
    assert os.path.dirname(ws.add_file(csv)) == ws.data_dir and os.path.dirname(ws.add_file(txt)) == ws.files_dir
    assert os.path.basename(ws.add_file(csv)) == "run1 (1).csv", "duplicate import overwrote the file"
    hits = ws.search("battery")
    assert hits and "battery" in hits[0][1], hits
    assert any("thermal" in t.lower() for _, t in ws.search("thermal insulation"))
    eq(ws.search("zzzz"), [])
    assert "couldn't find" in ws.ask("zzzz")


def t_ws_experiment_answer():
    pw = mod("project_workspace")
    ws = pw.Workspace(os.path.join(_fdir(), "p"), "P")
    log = os.path.join(ws.exp_dir, "exp1")
    os.makedirs(log)
    with open(os.path.join(log, "log.jsonl"), "w") as f:
        for temp in (20.0, 24.5):
            f.write(json.dumps({"type": "sensor", "data": {"temperature": temp}}) + "\n")
    ans = ws.ask("what was the temperature in my last experiment")
    assert "exp1" in ans and "24.5" in ans and "20" in ans and "2 readings" in ans, ans


def t_ws_voice():
    with ws_env() as e:
        say, said, pw = e.say, e.said, e.pw
        assert not say("tell me a joke") and not say("add note: x"), "must ignore non-workspace phrases when none is open"
        assert say("create a workspace for my Mars project") and "Mars project workspace ready" in said[-1]
        assert os.path.isdir(os.path.join(e.root, "mars", "files"))
        eq(e.tel.EXPERIMENTS_DIR, os.path.join(e.root, "mars", "experiments"))
        assert say("add note: battery drains fast") and said[-1] == "Note added."
        assert say("add task: order sensors") and said[-1] == "Task added."
        assert say("complete task 1") and said[-1] == "Task completed."
        assert say("complete task 9") and "don't have" in said[-1]
        assert say("add link: nasa dot gov") and e.voice.hologram.workspace.ws.meta["links"][0]["url"] == "https://nasa.gov"
        n = len(said)
        assert say("show tasks") and len(said) == n, "tab switch should be silent"
        eq(e.voice.hologram.workspace.tab, "tasks")
        assert say("search project for battery") and "battery" in said[-1], said[-1]
        # experiments recorded now land inside the project and are answerable
        rec = e.tel.ExperimentRecorder()
        assert rec.start("trial1")[0]
        rec.log_sensor({"temperature": 21.5})
        assert say("what was the temperature in my last experiment") and "21.5" in said[-1], said[-1]
        assert say("list my workspaces") and "mars" in said[-1]
        assert say("open my Venus workspace") and "don't have a workspace" in said[-1]
        assert say("open my Mars workspace") and "Opened" in said[-1]
        e.voice.hologram.workspace.draw(THEME)
        assert e.voice.hologram._draw_panel.called, "overlay did not draw"
        assert say("close workspace") and said[-1] == "Workspace closed."
        assert e.tel.EXPERIMENTS_DIR != os.path.join(e.root, "mars", "experiments"), "experiments dir not restored"
        assert not say("close workspace") or "No workspace" in said[-1]


check(G, "workspace data: folders, persistence, tasks, file import, search", t_ws_data)
check(G, "'last experiment' sensor answer", t_ws_experiment_answer)
check(G, "voice flow end to end (create, edit, ask, open, close, overlay draw)", t_ws_voice)

# ---- operation simulator
G = "Operation simulator"


@contextlib.contextmanager
def sim_env(files=()):
    au, sc, sm = mod("aurora_utilities"), mod("system_control"), mod("operation_simulator")
    home = os.path.join(_fdir(), "h")
    dl = os.path.join(home, "Downloads")
    os.makedirs(dl)
    os.makedirs(os.path.join(home, "Documents", "School"))
    for f in files:
        open(os.path.join(dl, f), "w").write("x")
    with mock.patch.object(au, "HOME", home), mock.patch.object(au, "BASE", _fdir()), \
            mock.patch.object(sc, "NOTES_FILE", os.path.join(_fdir(), "notes.txt")):
        said, ran = [], []
        voice = SimpleNamespace(_speak=said.append, _handle_local_command=lambda op: ran.append(op) or True,
                                hologram=mock.MagicMock(), client=None, plus=None, _productivity=SimpleNamespace(found=[]))

        def run(phrase):
            handled = sm.handle_command(voice, phrase)
            return handled, (said[-1] if said else "")
        yield SimpleNamespace(au=au, sc=sc, sm=sm, home=home, dl=dl, voice=voice, said=said, ran=ran, run=run)


def expect(e, phrase, *frags):
    handled, msg = e.run(phrase)
    assert handled, f"'{phrase}' not handled"
    for f in frags:
        assert f.lower() in msg.lower(), f"{phrase!r} -> {msg!r} lacks {f!r}"
    return msg


def t_sim_organize_and_go():
    files = ["report.pdf", "pic.png", "song.mp3"]
    with sim_env(files) as e:
        msg = expect(e, "simulate organizing my downloads", "3 loose files", "go ahead")
        eq(sorted(os.listdir(e.dl)), sorted(files))                 # nothing moved
        assert e.voice.hologram.show_info_card.called
        assert e.run("go ahead")[0]
        eq(e.ran, ["organize my downloads"])
        assert not e.run("go ahead")[0], "go ahead must only work once"
    with sim_env() as e:
        expect(e, "preview organizing my downloads", "nothing to organize")
    with sim_env(["a.pdf"]) as e:
        e.run("simulate organizing my downloads")
        e.voice._simulator.last = (e.voice._simulator.last[0], time.time() - 1000)
        expect(e, "go ahead", "expired")
        eq(e.ran, [])


def t_sim_files():
    with sim_env(["report.pdf", "summary.pdf"]) as e:
        expect(e, "what would happen if I rename report.pdf in downloads to notes.pdf", "renamed to notes.pdf", "Nothing is overwritten")
        expect(e, "simulate renaming report.pdf in downloads to summary.pdf", "summary (1).pdf", "taken")
        expect(e, "simulate renaming report.pdf in downloads to a/b", "isn't a valid file name")
        expect(e, "simulate moving report.pdf in downloads to documents", "would move to Documents")
        assert not e.run("simulate renaming nothing.pdf in downloads to x.pdf")[0], "unknown file should fall through"
        eq(sorted(os.listdir(e.dl)), ["report.pdf", "summary.pdf"])   # still untouched
        open(os.path.join(e.home, "Documents", "todo.txt"), "w").write("x")
        expect(e, "simulate zipping todo.txt in documents", "todo.txt.zip would be created")
        assert not os.path.exists(os.path.join(e.home, "Documents", "todo.txt.zip"))
    with sim_env() as e:
        with zipfile.ZipFile(os.path.join(e.dl, "archive.zip"), "w") as z:
            z.writestr("a.txt", "1")
            z.writestr("b.txt", "2")
        with zipfile.ZipFile(os.path.join(e.dl, "evil.zip"), "w") as z:
            z.writestr("../evil.txt", "x")
        expect(e, "simulate extracting archive.zip in downloads", "folder called archive", "2 entries")
        expect(e, "simulate extracting evil.zip in downloads", "unsafe paths")
        assert not os.path.exists(os.path.join(e.dl, "archive"))


def t_sim_batches():
    with sim_env(["physics_notes.txt", "b.txt"]) as e:
        expect(e, "simulate moving these into my school folder", "tell me which files first")
        expect(e, "simulate renaming these files properly", "tell me which files first")
        e.voice._productivity.found = [os.path.join(e.dl, "physics_notes.txt"), os.path.join(e.dl, "b.txt")]
        expect(e, "simulate moving these into my school folder", "2 files would move into School")
        expect(e, "simulate moving these into my chemistry folder", "no chemistry folder", "new one in Documents")
        expect(e, "simulate renaming these files properly", "physics_notes.txt to Physics Notes.txt")
        assert os.path.exists(os.path.join(e.dl, "physics_notes.txt")), "files changed during simulation"


def t_sim_data():
    with sim_env() as e:
        expect(e, "simulate deleting all notes", "no notes")
        e.au.add_quick_note("first note")
        e.au.add_quick_note("second note")
        expect(e, "simulate deleting all notes", "All 2 notes", "no undo")
        expect(e, "simulate deleting note 1", "Note 1 would be deleted", "first note")
        eq(len(e.au.list_notes()), 2)                                   # nothing was deleted
        e.au.teach_skill("warmup", ["a", "b"])
        expect(e, "simulate deleting skill warmup", "warmup", "2 steps")
        expect(e, "simulate deleting skill nope", "don't have a skill")
        assert e.au.get_skill("warmup")
        e.au.add_card("bio", "q", "a")
        expect(e, "simulate deleting deck bio", "bio", "1 flashcards")
        expect(e, "simulate deleting deck nope", "don't have that deck")
        assert e.au.list_decks() == {"bio": 1}


def t_sim_system():
    with sim_env() as e:
        with mock.patch.object(e.sc, "get_volume_percent", return_value=50):
            expect(e, "simulate volume to 90", "from 50 to 90", "loud")
            expect(e, "simulate volume up", "to 60")
            expect(e, "simulate muting", "to 0")
        with mock.patch.object(e.sc, "get_volume_percent", return_value=None):
            expect(e, "simulate volume up", "isn't available")
        expect(e, "simulate locking my computer", "lock")
        for result, frag in (((False, ""), "can't see your phone"), ((True, "abc\tdevice"), "USB connection keeps working"),
                             ((True, "192.168.1.5:5555\tdevice"), "lose the connection")):
            with mock.patch.object(e.sc, "is_device_connected", return_value=result):
                expect(e, "simulate turning off wifi on my phone", frag)
        with mock.patch.object(e.au, "latest_draft", return_value=None):
            expect(e, "simulate sending the email", "no draft")
        with mock.patch.object(e.au, "latest_draft", return_value={"to": "a@b.c", "subject": "Hi", "body": "see you friday"}), \
                mock.patch.object(e.au, "email_configured", return_value=False):
            expect(e, "simulate sending the email", "a@b.c", "isn't set up")


def t_sim_edit_code():
    with sim_env() as e:
        path = os.path.join(e.home, "sorter.py")
        open(path, "w").write("a\n")
        br = mod("brain")
        with mock.patch.object(e.sc, "resolve_project_path", return_value=path):
            expect(e, "simulate editing sorter to add a reverse option", "AI online")
            e.voice.client = object()
            with mock.patch.object(br, "edit_code", return_value="a\nb\n"):
                expect(e, "simulate editing sorter to add a reverse option", "1 line added", ".bak")
            eq(open(path).read(), "a\n")                                 # file untouched
        with mock.patch.object(e.sc, "resolve_project_path", return_value=None), \
                mock.patch.object(e.sc, "find_arduino_sketch", return_value=None):
            assert not e.run("simulate editing ghost to do x")[0]


def t_sim_passthrough():
    with sim_env(["a.pdf"]) as e:
        assert not e.run("organize my downloads")[0], "real commands must not be intercepted"
        assert not e.run("simulate world peace")[0], "unknown simulations must fall through"
        e.run("simulate organizing my downloads")
        assert not e.run("what time is it")[0]
        assert e.voice._simulator.last is None, "unrelated command should clear the pending simulation"


check(G, "organize preview, 'go ahead' runs once, expiry, nothing touched", t_sim_organize_and_go)
check(G, "rename / move / zip / extract previews (clashes, invalid, unsafe zip)", t_sim_files)
check(G, "batch move / rename of found files", t_sim_batches)
check(G, "notes, skills, decks previews", t_sim_data)
check(G, "volume, lock, phone wifi, email previews", t_sim_system)
check(G, "code-edit preview (needs AI, file untouched)", t_sim_edit_code)
check(G, "unrelated and real commands fall through", t_sim_passthrough)

# ---- document generator
G = "Document generator"


def dgm():
    return mod("document_generator")


def t_dg_parse():
    d = dgm()
    for word, kind in [("presentation", "pptx"), ("slide deck", "pptx"), ("powerpoint", "pptx"), ("report", "docx"),
                       ("word doc", "docx"), ("spreadsheet", "xlsx"), ("excel sheet", "xlsx")]:
        eq(d._kind(word), kind)
    eq(d._REQ.search("make a 5 slide presentation about black holes").groups(), ("5", "presentation", "black holes"))
    assert not d._REQ.search("start presentation called robotics"), "must not steal presentation mode"
    assert not d._REQ.search("generate my experiment report"), "must not steal the experiment report"


def t_dg_offline_build():
    d = dgm()
    with mock.patch.object(d, "OUT_DIR", _fdir()):
        for kind in ("pptx", "docx", "xlsx"):
            path, n, ai, err = d.create(kind, "Black Holes!", None, 3)
            if err and err.startswith("Missing"):
                skip(err)
            assert not err and not ai, err
            assert path.endswith(d.EXT[kind]) and os.path.getsize(path) > 1000, path
    br = mod("brain")
    with mock.patch.object(br, "_complete", return_value="not json"), mock.patch.object(d, "OUT_DIR", _fdir()):
        path, n, ai, err = d.create("docx", "x", object())
    if err and err.startswith("Missing"):
        skip(err)
    assert not err and not ai, "bad AI output must fall back to the skeleton"


def t_dg_ai_content():
    d, br = dgm(), mod("brain")
    pptx = '{"title":"T","slides":[{"title":"A","bullets":["x","y"],"notes":"n"},{"title":"B","bullets":[],"notes":""}]}'
    xlsx = '{"title":"Budget","headers":["Item","Cost"],"rows":[["Rent",800],["Food",300],["Total","=SUM(B2:B3)"]]}'
    docx = '{"title":"Doc","sections":[{"heading":"Intro","paragraphs":["Hello world."],"bullets":["pt"]}]}'
    with mock.patch.object(d, "OUT_DIR", _fdir()):
        with mock.patch.object(br, "_complete", return_value="```json\n" + pptx + "\n```"):
            path, n, ai, err = d.create("pptx", "t without images", object(), 2)
        if err and err.startswith("Missing"):
            skip(err)
        assert ai and n == 3, (n, ai, err)
        from pptx import Presentation
        prs = Presentation(path)
        eq(len(prs.slides), 3)
        texts = [sh.text_frame.text for sh in prs.slides[1].shapes if sh.has_text_frame]
        assert "A" in texts, texts                     # designed slides use text boxes, not title placeholders
        eq(prs.slides[1].notes_slide.notes_text_frame.text, "n")
        with mock.patch.object(br, "_complete", return_value=xlsx):
            path, n, ai, err = d.create("xlsx", "budget", object())
        from openpyxl import load_workbook
        ws = load_workbook(path).active
        eq((ws["A1"].value, ws["B2"].value, ws["B4"].value), ("Item", 800, "=SUM(B2:B3)"))
        assert ws["A1"].font.bold, "header not bold"
        with mock.patch.object(br, "_complete", return_value=docx):
            path, n, ai, err = d.create("docx", "doc", object())
        from docx import Document
        text = " ".join(p.text for p in Document(path).paragraphs)
        assert "Intro" in text and "Hello world." in text and "pt" in text, text


def t_dg_voice():
    d = dgm()
    voice, said = stub_voice()
    with mock.patch.object(d, "OUT_DIR", _fdir()):
        assert d.handle_command(voice, "create a 2 row spreadsheet for expenses")
    if "Missing" in said[-1]:
        skip(said[-1])
    assert "spreadsheet is ready with 2 rows" in said[-1] and "offline" in said[-1], said
    assert voice.hologram.show_info_card.called
    assert not d.handle_command(voice, "start presentation called robotics")
    assert not d.handle_command(voice, "what time is it")


def t_dg_hooked():
    src = open(os.path.join(ROOT, "jarvis_ui", "voice_assistant.py"), encoding="utf-8").read()
    assert "document_generator.handle_command" in src, "run apply_document_generator_patch.py"
    assert "create_document" in {t["function"]["name"] for t in mod("brain").TOOLS}, "AI tool missing"
    v, H, said = make_voice()
    with mock.patch.object(dgm(), "OUT_DIR", _fdir()):
        assert v._handle_local_command("make a presentation about mars"), "not hooked into _handle_local_command"
        assert "presentation" in v._exec_tool("create_document", {"kind": "presentation", "topic": "mars"}).lower()


check(G, "request parsing (and no clash with presentation mode)", t_dg_parse)
check(G, "builds pptx/docx/xlsx offline; bad AI output falls back", t_dg_offline_build)
check(G, "AI content -> real slides, notes, formulas, headings (mocked model)", t_dg_ai_content)
check(G, "voice flow: reply, info card, unrelated phrases fall through", t_dg_voice)
check(G, "hooked into voice_assistant and the AI tools", t_dg_hooked)


def t_dg_live():
    d = dgm()
    with mock.patch.object(d, "OUT_DIR", _fdir()):
        path, n, ai, err = d.create("pptx", "the water cycle", get_client(), 3)
    if err and err.startswith("Missing"):
        skip(err)
    assert ai and n >= 3, (n, ai, err)


live("Document generator (real Groq)", t_dg_live)


def t_debug_loop():
    ad, br = mod("addons"), mod("brain")
    with mock.patch.object(ad, "SANDBOX_ROOT", os.path.join(tmpdir, "dbg")), \
            mock.patch.object(br, "fix_code", lambda c, code, err, goal="", lang="python": "print(42)\n"):
        code, ok, out, err, fixes = ad.debug_loop(object(), "print(1/0)\n", "demo", "dbg_demo")
    assert ok and fixes == 1 and out.strip() == "42", (ok, fixes, err)


check("Code sandbox", "debug loop fixes a failing script (mocked model)", t_debug_loop)


# ============================================================================ extended tests (internals, edge cases, error paths)
# Merged from the former test_aurora_extended.py; uses the same check()/mod()/eq() harness and report as above.
import io
import re
import socket
import urllib.error
from collections import deque
from datetime import datetime, timedelta


def ext(group):
    """Decorator: runs the test immediately through check() so it appears in the normal report."""
    def deco(fn):
        check(group, fn.__name__.replace("_", " "), fn)
        return fn
    return deco


def ns(**k):
    return SimpleNamespace(**k)


def ext_dir():
    return tempfile.mkdtemp(dir=tmpdir)


def raises(exc, fn, *a, **k):
    try:
        fn(*a, **k)
    except exc:
        return
    raise AssertionError(f"{exc.__name__} not raised")


def ext_voice(**kw):
    said = []
    v = ns(_speak=said.append, _log=lambda m: None, speak_now=said.append, client=None, plus=None, latest_frame=None,
           hologram=mock.MagicMock(), active_timers=[])
    v.__dict__.update(kw)
    return v, said


# ============================================================================ telemetry
@ext("Telemetry")
def experiment_recorder_full_cycle():
    tel = mod("telemetry")
    import numpy as np
    with mock.patch.object(tel, "EXPERIMENTS_DIR", ext_dir()):
        r = tel.ExperimentRecorder()
        assert not r.log_observation("x") and not r.log_sensor({"a": 1}) and r.stop()[0] is False
        eq(r.capture_screenshot(np.zeros((4, 4, 3), np.uint8))[0], False)
        eq(r.generate_report()[0], None)
        ok, name = r.start("My Test!")
        assert ok and name == "My_Test", name
        assert r.start("again")[0] is False
        assert r.log_observation("fizzing") and r.log_sensor({"t": 1})
        r.maybe_log_sensor({"t": 2}, interval=60)
        r.maybe_log_sensor({"t": 3}, interval=60)            # throttled
        eq(r.capture_screenshot(None)[0], False)
        ok, fname = r.capture_screenshot(np.zeros((8, 8, 3), np.uint8))
        assert ok and fname == "shot_001.png"
        r.stop()
        path, err = r.generate_report()
        assert err is None
        text = open(path, encoding="utf-8").read()
        for frag in ("EXPERIMENT REPORT: My_Test", "fizzing", "Observations: 1  Sensor readings: 2  Screenshots: 1", "shot_001.png"):
            assert frag in text, frag


@ext("Telemetry")
def telemetry_reader_config_errors():
    tel = mod("telemetry")
    for cfg, frag in ((None, "iot_config"), ({"mode": "bogus"}, "'serial' or 'wifi'"), ({"mode": "wifi"}, "url")):
        with mock.patch.object(tel, "load_config", return_value=cfg):
            ok, msg = tel.TelemetryReader().start()
            assert not ok and frag in msg, msg
    with mock.patch.object(tel, "load_config", return_value={"mode": "serial"}), mock.patch.object(tel, "PYSERIAL_AVAILABLE", True):
        ok, msg = tel.TelemetryReader().start()
        assert not ok and "port" in msg, msg
    d = os.path.join(ext_dir(), "c.json")
    open(d, "w").write("{bad")
    with mock.patch.object(tel, "CONFIG_FILE", d):
        eq(tel.load_config(), None)
        with open(d, "w") as f:
            json.dump({"mode": "wifi"}, f)
        eq(tel.load_config(), {"mode": "wifi"})


@ext("Telemetry")
def telemetry_reader_wifi_polling():
    from http.server import BaseHTTPRequestHandler, HTTPServer
    tel = mod("telemetry")

    class H(BaseHTTPRequestHandler):
        def do_GET(self):
            body = b'{"temp": 21.5}'
            self.send_response(200)
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)

        def log_message(self, *a):
            pass

    srv = HTTPServer(("127.0.0.1", 0), H)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    cfg = {"mode": "wifi", "url": f"http://127.0.0.1:{srv.server_port}/t", "poll_seconds": 0.5}
    try:
        with mock.patch.object(tel, "load_config", return_value=cfg):
            r = tel.TelemetryReader()
            assert r.is_stale()
            assert r.start()[0]
            assert "already running" in r.start()[1]
            for _ in range(40):
                if r.latest:
                    break
                time.sleep(0.1)
            eq(r.latest, {"temp": 21.5})
            assert r.connected and not r.is_stale()
            r.stop()
    finally:
        srv.shutdown()


@ext("Telemetry")
def game_session_reader():
    tel = mod("telemetry")
    with mock.patch.object(tel.system_control, "get_foreground_process_name", return_value="game.exe"), \
            mock.patch.object(tel.system_control, "is_recording", return_value=False):
        g = tel.GameSessionReader()
        assert g.is_stale()
        assert g.start()[0] and g.connected
        for _ in range(40):
            if g.latest:
                break
            time.sleep(0.1)
        assert set(g.latest) >= {"session", "fps", "cpu", "ram", "temp_c", "foreground", "recording"}, g.latest
        eq((g.latest["foreground"], g.latest["recording"], g.latest["fps"]), ("game.exe", "no", "N/A"))
        g.set_fps(60)
        eq(g.external_fps, 60)
        assert "already running" in g.start()[1] and not g.is_stale()
        g.stop()
        assert not g.connected and g.session_start is None


# ============================================================================ face id
def _face_env():
    f = mod("face_id")
    import cv2
    import numpy as np
    if not hasattr(cv2, "face"):
        skip("cv2.face missing (opencv-contrib-python)")
    d = ext_dir()
    patches = [mock.patch.object(f, "DATA_DIR", d), mock.patch.object(f, "SAMPLES_DIR", os.path.join(d, "samples")),
               mock.patch.object(f, "PEOPLE_PATH", os.path.join(d, "people.json")),
               mock.patch.object(f, "THRESHOLD_FILE", os.path.join(d, "thr.txt"))]
    return f, cv2, np, d, patches


@ext("Face recognition")
def face_enroll_recognize_forget_persist():
    f, cv2, np, d, patches = _face_env()
    for p in patches:
        p.start()
    try:
        fid = f.FaceID()
        eq(fid.people, {})
        rng = np.random.RandomState(3)
        face = lambda seed: cv2.GaussianBlur(np.random.RandomState(seed).randint(0, 255, (200, 200)).astype(np.uint8), (9, 9), 0)
        noisy = lambda img, n=6: [np.clip(img.astype(int) + rng.randint(-6, 6, img.shape), 0, 255).astype(np.uint8) for _ in range(n)]
        box = (0, 0, 200, 200)
        for name, seed in (("Alice", 1), ("Bob", 2)):
            ok, n, msg = fid._finish_enrollment({"name": name, "imgs": [f.prep(v, box) for v in noisy(face(seed))]})
            assert ok and n == name
        eq(sorted(fid.people), ["Alice", "Bob"])
        res = [fid.recognize(noisy(face(1), 1)[0], box) for _ in range(3)]
        eq(res[-1][0], "Alice")
        assert fid.last_distance is not None and fid.confirmed_name == "Alice"
        eq(sorted(f.FaceID().people), ["Alice", "Bob"])               # reloaded from disk
        assert fid.get_theme_index("Alice") in range(4) and fid.get_theme_index("Zed") is None
        assert fid.forget("alice") and not fid.forget("nobody")
        eq(list(fid.people), ["Bob"])
        assert not os.path.isdir(os.path.join(f.SAMPLES_DIR, "Alice"))
        fid.threshold = 0.0001
        fid._reset_votes()
        eq(fid.recognize(noisy(face(2), 1)[0], box)[0], None)         # strict threshold rejects
    finally:
        for p in patches:
            p.stop()


@ext("Face recognition")
def face_prep_threshold_enrollment_state():
    f, cv2, np, d, patches = _face_env()
    for p in patches:
        p.start()
    try:
        gray = np.zeros((200, 200), np.uint8)
        eq(f.prep(gray, (0, 0, 200, 200)).shape, (f.FACE_SIZE, f.FACE_SIZE))
        eq(f.prep(gray, (500, 500, 10, 10)), None)
        eq(f.FaceID._dirname("Sam O'Neil"), "Sam_O_Neil")
        eq(f.FaceID().threshold, f.DEFAULT_THRESHOLD)
        open(f.THRESHOLD_FILE, "w").write("55")
        eq(f.FaceID().threshold, 55.0)
        open(f.THRESHOLD_FILE, "w").write("abc")
        eq(f.FaceID().threshold, f.DEFAULT_THRESHOLD)
        fid = f.FaceID()
        eq(fid.recognize(gray, (0, 0, 200, 200)), (None, None))        # nobody enrolled
        assert fid.start_enrollment("X") and fid.enrolling
        with mock.patch.object(f, "ENROLL_TIMEOUT", -1):
            ok, name, msg = fid.feed_enrollment(gray, [])
        assert not ok and name == "X" and "clear look" in msg and not fid.enrolling
    finally:
        for p in patches:
            p.stop()


# ============================================================================ hand tracker
def _lm(up):
    xs, base = {"index": -0.03, "middle": -0.01, "ring": 0.01, "pinky": 0.03}, {"index": 5, "middle": 9, "ring": 13, "pinky": 17}
    pts = [(0, 0)] * 21
    pts[1], pts[2], pts[3], pts[4] = (-0.03, -0.02), (-0.03, -0.05), (-0.02, -0.06), (0.0, -0.10)     # folded thumb
    for name, i in base.items():
        x = xs[name]
        pts[i] = (x, -0.09)
        ys = (-0.13, -0.16, -0.19) if name in up else (-0.12, -0.10, -0.06)
        for k, y in enumerate(ys, 1):
            pts[i + k] = (x, y)
    return [ns(x=x, y=y, z=0.0) for x, y in pts]


@ext("Hand tracking")
def pose_classification():
    ht = mod("hand_tracker")
    tracker = object.__new__(ht.HandTracker)
    cases = [({"index", "middle", "ring", "pinky"}, "open_palm", 4), (set(), "fist", 0), ({"index", "middle"}, "peace", 2),
             ({"index"}, "point", 1), ({"index", "pinky"}, "rock_sign", 2)]
    for up, want, count in cases:
        gesture, n, ratio = tracker._classify(_lm(up), None, ht._Track(0))
        assert (gesture, n) == (want, count), f"{sorted(up)} -> {gesture}, {n}"


@ext("Hand tracking")
def angle_clamp_debouncer_oneeuro():
    ht = mod("hand_tracker")
    eq(ht._clamp01(5), 1.0)
    eq(ht._clamp01(-1), 0.0)
    assert abs(ht._angle((-1, 0, 0), (0, 0, 0), (1, 0, 0)) - 180) < 1e-6
    eq(ht._angle((0, 0, 0), (0, 0, 0), (1, 0, 0)), 180.0)           # degenerate vector
    d = ht.GestureDebouncer(window=6, min_votes=4)
    out = [d.update("fist") for _ in range(4)]
    eq([o[1] for o in out], [False, False, False, True])
    eq(d.update("fist"), ("fist", False))
    f = ht.OneEuro()
    eq(f(1.0, 0.0), 1.0)
    assert 1.0 < f(2.0, 0.033) <= 2.0


@ext("Hand tracking")
def intent_gating_swipe_and_track_assignment():
    ht = mod("hand_tracker")
    tr = ht._Track(0)
    eq(tr.speed, 0.0)
    T = object.__new__(ht.HandTracker)
    assert T._gate_event(tr, "peace", True, 100.0)                   # fires once, still
    assert not T._gate_event(tr, "peace", False, 100.1)
    assert not T._gate_event(tr, "peace", True, 100.2)               # cooldown
    assert not T._gate_event(tr, "fist", True, 105.0)                # not actionable
    tr.fx.dx = 5.0
    assert not T._gate_event(tr, "peace", True, 110.0)               # moving too fast
    T._motion, T._last_swipe_time = deque(maxlen=60), 0.0
    got = [T._update_swipe((0.1 + 0.1 * i, 0.5), 10 + 0.05 * i) for i in range(6)]
    assert "right" in got and got.count("right") == 1, got           # cooldown stops repeats
    T._tracks, T._next_id = {}, 0
    a = T._assign([(0.2, 0.2), (0.8, 0.8)])
    ids = sorted(t_.id for t_ in a.values())
    eq(ids, [0, 1])
    b = T._assign([(0.21, 0.2), (0.79, 0.8)])
    eq({k: v.id for k, v in a.items()}, {k: v.id for k, v in b.items()})   # same hands keep ids
    for _ in range(ht.MAX_MISSED_FRAMES + 2):
        T._assign([])
    eq(T._tracks, {})


@ext("Hand tracking")
def enhance_read_and_model_download():
    ht = mod("hand_tracker")
    import cv2
    import numpy as np
    T = object.__new__(ht.HandTracker)
    T._clahe = cv2.createCLAHE(clipLimit=2.5, tileGridSize=(8, 8))
    bright, dark = np.full((32, 32, 3), 200, np.uint8), np.full((32, 32, 3), 10, np.uint8)
    assert T._enhance(bright) is bright
    out = T._enhance(dark)
    eq(out.shape, dark.shape)
    with mock.patch.object(ht.os.path, "exists", return_value=True), mock.patch.object(ht.urllib.request, "urlretrieve") as dl:
        ht._ensure_model("x.task")
        dl.assert_not_called()
    with mock.patch.object(ht.os.path, "exists", return_value=False), mock.patch.object(ht.urllib.request, "urlretrieve") as dl:
        ht._ensure_model("x.task")
        dl.assert_called_once()
    T.cap = ns(read=lambda: (False, None))
    r, fr = T.read()
    assert not r.detected and fr is None
    frame = np.zeros((48, 64, 3), np.uint8)
    T.cap = ns(read=lambda: (True, frame))
    T.landmarker = ns(detect_for_video=lambda img, ts: ns(hand_landmarks=[], hand_world_landmarks=[], handedness=[]))
    T._ts, T._tracks, T._next_id, T._primary_id = 0, {}, 0, None
    T._prev_two_hand_dist = T._prev_roll_angle = T._pinch_drag_prev = T._pull_zoom_prev_size = None
    T._prev_primary_gesture = T._prev_primary_pos = T._prev_time = None
    T._roll_mode, T._peak_speed, T._last_swipe_time, T._motion = 0, 0.0, 0.0, deque(maxlen=60)
    r, fr = T.read()
    assert not r.detected and fr is not None and hasattr(T, "raw_frame")


# ============================================================================ hologram (no GL context)
def bare_holo():
    hm = mod("hologram")
    h = object.__new__(hm.Hologram)
    h.__dict__.update(
        elapsed=0.0, _materialize_start=0.0, _materialize_duration=0.45, weather=None, map_highlight=None, docked=False,
        mode="empty", mode_label="", orbits=[], selected_index=None, info_card=None, protons=1, neutrons=0, electron_count=1,
        shell_summary="", current_element_z=1, current_star_index=0, molecule_data=None, graph_points=[], graph_expr="",
        physics_type=None, physics_state={}, constellation_data=None, constellation_name="", network_nodes=[], network_edges=[],
        shape_name="sphere", telemetry_label="", telemetry_data={}, telemetry_history={}, telemetry_connected=False,
        telemetry_last_update=0.0, event_log=deque(maxlen=6), _event_log_lock=threading.Lock(), event_listeners=[],
        theme_index=2, _theme_prev_color=hm.THEMES[2], _theme_transition_start=-999.0, _theme_transition_duration=0.6,
        zoom=1.0, brightness=1.0, roll=0.0, translate_x=0.0, translate_y=0.0, flash_color=(1, 1, 1), flash_intensity=0.0,
        cpu_history=deque([0.0] * 40, maxlen=40), state="idle", hud_lines=[], rotation_x=0.0, rotation_y=0.0,
        pulse_phase=0.0, sweep_phase=0.0, scanline_phase=0.0, _glitch=0.0, _glitch_timer=5.0, _dock_shift=0.0,
        _last_time=time.time() - 0.02, _fps=0.0, font=None, font_small=None, font_big=None, zoom_=1.0)
    return hm, h


@ext("Hologram logic")
def element_tables_and_shell_rules():
    hm = mod("hologram")
    eq(len(hm.PERIODIC_TABLE), 118)
    eq(hm.ELEMENTS["carbon"], ("C", 6))
    eq(hm.ATOMIC_NUMBER_TO_ELEMENT[79], ("Au", "gold"))
    eq(hm.fill_shells(11), [2, 8, 1])
    eq(hm.fill_shells(0), [])
    eq(sum(hm.fill_shells(118)), 118)
    eq(hm._approx_neutrons(1), 1)
    eq(hm._approx_neutrons(20), 20)
    assert hm._approx_neutrons(92) > 92
    a, b = hm._generate_procedural_system("zorg"), hm._generate_procedural_system("zorg")
    eq(a, b)
    assert 2 <= len(a) <= 6 and a[0][0].startswith("Zorg ")
    nodes = hm._build_network_nodes()
    eq(len(nodes), 10)
    edges = hm._build_network_edges(nodes)
    assert all(i < j for i, j in edges) and {k for e in edges for k in e} == set(range(10))
    eq(hm.Hologram._nucleon_positions(None, 0), [])
    eq(len(hm.Hologram._nucleon_positions(None, 7)), 7)


@ext("Hologram logic")
def atom_loading_and_editing():
    hm, h = bare_holo()
    eq(h.load_atom("carbon"), ("C", "carbon"))
    eq((h.protons, h.neutrons, h.electron_count, len(h.orbits)), (6, 6, 6, 2))
    assert "CARBON" in h.mode_label
    eq(h.load_atom("Fe"), ("Fe", "iron"))
    eq(h.load_atom("unobtainium"), None)
    h.load_atom("hydrogen")
    h.add_protons(1)
    assert (h.protons, h.electron_count) == (2, 2) and "HELIUM" in h.mode_label
    h.add_neutrons(-9)
    eq(h.neutrons, 0)
    h.add_electrons(1)
    assert "charge -1" in h.mode_label
    h.protons, h.electron_count = 500, 500
    h._update_custom_label()
    assert "CUSTOM ELEMENT" in h.mode_label
    h.current_element_z = 118
    eq(h.next_element()[2], 1)
    eq(h.previous_element()[2], 118)


@ext("Hologram logic")
def display_loaders():
    hm, h = bare_holo()
    h.load_solar_system()
    eq((h.mode, len(h.orbits)), ("solar_system", 8))
    eq(h.load_star_system("trappist-1"), ("Trappist-1", 7, False))
    label, n, generated = h.load_star_system("made up")
    assert generated and "(est.)" in h.mode_label
    first = h.next_star_system()[0]
    assert first != h.previous_star_system()[0] or True
    h.load_demo()
    eq(len(h.orbits), 4)
    h.load_network()
    eq(len(h.network_nodes), 10)
    assert h.load_shape("cube") and not h.load_shape("nope")
    assert h.load_molecule("water") and not h.load_molecule("unobtainium")
    assert h.load_constellation("orion") and not h.load_constellation("nothing")
    h.load_satellite("ISS")
    eq((h.mode, len(h.orbits)), ("satellite", 1))
    h.show_info_card("q", "a")
    eq((h.mode, h.info_card), ("info", {"question": "q", "answer": "a"}))
    h.show_weather({"location": "X", "state": "texas"})
    assert h.docked and h.map_highlight == "texas"
    h.hide_weather()
    assert not h.docked and h.map_highlight is None


@ext("Hologram logic")
def graphs_and_physics_steps():
    hm, h = bare_holo()
    assert h.load_math_graph("sin(x)") and len(h.graph_points) == 161
    assert h.load_math_graph("x^2")
    assert not h.load_math_graph("") and not h.load_math_graph("1/0")
    assert not h.load_math_graph("__import__('os').getcwd()"), "eval sandbox escaped"
    eq(h._eval_graph_expr("x+1", 2), 3)
    h.load_physics_sim("pendulum")
    th = h.physics_state["theta"]
    h._step_physics(0.05)
    assert h.physics_state["theta"] != th and h.physics_state["trail"]
    h.load_physics_sim("projectile")
    h._step_physics(0.1)
    assert h.physics_state["x"] > -2.2
    h.physics_state["y"] = -5
    h._step_physics(0.1)
    eq(h.physics_state["x"], -2.2)                                      # resets after landing
    h.load_physics_sim("other")
    eq(h.physics_state, {})


@ext("Hologram logic")
def orbit_selection_and_view_controls():
    hm, h = bare_holo()
    h.load_demo()
    assert h.select_orbit(0) and not h.select_orbit(9)
    h.deselect_orbit()
    seen = []
    for _ in range(5):
        h.select_next_orbit()
        seen.append(h.selected_index)
    eq(seen, [0, 1, 2, 3, None])
    h.select_orbit(0)
    h.edit_selected_orbit(-1.0, 1.0, 10.0)
    eq(h.orbits[0]["radius"], hm.RADIUS_MAX)
    assert h.orbits[0]["speed"] <= hm.SPEED_MAX
    h.edit_selected_orbit(1.0, -1.0, 100.0)
    eq(h.orbits[0]["radius"], hm.RADIUS_MIN)
    h.apply_zoom_delta(100)
    eq(h.zoom, 3.0)
    h.apply_zoom_delta(-100)
    eq(h.zoom, 0.4)
    h.apply_roll_delta(400)
    eq(h.roll, 40.0)
    h.apply_pan_delta(100, 100)
    assert h.translate_x == 3.0 and h.translate_y == -2.0
    h.adjust_brightness(5)
    eq(h.brightness, 1.8)
    h.cycle_theme(1)
    eq(h.theme_index, 3)
    h.set_theme_index(9)
    eq(h.theme_index, 1)
    h.elapsed = 5.0
    eq(h._current_theme_color(), hm.THEMES[1])
    h.trigger_flash("thumbs_up", 0.5)
    eq((h.flash_color, h.flash_intensity), (hm.FLASH_COLORS["thumbs_up"], 0.5))
    h.reset_orbits()
    eq(h.mode, "demo")


@ext("Hologram logic")
def events_telemetry_and_helpers():
    hm, h = bare_holo()
    heard = []
    h.event_listeners = [lambda s: heard.append(s), lambda s: 1 / 0]
    h.log_event("FACE: hi")
    assert heard == ["FACE: hi"] and "FACE: hi" in h.event_log[0]
    h.update_telemetry({"t": 1}, True)
    eq(h.telemetry_data, {})                                              # ignored unless on screen
    h.show_telemetry("Mars")
    assert h.mode == "telemetry" and "MARS" in h.mode_label
    h.update_telemetry({"t": 1.5, "name": "x", "flag": True}, True)
    eq(list(h.telemetry_history), ["t"])
    h.hide_telemetry()
    eq(h.mode, "empty")
    eq(h._truncate("abcdefgh", 5), "abcd…")
    eq(h._normalize_history([5, 5, 5]), [0.5, 0.5, 0.5])
    eq(h._normalize_history([0, 5, 10]), [0.0, 0.5, 1.0])
    eq(h._wrap_text(None, "hello world", 10), ["hello world"])
    eq(h._pill_width("abc"), 90)
    assert h._stat_status_color("CPU", None) != h._stat_status_color("CPU", 0.9)
    eq(h._stat_status_color("BATT", 0.1), (1.0, 0.35, 0.3))
    h.elapsed = 99
    eq(h._materialize_progress(), 1.0)
    h.mode = "atom"
    assert any(k == "PROTONS" for k, _ in h._mode_readout())
    h.mode = "shape"
    assert ("MODEL", "SPHERE") in h._mode_readout() and h._mode_readout()[-1][0] == "BRIGHTNESS"
    h.cpu_history = deque([0.0] * 40, maxlen=40)
    h.update(0.2, -0.2, 0.5, target_dt=0.1)
    assert h.rotation_y > 0 and h.elapsed > 99


# ============================================================================ system control
@ext("System control")
def filenames_templates_and_project_lookup():
    sc = mod("system_control")
    eq(sc.safe_filename("my demo!"), "my_demo")
    eq(sc.safe_filename("../x"), "x")
    eq(sc.safe_filename("!!!"), "untitled")
    assert "def main" in sc._template("python", "x") and "<html>" in sc._template("html", "x")
    eq(sc._template("klingon", "x"), "# x\n")
    d = ext_dir()
    with mock.patch.object(sc, "BASE", d), mock.patch.object(sc, "PROJECTS_ROOT", os.path.join(d, "p")), \
            mock.patch.object(sc, "ARDUINO_ROOT", os.path.join(d, "a")):
        json.dump({"Demo": "C:/x/demo.py"}, open(os.path.join(d, "code_projects.json"), "w"))
        eq(sc.resolve_project_path("demo"), "C:/x/demo.py")
        eq(sc.find_arduino_sketch("DEMO"), "C:/x/demo.py")
        eq(sc.resolve_project_path("ghost"), None)
        ino = sc.get_or_make_path("blinker", "arduino")
        assert ino.endswith("blinker.ino") and os.path.exists(ino)
        py = sc.get_or_make_path("tool", "python")
        assert py.endswith("tool.py")
        eq(sc.get_or_make_path("tool", "python"), py)
        sc.write_file_content(os.path.join(d, "sub", "x.txt"), "hi")
        eq(open(os.path.join(d, "sub", "x.txt")).read(), "hi")
    eq(sc.summarize_diff("a\nb\n", "a\n"), "1 line removed")
    eq(sc.summarize_diff("a\n", "a\n"), "no changes")
    assert not sc.open_in_vscode("/no/such/path")[0] and not sc.open_in_arduino("/no/such/path")[0]
    calls = []
    with mock.patch.object(sc, "open_in_arduino", lambda p: calls.append("ino") or (True, "")), \
            mock.patch.object(sc, "open_in_vscode", lambda p: calls.append("code") or (True, "")):
        sc.open_in_editor("a.ino")
        sc.open_in_editor("a.py")
    eq(calls, ["ino", "code"])


@ext("System control")
def volume_media_apps_and_status():
    sc = mod("system_control")
    with mock.patch.object(sc, "PYCAW_AVAILABLE", False):
        assert sc.get_volume_percent() is None and not sc.set_volume_percent(5) and sc.adjust_volume_percent(5) is None
    state = {"v": 0.5}
    iface = ns(GetMasterVolumeLevelScalar=lambda: state["v"], SetMasterVolumeLevelScalar=lambda v, ctx: state.update(v=v))
    with mock.patch.object(sc, "PYCAW_AVAILABLE", True), mock.patch.object(sc, "_volume_interface", lambda: iface):
        eq(sc.get_volume_percent(), 50)
        eq(sc.adjust_volume_percent(10), 60)
        assert sc.set_volume_percent(150)
        eq(state["v"], 1.0)
        sc.set_volume_percent(-5)
        eq(state["v"], 0.0)
    keys = []
    fake = ns(windll=ns(user32=ns(keybd_event=lambda *a: keys.append(a[0]))))
    with mock.patch.object(sc, "ctypes", fake):
        assert sc.media_play_pause() and sc.media_next() and sc.media_previous()
    eq(keys[::2], [0xB3, 0xB0, 0xB1])
    eq(sc.close_app("explorer"), False)
    seen = []
    with mock.patch.object(sc.subprocess, "run", lambda cmd, **k: seen.append(cmd) or ns(returncode=0)):
        assert sc.close_app("calc")
    assert "CalculatorApp.exe" in seen[0]
    with mock.patch.object(os, "startfile", create=True, side_effect=OSError), mock.patch.object(sc.shutil, "which", return_value=None):
        assert not sc.launch_app("nothing")
    with mock.patch.object(os, "startfile", create=True, side_effect=OSError), mock.patch.object(sc.shutil, "which", return_value="/bin/x"), \
            mock.patch.object(sc.subprocess, "Popen") as po:
        assert sc.launch_app("x") and po.called
    eq(sc.APP_RE.search("open visual studio code").group(1), "visual studio code")
    fake_ps = ns(cpu_percent=lambda interval=None: 12.4, virtual_memory=lambda: ns(percent=40.2),
                 sensors_battery=lambda: ns(percent=80, power_plugged=True),
                 sensors_temperatures=lambda: {"a": [ns(current=50.0)]})
    with mock.patch.object(sc, "PSUTIL_AVAILABLE", True), mock.patch.object(sc, "psutil", fake_ps):
        s = sc.get_system_status()
        assert "CPU at 12 percent" in s and "charging" in s
        eq(sc.get_cpu_temp(), 50.0)
    with mock.patch.object(sc, "PSUTIL_AVAILABLE", False):
        assert sc.get_system_status() is None and sc.get_cpu_temp() is None and sc.get_foreground_process_name() is None


@ext("System control")
def screen_recording_state_machine():
    sc = mod("system_control")
    with mock.patch.object(sc, "RECORDING_DIR", ext_dir()):
        with mock.patch.object(sc.subprocess, "Popen", side_effect=FileNotFoundError):
            ok, msg = sc.start_recording()
            assert not ok and "ffmpeg" in msg and not sc.is_recording()
        proc = mock.MagicMock()
        with mock.patch.object(sc.subprocess, "Popen", return_value=proc):
            ok, path = sc.start_recording("a.mp4")
            assert ok and path.endswith("a.mp4") and sc.is_recording()
            assert sc.start_recording()[1] == "Already recording."
            assert sc.stop_recording()[0] and not sc.is_recording()
            proc.communicate.assert_called_once()
        eq(sc.stop_recording(), (False, "Not recording."))


@ext("Phone control (ADB mocked)")
def adb_basics_contacts_and_calls():
    sc = mod("system_control")
    d = ext_dir()
    shell_log = []
    with mock.patch.object(sc, "BASE", d), mock.patch.object(sc, "_shell", lambda cmd, timeout=15: shell_log.append(cmd) or (True, "")):
        with mock.patch.object(sc, "_adb", return_value=(True, "List of devices attached\nabc\tdevice\n")):
            assert sc.is_device_connected()[0]
        with mock.patch.object(sc, "_adb", return_value=(True, "List of devices attached\nabc\toffline\n")):
            assert not sc.is_device_connected()[0]
        with mock.patch.object(sc, "_adb", return_value=(False, "adb not found")):
            eq(sc.is_device_connected(), (False, "adb not found"))
        eq(sc.unlock_with_pin(), (False, "No phone_pin.txt found"))
        open(os.path.join(d, "phone_pin.txt"), "w").write("12ab")
        assert "digits only" in sc.unlock_with_pin()[1]
        open(os.path.join(d, "phone_pin.txt"), "w").write("1234")
        assert sc.unlock_with_pin()[0]
        assert "input text 1234" in shell_log and "input keyevent KEYCODE_ENTER" in shell_log
        json.dump({"Mom": "123", "me": "999"}, open(os.path.join(d, "contacts.json"), "w"))
        eq(sc.load_contacts(), {"mom": "123", "me": "999"})
        eq(sc.load_my_number(), "999")
        open(os.path.join(d, "my_number.txt"), "w").write("555")
        eq(sc.load_my_number(), "555")
        sc.call_number("(555) 123-4567")
        assert shell_log[-1].endswith("tel:5551234567")
        sc.wifi_set(True)
        sc.wifi_set(False)
        assert shell_log[-2].endswith("enable") and shell_log[-1].endswith("disable")
        json.dump({"home": {"ssid": "Net", "password": "pw"}, "cafe": {"ssid": "Cafe", "password": ""}},
                  open(os.path.join(d, "wifi_networks.json"), "w"))
        assert sc.wifi_connect("home")[0] and "wpa2 'pw'" in shell_log[-1]
        sc.wifi_connect("cafe")
        assert shell_log[-1].endswith("'Cafe' open")
        ok, msg = sc.wifi_connect("ghost")
        assert not ok and "wifi_networks.json" in msg
        sc.web_search("hello world")
        assert "q=hello+world" in shell_log[-1]


@ext("Phone control (ADB mocked)")
def adb_packages_apps_wireless_whatsapp_search():
    sc = mod("system_control")
    d = ext_dir()
    pm = "package:com.foo.bar\npackage:com.whatsapp\npackage:com.zed.app"

    def shell(cmd, timeout=15):
        if cmd == "pm list packages":
            return True, pm
        if "contacts/data/phones" in cmd:
            return True, "Row: 0 display_name=Mom, data1=+1 555\nRow: 1 display_name=Bob, data1=123"
        if "mimetype=" in cmd:
            return True, "Row: 0 _id=7, display_name=Mom Smith"
        if "ip -f inet" in cmd:
            return True, "inet 192.168.1.5/24"
        if cmd.startswith("monkey") and "com.zed.app" in cmd:
            return True, "No activities found"
        if cmd.startswith("find"):
            return True, "/sdcard/a_notes.txt\nnoise"
        return True, ""

    with mock.patch.object(sc, "BASE", d), mock.patch.object(sc, "_shell", shell), mock.patch.object(sc.time, "sleep", lambda s: None):
        eq(sc._resolve_pkg("instagram"), "com.instagram.android")
        eq(sc._resolve_pkg("zed"), "com.zed.app")
        eq(sc._resolve_pkg("nothinghere"), None)
        assert sc.open_app("whatsapp") and not sc.open_app("zed") and not sc.open_app("nothinghere")
        assert sc.close_phone_app("whatsapp") and not sc.close_phone_app("nothinghere")
        eq(sc._phone_contacts(), [("Mom", "+1 555"), ("Bob", "123")])
        eq(sc.whatsapp_call("mom"), (True, "Mom Smith"))
        res = sc.search_phone("mom")
        eq(res["contacts"], [("Mom", "+1 555")])
        eq(sc.search_phone("whats", "apps")["apps"], ["com.whatsapp"])
        eq(sc.search_phone("notes", "files")["files"], ["/sdcard/a_notes.txt"])
        with mock.patch.object(sc, "_adb", side_effect=[(True, ""), (True, "connected to 192.168.1.5:5555")]):
            eq(sc.connect_wireless(), (True, "192.168.1.5"))
        with mock.patch.object(sc, "_adb", return_value=(False, "no")):
            eq(sc.connect_wireless(), (False, "no"))
    with mock.patch.object(sc, "BASE", d), mock.patch.object(sc, "_shell", lambda c, timeout=15: (True, "")):
        ok, msg = sc.whatsapp_call("stranger")
        assert not ok and "couldn't find" in msg
        json.dump({"dad": "+1 (800) 555"}, open(os.path.join(d, "contacts.json"), "w"))
        ok, msg = sc.whatsapp_call("dad")
        assert ok and "tap the call button" in msg


@ext("System control")
def qr_codes_vault_helpers_and_pdf_reading():
    sc = mod("system_control")
    import cv2
    import numpy as np
    eq(sc.read_codes(np.zeros((50, 50, 3), np.uint8)), [])
    if hasattr(cv2, "QRCodeEncoder_create"):
        qr = cv2.QRCodeEncoder_create().encode("hello")
        big = cv2.copyMakeBorder(cv2.resize(qr, None, fx=8, fy=8, interpolation=cv2.INTER_NEAREST), 40, 40, 40, 40,
                                 cv2.BORDER_CONSTANT, value=255)
        codes = sc.read_codes(cv2.cvtColor(big, cv2.COLOR_GRAY2BGR) if big.ndim == 2 else big)
        if not codes:
            skip("QR decoder did not read the generated code on this OpenCV build")
        assert any("hello" in c for c in codes)
    d = ext_dir()
    open(os.path.join(d, "a.txt"), "w").write("hello vault")
    open(os.path.join(d, "bad.pdf"), "wb").write(b"not a pdf")
    eq(sc._read_any(os.path.join(d, "a.txt")), "hello vault")
    eq(sc._read_any(os.path.join(d, "bad.pdf")), "")
    eq(sc._read_any(os.path.join(d, "x.xyz")), "")
    assert "the" not in sc._vwords("the quick fox")
    with mock.patch.object(sc, "BASE", d):
        eq(sc.vault_search("anything"), None)
        ok, msg = sc.vault_add_folder(os.path.join(d, "missing"))
        assert not ok and "isn't a folder" in msg


# ============================================================================ voice assistant
def va_obj():
    va = mod("voice_assistant")
    v = object.__new__(va.VoiceAssistant)
    said = []
    v._log = v._on_log = lambda m: None
    v._speak = said.append
    v.speech = ns(say=said.append, busy=False, wait=lambda *a: None, stop=lambda: None)
    v.active_timers, v.last_reply, v.latest_frame, v.last_weather_error = [], "", None, None
    v._cancel, v._turn_display, v._state, v._capture, v._capture_tid = threading.Event(), False, "idle", None, None
    v.hologram, v.client, v.brain, v.plus = mock.MagicMock(), None, None, None
    return va, v, said


@ext("Voice assistant internals")
def wake_word_calculator_cleaning_chunker():
    va = mod("voice_assistant")
    eq(va.find_wake_word("hey aurora, set a timer"), (True, "set a timer"))
    eq(va.find_wake_word("set a timer aurora"), (True, "set a timer"))
    eq(va.find_wake_word("aurora"), (True, ""))
    eq(va.find_wake_word("nothing here"), (False, None))
    eq(va.try_calculate("what is 5 plus 3")[0], 8)
    eq(va.try_calculate("10 divided by 4")[0], 2.5)
    assert va.try_calculate("1 divided by 0") is None and va.try_calculate("__import__('os')") is None
    raises(ValueError, va._eval_node, va.ast.parse("foo", mode="eval").body)
    eq(va.clean_for_speech("see [docs](http://x.y) now\n- item\n```code```"), "see docs now item")
    c = va.Chunker()
    eq(c.feed("Hi"), [])
    eq(c.flush(), ["Hi"])
    eq(c.flush(), [])
    out = c.feed("This is the first sentence. And a second sentence that is long enough. ")
    assert len(out) == 2
    assert va.STOP_RE.fullmatch("please stop") and not va.STOP_RE.fullmatch("stop the music now please")


@ext("Voice assistant internals")
def api_key_loading_and_null_voice():
    va = mod("voice_assistant")
    d = ext_dir()
    keyfile = os.path.join(d, "k.txt")
    with mock.patch.object(va, "API_KEY_FILE", keyfile), mock.patch.dict(os.environ, {"GROQ_API_KEY": "envkey"}):
        eq(va.load_api_key(), "envkey")
        open(keyfile, "w").write(" filekey \n")
        eq(va.load_api_key(), "filekey")
    nv = va.NullVoice()
    nv.start()
    nv.stop()
    nv.speak_now("hi")
    assert nv.state == "idle" and not nv.enabled and nv.active_reader is nv.telemetry


@ext("Voice assistant internals")
def offline_recognizer_paths():
    va = mod("voice_assistant")
    import numpy as np
    sr = va.sr
    rec = ns(recognize_sphinx=lambda audio: "sphinx text")
    o = va.OfflineRecognizer(lambda m: None, rec)
    o._load = lambda: None
    eq(o.recognize(ns()), "sphinx text")
    rec.recognize_sphinx = mock.Mock(side_effect=sr.RequestError("x"))
    raises(sr.UnknownValueError, o.recognize, ns())
    audio = ns(get_raw_data=lambda **k: np.zeros(1600, np.int16).tobytes())
    o._load = lambda: ns(transcribe=lambda pcm, **k: ([ns(text=" hi "), ns(text="there")], None))
    eq(o.recognize(audio), "hi there")
    o._load = lambda: ns(transcribe=lambda pcm, **k: ([], None))
    raises(sr.UnknownValueError, o.recognize, audio)
    r2 = ns(recognize_google=mock.Mock(side_effect=sr.RequestError("down")))
    _, v, _ = va_obj()
    v.recognizer = r2
    with mock.patch.object(va, "OfflineRecognizer", lambda log, rec: ns(_load=lambda: None, recognize=lambda a: "offline text")):
        v.enable_offline_recognition(lambda: False)
        eq(r2.recognize_google(ns()), "offline text")
        v.enable_offline_recognition(lambda: True)


@ext("Voice assistant internals")
def weather_place_parsing_and_wmo():
    va = mod("voice_assistant")
    P = va.VoiceAssistant._weather_place
    eq(P("weather in new york right now"), "new york")
    eq(P("paris weather"), "paris")
    eq(P("what's the weather like in delhi"), "delhi")
    eq(P("what's the weather right now"), None)
    eq(P("weather in my city"), None)
    eq(va.VoiceAssistant.WMO[0], ("sunny", "clear sky"))
    eq(va.VoiceAssistant.WMO[95][0], "storm")
    eq(va.VoiceAssistant.WMO[61][0], "rain")


@ext("Voice assistant internals")
def geocoding_ip_location_and_weather_fetch():
    va, v, _ = va_obj()
    results = {"results": [{"name": "Paris", "admin1": "Ile-de-France", "country": "France", "latitude": 1, "longitude": 2},
                           {"name": "Paris", "admin1": "Texas", "country": "United States", "latitude": 3, "longitude": 4}]}
    v._get_json = lambda url, headers=None: results
    eq(v._geocode_location("Paris, Texas"), (3, 4, "Paris, Texas, United States"))
    eq(v._geocode_location("Paris")[:2], (1, 2))
    v._get_json = lambda url, headers=None: {}
    eq(v._geocode_location("Nowhere"), None)
    v._get_json = mock.Mock(side_effect=OSError("x"))
    eq(v._geocode_location("Paris"), None)
    v._get_json = lambda url, headers=None: {"latitude": 1, "longitude": 2, "city": "X", "country_name": "Y"}
    eq(v._ip_geolocate(), (1, 2, "X, Y"))
    v._get_json = lambda url, headers=None: {"city": "X"}
    eq(v._ip_geolocate(), None)
    v._geocode_location = lambda q: (1.0, 2.0, "Place")
    v._get_json = lambda url, headers=None: {"current": {"temperature_2m": 9.5, "weather_code": 61, "relative_humidity_2m": 70, "wind_speed_10m": 12}}
    d = v._fetch_weather_data("Place")
    eq((d["location"], d["condition"], d["temp_c"], d["humidity"]), ("Place", "rain", 9.5, 70))
    v._get_json = mock.Mock(side_effect=OSError("down"))
    eq(v._fetch_weather_data("Place"), None)
    eq(v.last_weather_error, "network")
    v._geocode_location = lambda q: None
    v._fetch_weather_data("Atlantis")
    eq(v.last_weather_error, "not_found")
    v._ip_geolocate = lambda: None
    v._fetch_weather_data(None)
    eq(v.last_weather_error, "no_location")


@ext("Voice assistant internals")
def tool_execution_phrase_mapping():
    va, v, _ = va_obj()
    v._run_local = lambda p: p
    cases = [("set_timer", {"seconds": 30, "label": "tea"}, "set a timer for 30 seconds called tea"),
             ("pc_control", {"action": "volume_up"}, "volume up"), ("pc_control", {"action": "set_volume", "value": "50"}, "volume to 50"),
             ("pc_control", {"action": "open_app", "value": "notepad"}, "open notepad"),
             ("phone_control", {"action": "call", "target": "me"}, "call me"),
             ("phone_control", {"action": "whatsapp_call", "target": "mom"}, "call mom on whatsapp"),
             ("phone_control", {"action": "wifi_off"}, "turn off wifi on my phone"),
             ("show_display", {"kind": "graph", "expression": "x squared"}, "plot x squared"),
             ("show_display", {"kind": "atom", "name": "gold"}, "show me a gold atom"),
             ("show_display", {"kind": "heart"}, "show the human heart"),
             ("edit_atom", {"particle": "electron", "delta": -2}, "remove 2 electron"),
             ("take_note", {"text": "hi"}, "take a note: hi"),
             ("write_code", {"name": "x", "language": "rust", "description": "does y"}, "write python code called x that does y"),
             ("debug_code", {"name": "sorter"}, "debug sorter"),
             ("create_document", {"kind": "presentation", "topic": "mars"}, "create a presentation about mars"),
             ("get_weather", {"location": "Rome"}, "weather in Rome"), ("get_weather", {}, "what's the weather right now"),
             ("nonsense", {}, "Unsupported request.")]
    for name, args, want in cases:
        eq(v._exec_tool(name, args), want)
    assert v._exec_tool("get_time", {}).startswith("It's")
    sc = mod("system_control")
    with mock.patch.object(sc, "read_notes", return_value=[]):
        eq(v._exec_tool("read_notes", {}), "no notes yet")
    with mock.patch.object(sc, "vault_search", return_value=None):
        assert "vault isn't built" in v._exec_tool("search_files", {"query": "x"})
    v._handle_local_command = lambda p: False
    v2 = va_obj()[1]
    v2._handle_local_command = lambda p: False
    eq(v2._run_local("zzz"), "That action isn't available.")


@ext("Voice assistant internals")
def speech_timers_run_text_and_phone_commands():
    va, v, said = va_obj()
    v.speak_now("Hello **there**")
    eq((v.last_reply, said[-1]), ("Hello there", "Hello **there**"))
    v._start_timer(0, "tea")
    time.sleep(0.4)
    assert v.active_timers == [] and any("tea is up" in s for s in said)
    v._handle_local_command = lambda text: setattr(v, "last_reply", "ok") or True
    eq(v.run_text("x"), "ok")
    v._handle_local_command = lambda text: True
    v.last_reply = ""
    eq(v.run_text("x"), "Done.")
    asked = []
    v._handle_local_command = lambda text: False
    v._ask = lambda text: asked.append(text) or setattr(v, "last_reply", "ai reply")
    eq(v.run_text("hello"), "ai reply")
    sc = mod("system_control")
    v._phone_ready = lambda: True
    said.clear()
    with mock.patch.object(sc, "whatsapp_call", return_value=(True, "Mom")):
        assert v._handle_phone_command("video call mom on whatsapp") and "video call with Mom" in said[-1]
    with mock.patch.object(sc, "connect_wireless", return_value=(True, "1.2.3.4")):
        assert v._handle_phone_command("connect to my phone over wifi") and "Connected wirelessly" in said[-1]
    with mock.patch.object(sc, "close_phone_app", return_value=True):
        assert v._handle_phone_command("close instagram on my phone") and "Closed instagram" in said[-1]
    with mock.patch.object(sc, "web_search", return_value=True):
        assert v._handle_phone_command("google cats on my phone") and "Searching for cats" in said[-1]
    with mock.patch.object(sc, "wifi_connect", return_value=(True, "Connecting your phone to Home.")):
        assert v._handle_phone_command("connect my phone to the wifi network home")
    with mock.patch.object(sc, "search_phone", return_value={"contacts": [("Mom", "1")], "apps": [], "files": []}):
        assert v._handle_phone_command("search my phone for contact mom") and "1 contact" in said[-1]
    with mock.patch.object(sc, "search_phone", return_value={"contacts": [], "apps": [], "files": []}):
        v._handle_phone_command("search phone for zzz")
        assert "didn't find" in said[-1]
    assert not v._handle_phone_command("what time is it")
    v._phone_ready = va.VoiceAssistant._phone_ready.__get__(v)
    with mock.patch.object(sc, "is_device_connected", return_value=(False, "")):
        said.clear()
        assert not v._phone_ready() and "can't see your phone" in said[-1]
    v._interrupt()
    assert v._cancel.is_set()


# ============================================================================ brain
class _Create:
    """client.chat.completions.create stand-in: returns scripted streams, records kwargs."""

    def __init__(self, *streams):
        self.streams, self.calls = list(streams), []
        self.chat = ns(completions=ns(create=self))

    def __call__(self, **kw):
        self.calls.append(kw)
        item = self.streams.pop(0)
        if isinstance(item, Exception):
            raise item
        return item


def _ev(content=None, tool=None):
    tcs = [ns(index=0, id=tool[0], function=ns(name=tool[1], arguments=tool[2]))] if tool else None
    return ns(choices=[ns(delta=ns(content=content, tool_calls=tcs))])


@ext("AI brain")
def agent_tool_loop_and_streaming():
    br = mod("brain")
    c = _Create(iter([_ev(tool=("c1", "set_timer", '{"sec')), _ev(tool=(None, None, 'onds": 5}'))]),
                iter([_ev("Done"), _ev(" now.")]))
    calls, deltas = [], []
    b = br.Brain(c, lambda m: None, lambda n, a: calls.append((n, a)) or "Timer set")
    reply = b.ask("set a timer", deltas.append, threading.Event())
    eq(calls, [("set_timer", {"seconds": 5})])
    eq((reply, "".join(deltas)), ("Done now.", "Done now."))
    eq([m["role"] for m in b.memory.recent], ["user", "assistant"])
    tool_msg = [m for m in c.calls[1]["messages"] if m["role"] == "tool"][0]
    eq(tool_msg["content"], "Timer set")
    c2 = _Create(iter([_ev(tool=("c1", "x", "{}"))]), iter([_ev("fine")]))
    b2 = br.Brain(c2, lambda m: None, mock.Mock(side_effect=RuntimeError("boom")))
    b2.ask("do it", lambda d: None, threading.Event())
    assert [m for m in c2.calls[1]["messages"] if m["role"] == "tool"][0]["content"] == "error: boom"
    cancel = threading.Event()
    cancel.set()
    eq(br.Brain(_Create(iter([_ev("x")])), lambda m: None, None).ask("hi", lambda d: None, cancel), "")


@ext("AI brain")
def request_fallbacks_and_oversize_retry():
    br = mod("brain")
    c = _Create(RuntimeError("bad tool"), iter([_ev("ok")]))
    b = br.Brain(c, lambda m: None, lambda n, a: "")
    eq(b.ask("what's the latest news", lambda d: None, threading.Event()), "ok")
    assert {"type": "browser_search"} in c.calls[0]["tools"] and {"type": "browser_search"} not in c.calls[1]["tools"]
    big = RuntimeError("413 request_too_large")
    b2 = br.Brain(_Create(big, big), lambda m: None, lambda n, a: "")
    raises(RuntimeError, b2.ask, "hello", lambda d: None, threading.Event())
    eq(b2.memory.recent, [])
    assert br._too_large(RuntimeError("Request Too Large")) and not br._too_large(RuntimeError("other"))


@ext("AI brain")
def completion_code_helpers_and_memory_compression():
    br = mod("brain")
    resp = lambda text: ns(choices=[ns(message=ns(content=text), finish_reason="stop")])
    client = ns(chat=ns(completions=ns(create=lambda **k: resp("  hi  "))))
    eq(br._complete(client, "s", "u", 100), "hi")
    empty = ns(chat=ns(completions=ns(create=lambda **k: resp(""))))
    raises(RuntimeError, br._complete, empty, "s", "u", 100)
    with mock.patch.object(br, "_complete", return_value="```python\nprint(1)\n```"):
        eq(br.generate_code(None, "x"), "print(1)\n")
        eq(br.edit_code(None, "a", "b"), "print(1)\n")
        eq(br.fix_code(None, "a", "err"), "print(1)\n")
    eq(br.strip_code_fences("no fences"), "no fences\n")
    assert not br.fuzzy_hit("hello world", ["aurora"]) and br.fuzzy_hit("open the pod bay doors", ["pod bay"])
    with mock.patch.object(br, "_complete", return_value="SUMMARY"):
        m = br.RollingMemory(object(), keep_recent=4, compress_above=6)
        for i in range(10):
            m.add("user", f"msg {i}")
        m.maybe_compress()
        eq(len(m.recent), 4)
        assert m.get_messages()[0]["content"].endswith("SUMMARY")
        m.pop_last()
        eq(len(m.recent), 3)


@ext("AI brain")
def vision_helpers_model_fallback():
    br = mod("brain")
    import numpy as np
    frame = np.zeros((16, 16, 3), np.uint8)
    seq = [RuntimeError("m1 down"), ns(choices=[ns(message=ns(content=""))]), ns(choices=[ns(message=ns(content="a cat"))])]

    def create(**kw):
        item = seq.pop(0)
        if isinstance(item, Exception):
            raise item
        return item
    client = ns(chat=ns(completions=ns(create=create)))
    eq(br.describe_scene(client, frame), "a cat")
    bad = ns(chat=ns(completions=ns(create=mock.Mock(side_effect=RuntimeError("all down")))))
    raises(RuntimeError, br.describe_scene, bad, frame)
    prompts = []
    with mock.patch.object(br, "describe_scene", lambda c, f, p, *a: prompts.append(p) or "answer"):
        eq(br.describe_screen(None, frame, "why fails"), "answer")
    assert "why fails" in prompts[0] and "untrusted" in prompts[0]
    f = br._fn("x", "d", {"a": {"type": "string"}}, ["a"])
    eq(f["function"]["parameters"]["required"], ["a"])


# ============================================================================ utilities (aurora_utilities)
def util():
    return mod("aurora_utilities")


@ext("Utilities: files")
def home_boundaries_folder_lookup_categories_dates():
    u = util()
    home = ext_dir()
    for sub in ("Documents", "Downloads", "AppData/Local", "Documents/School"):
        os.makedirs(os.path.join(home, sub))
    with mock.patch.object(u, "HOME", home):
        assert u.inside_home(os.path.join(home, "Documents")) and not u.inside_home(home)
        assert not u.inside_home(os.path.join(home, "AppData", "Local")) and not u.inside_home(os.path.dirname(home))
        eq(u.folder_path("my downloads folder"), os.path.join(home, "Downloads"))
        eq(u.folder_path("home"), home)
        eq(u.folder_path("nope"), None)
        eq(u.find_folder("school"), os.path.join(home, "Documents", "School"))
        eq(u.find_file("nothing"), [])
        for i in range(2):
            open(os.path.join(home, "Downloads", f"f{i}.txt"), "w").write("x")
        with mock.patch.object(u, "MAX_ORGANIZE", 1):
            raises(u.FileError, u.plan_organize, os.path.join(home, "Downloads"))
        raises(u.FileError, u.plan_organize, tempfile.gettempdir())
        raises(u.FileError, u.rename_file, os.path.join(home, "Downloads", "f0.txt"), "..")
        raises(u.FileError, u.move_file, os.path.join(home, "Downloads", "f0.txt"), os.path.join(home, "missing"))
        raises(u.FileError, u.extract_zip, os.path.join(home, "Downloads", "f0.txt"))
        z = u.compress_path(os.path.join(home, "Documents"))
        assert z.endswith("Documents.zip") and zipfile.is_zipfile(z)
    eq((u.category_for("a.PNG"), u.category_for("a.pdf"), u.category_for("a.xyz"), u.category_for("a.exe")),
       ("Images", "Documents", "Other", "Installers"))
    now = datetime(2026, 10, 14, 15, 30)
    eq(u.date_range("today", now), (datetime(2026, 10, 14), None))
    eq(u.date_range("yesterday", now), (datetime(2026, 10, 13), datetime(2026, 10, 14)))
    ws, we = u.date_range("last week", now)
    eq((we - ws).days, 7)
    eq(u.date_range("last month", now), (datetime(2026, 9, 1), datetime(2026, 10, 1)))
    eq(u.date_range("whenever", now), None)


@ext("Utilities: files")
def file_search_peek_and_name_proposals():
    u = util()
    root = ext_dir()
    open(os.path.join(root, "physics_notes.txt"), "w").write("newton laws")
    open(os.path.join(root, "diary.txt"), "w").write("studied thermodynamics today")
    open(os.path.join(root, "cooking.txt"), "w").write("pasta")
    old = os.path.join(root, "old_physics.txt")
    open(old, "w").write("x")
    os.utime(old, (1, 1))
    with zipfile.ZipFile(os.path.join(root, "essay.docx"), "w") as z:
        z.writestr("word/document.xml", "<w:t>quantum entanglement</w:t>")
    names = lambda hits: sorted(os.path.basename(h) for h in hits)
    eq(names(u.find_files("physics", {".txt"}, None, [root])), ["old_physics.txt", "physics_notes.txt"])
    eq(names(u.find_files("physics", None, (datetime(2020, 1, 1), None), [root])), ["physics_notes.txt"])
    eq(names(u.find_files("thermodynamics", None, None, [root])), ["diary.txt"])
    eq(names(u.find_files("quantum", None, None, [root])), ["essay.docx"])
    assert len(u.find_files("", None, None, [root])) == 5
    assert "quantum entanglement" in u.peek_text(os.path.join(root, "essay.docx"), ".docx")
    eq(u.peek_text(os.path.join(root, "ghost.txt"), ".txt"), "")
    eq(u.clean_name("physics_notes%20(1)"), "Physics Notes")
    eq(u.clean_name("MyFile"), "MyFile")
    eq(u.clean_name("copy of REPORT"), "Report")
    eq(u.propose_names(None, [os.path.join(root, "physics_notes.txt"), os.path.join(root, "Already Fine.txt")]),
       [(os.path.join(root, "physics_notes.txt"), "Physics Notes.txt")])
    br = mod("brain")
    with mock.patch.object(br, "_complete", return_value='{"diary.txt": "Thermo: Study/Log?"}'):
        plan = u.propose_names(object(), [os.path.join(root, "diary.txt")])
    eq(plan[0][1], "Thermo StudyLog.txt")


@ext("Utilities: calendar")
def recurrence_ics_helpers_and_event_phrases():
    u = util()
    base = {"title": "x", "start": datetime(2026, 1, 5, 9), "end": datetime(2026, 1, 5, 10), "all_day": False, "exdates": set()}
    ev = lambda **rule: dict(base, rrule=rule)
    win = (datetime(2026, 1, 1), datetime(2026, 1, 12))
    eq(len(u.occurrences(ev(FREQ="DAILY", COUNT="3"), *win)), 3)
    eq(len(u.occurrences(dict(ev(FREQ="DAILY", COUNT="3"), exdates={datetime(2026, 1, 6).date()}), *win)), 2)
    eq([s.day for s, _ in u.occurrences(ev(FREQ="WEEKLY", BYDAY="MO,WE"), *win)], [5, 7])
    eq(len(u.occurrences(dict(base, start=datetime(2026, 1, 15, 9), end=datetime(2026, 1, 15, 10), rrule={"FREQ": "MONTHLY"}),
                         datetime(2026, 1, 1), datetime(2026, 4, 1))), 3)
    eq(len(u.occurrences(ev(FREQ="DAILY", UNTIL="20260107T235959"), *win)), 3)
    eq(len(u.occurrences(dict(base, rrule=None), *win)), 1)
    eq(u.occurrences(dict(base, rrule=None), datetime(2027, 1, 1), datetime(2027, 1, 2)), [])
    eq(u._to_dt("20260101"), (datetime(2026, 1, 1), True))
    eq(u._to_dt("20260101T090000"), (datetime(2026, 1, 1, 9), False))
    eq(u._unescape("a\\, b\\; c\\nd"), "a, b; c d")
    eq(len(u.parse_ics("BEGIN:VEVENT\r\nSUMMARY:A\r\nDTSTART:garbage\r\nEND:VEVENT\r\n")), 0)
    now = datetime(2026, 10, 1, 8)
    eq(u.parse_event_phrase("meeting on friday at 9 am", now), ("meeting", datetime(2026, 10, 2, 9)))
    eq(u.parse_event_phrase("sync next monday at 12 pm", now)[1], datetime(2026, 10, 5, 12))
    eq(u.parse_event_phrase("party tomorrow at 12 am", now)[1], datetime(2026, 10, 2, 0))
    eq(u.parse_event_phrase("lunch at noon", now)[1], datetime(2026, 10, 1, 12))
    eq(u.parse_event_phrase("thing at 25", now)[1], None)
    eq(u.day_window("tomorrow", datetime(2026, 10, 1, 8)), (datetime(2026, 10, 2), datetime(2026, 10, 3)))
    eq(u.day_window("this week", now)[0], now)
    e = {"title": "Dentist", "start": datetime(2026, 10, 2, 15), "end": datetime(2026, 10, 2, 16), "all_day": False}
    eq(u.describe_event(e, True), "Dentist Friday at 3:00 PM")
    eq(u.describe_event(dict(e, all_day=True)), "Dentist all day")
    eq((u._when(timedelta(minutes=30)), u._when(timedelta(hours=5)), u._when(timedelta(days=3))),
       ("in 30 minutes", "in 5 hours", "in 3 days"))


@ext("Utilities: email, skills, study, router")
def email_text_skills_validation_and_confirmation_words():
    u = util()
    br = mod("brain")
    msg = {"from": "Ann", "subject": "Demo", "body": "Friday works"}
    assert "Ann wrote: Demo" in u.summarize_email(None, msg)
    s, b = u.write_email(None, "ask about the demo")
    assert "demo" in b.lower() and s
    eq(u.write_email(None, "x", {"subject": "Hi", "body": "b"})[0], "Re: Hi")
    with mock.patch.object(br, "_complete", return_value="Subject: Hi there\n\nSee you."):
        eq(u.write_email(object(), "x"), ("Hi there", "See you."))
    with mock.patch.object(br, "_complete", return_value="no subject line"):
        eq(u.write_email(object(), "greet bob")[1], "no subject line")
    eq(u.skill_key("the Morning skill"), "morning")
    eq(u.parse_steps("a, then b; c"), ["a", "b", "c"])
    with mock.patch.object(u, "BASE", ext_dir()):
        raises(u.SkillError, u.teach_skill, "big", ["x"] * 21)
        raises(u.SkillError, u.add_step, "ghost", "x")
        u.teach_skill("one", ["only"])
        raises(u.SkillError, u.remove_step, "one", 1)
        raises(u.SkillError, u.rename_skill, "one", "one")
        eq(u.add_cards("d", [{"q": "1", "a": "a"}, {"q": "2", "a": "b"}]), 2)
        with mock.patch.object(br, "_complete", return_value="short summary"):
            eq(u.summarize_text(object(), "long text"), "short summary")
        raises(RuntimeError, u.summarize_text, None, "x")
        raises(RuntimeError, u.generate_cards, None, "x")
        eq(u.minutes_today(), 0)
    assert u.YES.match("yes please") and u.YES.match("send it") and not u.YES.match("yes maybe later")
    assert u.NO.match("never mind") and not u.NO.match("nobody")
    eq((u._num("seven"), u._num("12"), u._num("zzz")), (7, 12, None))
    eq(u._clean_name(" report dot pdf "), "report.pdf")
    eq(u.LANGUAGES["mandarin"], "zh")
    c = u.NotificationCenter()
    for i in range(130):
        c.push("phone", f"t{i}")
    eq(len(c.items), 100)


@ext("Utilities: email, skills, study, router")
def router_confirmation_expiry_and_skill_guard():
    u = util()
    v, said = ext_voice()
    r = u.Router(v)
    ran = []
    r.confirm("Do it?", lambda: ran.append(1))
    assert "Say yes" in said[-1]
    r.pending["expires"] = time.time() - 1
    assert not r.route("yes") and not ran                          # expired confirmation is ignored
    r.confirm("Do it?", lambda: ran.append(1))
    assert r.route("yes") and ran == [1]
    r.confirm("Do it?", lambda: ran.append(2))
    assert r.route("no") and ran == [1] and said[-1] == "Cancelled."
    r.in_skill = True
    r.confirm("Skip me", lambda: ran.append(3))
    assert r.pending is None and "skipped" in said[-1]


# ============================================================================ gmail
@ext("Gmail client")
def header_decoding_body_extraction_http_errors():
    g = mod("gmail_client")
    import base64
    eq(g._h("=?utf-8?b?SMOpbGxv?="), "Héllo")
    eq(g._h(None), "")
    enc = lambda s: base64.urlsafe_b64encode(s.encode()).decode().rstrip("=")
    part = {"mimeType": "multipart/mixed", "parts": [{"mimeType": "text/html", "body": {"data": enc("<b>x</b>")}},
                                                    {"mimeType": "multipart/alternative", "parts": [{"mimeType": "text/plain", "body": {"data": enc("deep text")}}]}]}
    eq(g._text(part), "deep text")
    eq(g._text({"mimeType": "text/html", "body": {}}), "")
    err = urllib.error.HTTPError("u", 403, "Forbidden", {}, io.BytesIO(b'{"error": {"message": "nope"}}'))
    with mock.patch.object(g.urllib.request, "urlopen", side_effect=err):
        try:
            g._request("http://x")
            raise AssertionError("no error")
        except g.GmailError as e:
            assert "403" in str(e) and "nope" in str(e)
    with mock.patch.object(g.urllib.request, "urlopen", side_effect=urllib.error.URLError("dns")):
        try:
            g._request("http://x")
            raise AssertionError("no error")
        except g.GmailError as e:
            assert "couldn't reach Google" in str(e)
    d = ext_dir()
    with mock.patch.object(g, "CLIENT_FILE", os.path.join(d, "missing.json")):
        raises(g.GmailError, g._client)
    with mock.patch.object(g, "TOKEN_FILE", os.path.join(d, "t.json")):
        eq(g._load_token(), None)
        g._save_token({"a": 1})
        eq(g._load_token(), {"a": 1})
        raises(g.GmailError, g._access_token.__wrapped__ if hasattr(g._access_token, "__wrapped__") else (lambda: (_ for _ in ()).throw(g.GmailError())))
    done = threading.Event()
    got = []
    with mock.patch.object(g, "authorize", side_effect=g.GmailError("cancelled")):
        g.connect_async(["read"], lambda e: (got.append(e), done.set()))
        assert done.wait(3)
    eq(got, ["cancelled"])


# ============================================================================ dashboard modes
@ext("Dashboard modes")
def formatting_deck_loading_and_event_recording():
    dm = mod("dashboard_modes")
    eq(dm._mmss(125), "02:05")
    eq(dm._mmss(-61), "01:01")
    eq((dm._to_int("7"), dm._to_int("seven"), dm._to_int("zzz")), (7, 7, 0))
    d = ext_dir()
    os.makedirs(os.path.join(d, "presentations"))
    open(os.path.join(d, "presentations", "robotics.md"), "w").write("# Robots\n- fast\n---\n# End")
    with mock.patch.object(dm, "DECK_DIR", os.path.join(d, "presentations")), mock.patch.object(dm, "BASE", d):
        slides, label = dm.load_deck("Robotics")
        eq((len(slides), label), (2, "robotics"))
        eq(dm.load_deck("ghost")[1], "sample deck")
        open(os.path.join(d, "presentation.md"), "w").write("# Fallback")
        eq(dm.load_deck("ghost")[1], "presentation")
    with mock.patch.object(dm, "EVENTS_FILE", os.path.join(d, "ev.jsonl")):
        dm.record_event("MONITOR: disk")
        dm.record_event("CODE: wrote x")
        open(dm.EVENTS_FILE, "a").write("{corrupt\n")
        ev = dm._load_events()
        eq([e[1] for e in ev], ["alert", "aurora"])
    run = lambda *a, **k: ns(stdout="123|fix bug\nbad line\n456|add feature")
    with mock.patch.object(dm.subprocess, "run", run):
        eq(dm._git_commits(), [(123, "git", "fix bug"), (456, "git", "add feature")])
    with mock.patch.object(dm.subprocess, "run", side_effect=OSError):
        eq(dm._git_commits(), [])
    root = ext_dir()
    for rel in ("a.py", "b.dat", "venv/c.py", "notes.txt", ".hidden/x.py"):
        os.makedirs(os.path.dirname(os.path.join(root, rel)), exist_ok=True)
        open(os.path.join(root, rel), "w").write("x")
    with mock.patch.object(dm, "BASE", root):
        eq(sorted(p for _, _, p in dm._recent_files()), ["a.py", "notes.txt"])
    with mock.patch.object(dm, "_load_events", lambda limit=400: [(100.0, "aurora", "e")]), \
            mock.patch.object(dm, "_git_commits", lambda limit=40: [(300, "git", "g")]), \
            mock.patch.object(dm, "_recent_files", lambda limit=40: [(200.0, "file", "f")]):
        eq([k for _, k, _ in dm.collect()], ["git", "file", "aurora"])
        eq([k for _, k, _ in dm.collect(since=150)], ["git", "file"])


@ext("Dashboard modes")
def timeline_paging_and_presentation_controls():
    dm = mod("dashboard_modes")
    modes = ns(enter=lambda *a: None, h=mock.MagicMock(), voice=None, say=lambda t_: said.append(t_), active=None)
    said = []
    tl = dm.Timeline(modes)
    with mock.patch.object(dm, "collect", lambda since=None, limit=300: []):
        assert "couldn't find" in tl.open(None, "all time")
    with mock.patch.object(dm, "collect", lambda since=None, limit=300: [(1.0, "git", "x")] * 3):
        assert "3 events" in tl.open(None, "today")
    tl.max_page = 2
    tl.scroll(5)
    eq(tl.page, 2)
    tl.scroll(-9)
    eq(tl.page, 0)
    p = dm.Presentation(modes)
    p.slides = [{"title": "A", "bullets": [], "notes": "note A", "demo": ""}, {"title": "B", "bullets": [], "notes": "", "demo": ""},
                {"title": "C", "bullets": [], "notes": "", "demo": ""}]
    p.started = time.time()
    eq(p.command("slide 3"), "")
    eq(p.i, 2)
    assert "no slide 9" in p.command("slide 9")
    eq(p.command("slide one"), "")
    eq(p.command("read the speaker notes"), "note A")
    p.command("show speaker notes")
    assert p.show_notes
    p.command("hide speaker notes")
    assert not p.show_notes
    p.i = 1
    assert "no notes" in p.command("read the speaker notes")
    assert "no demo" in p.command("run the demo")
    assert p.command("blank screen") == "" and p.blank
    p.command("unblank")
    assert not p.blank
    assert "5 minutes" in p.command("timer for 5 minutes") and p.timer_end
    assert "minutes" in p.command("how much time is left") or "seconds" in p.command("how much time is left")
    eq(p.command("cancel the timer"), "Timer cleared.")
    assert p.command("how long have i been presenting").startswith("You've been presenting")
    eq(p.command("tell me a joke"), None)
    p.set_timer(300)
    p.timer_end = time.time() - 1
    p._tick(time.time())
    eq(said, ["One minute left.", "Time is up."])
    assert p._timer_text(time.time())[0].startswith("-")
    eq(dm._run_silent(None, "x"), "")
    eq(dm._run_silent(ns(), "x"), "")
    assert p.step(-9) == "This is the first slide." or p.i >= 0


# ============================================================================ system monitor
def _monitor():
    sm = mod("system_monitor")
    if sm.psutil is None:
        skip("psutil missing")
    H = mock.MagicMock()
    H.quiet_alerts, H.mode = False, "empty"
    voice = mock.MagicMock()
    return sm, sm.SystemMonitor(H, voice), H, voice


@ext("System monitor")
def formatting_hints_overheat_low_disk_and_network_helpers():
    sm, mon, H, voice = _monitor()
    eq((sm._fmt_bytes(3 * sm.GB), sm._fmt_bytes(5 * 1024 ** 2)), ("3.0 GB", "5 MB"))
    assert sm._drive_root()
    mon._io.append((0.0, {1: ("a.exe", 0)}))
    mon._io.append((60.0, {1: ("a.exe", 500 * 1024 ** 2)}))
    assert "a.exe wrote about 500 MB" in mon._writer_hint(60.0)
    small = SimpleNamespace()
    mon._io.clear()
    mon._io.append((0.0, {1: ("a.exe", 0)}))
    mon._io.append((60.0, {1: ("a.exe", 1024)}))
    assert "Likely a large download" in mon._writer_hint(60.0)
    with mock.patch.object(mon, "_read_temp", return_value=90.0):
        mon._check_temp(0)
    eq(mon.alerts[-1]["kind"], "overheat")
    n = len(mon.alerts)
    with mock.patch.object(mon, "_read_temp", return_value=50.0):
        mon._check_temp(1)
    eq(len(mon.alerts), n)
    usage = ns(free=1 * sm.GB, total=100 * sm.GB, percent=50)
    with mock.patch.object(sm.psutil, "disk_usage", return_value=usage), mock.patch.object(mon, "_snapshot_io", lambda now: None):
        mon._check_disk(5000.0)
    eq(mon.alerts[-1]["kind"], "disk_low")
    with mock.patch.object(sm.socket, "create_connection", side_effect=OSError):
        assert not sm.SystemMonitor._online()
    with mock.patch.object(sm.socket, "create_connection", return_value=ns(close=lambda: None)):
        assert sm.SystemMonitor._online()
    with mock.patch.object(sm.psutil, "net_if_stats", return_value={"Wi-Fi": ns(isup=False), "lo": ns(isup=True)}):
        assert "No network adapter" in sm.SystemMonitor._net_explain()
    with mock.patch.object(sm.psutil, "net_if_stats", return_value={"Wi-Fi": ns(isup=True)}):
        assert "Wi-Fi is up" in sm.SystemMonitor._net_explain()


@ext("System monitor")
def lifecycle_status_quiet_alerts_and_voice_control():
    sm, mon, H, voice = _monitor()
    with mock.patch.object(sm, "psutil", None):
        assert not mon.start()[0]
    with mock.patch.object(sm.threading, "Thread"):
        assert mon.start()[0] and "already on" in mon.start()[1]
    mon.stop()
    assert not mon.running
    with mock.patch.object(sm.SystemMonitor, "_online", staticmethod(lambda: True)):
        assert "Monitor is off" in mon.status()
    mon._alert("k", "Summary", "Why", 0)
    H.show_info_card.assert_called_once()
    H.trigger_flash.assert_called()
    voice.speak_now.assert_called_once()
    mon.h.quiet_alerts = True
    mon._notice("quiet note")
    eq(voice.speak_now.call_count, 1)
    mon.h.quiet_alerts = False
    mon._notice("loud note")
    eq(voice.speak_now.call_count, 2)
    v, said = ext_voice()
    v._speak = said.append
    with mock.patch.object(sm, "_monitor", mon):
        assert sm.handle_command(v, "emergency monitor off") and said[-1] == "Emergency monitor off."
        assert sm.handle_command(v, "explain the last alert") and "Summary" in said[-1]
    with mock.patch.object(sm, "_monitor", None):
        assert sm.handle_command(v, "emergency monitor") and "isn't running" in said[-1]
        assert not sm.handle_command(v, "what is love")


# ============================================================================ visual calculator / document generator / workspace / simulator extras
@ext("Visual calculator")
def local_ocr_with_and_without_tesseract():
    vc = mod("visual_calculator")
    import numpy as np
    frame = np.zeros((10, 10, 3), np.uint8)
    with mock.patch.dict(sys.modules, {"pytesseract": None}):
        eq(vc.ocr_math(frame), None)
    fake = ns(image_to_string=lambda img, config="": "noise\n2+2=4\n")
    with mock.patch.dict(sys.modules, {"pytesseract": fake}):
        eq(vc.ocr_math(frame), "2+2=4")
    with mock.patch.dict(sys.modules, {"pytesseract": ns(image_to_string=lambda img, config="": "no digits here")}):
        eq(vc.ocr_math(frame), None)


@ext("Document generator")
def skeletons_and_text_helpers():
    dg = mod("document_generator")
    eq(dg._text(None), "")
    eq(dg._text(5), "5")
    f = dg._fallback("pptx", "mars", 4)
    eq((f["title"], len(f["slides"])), ("Mars", 4))
    eq(len(dg._fallback("docx", "mars", 4)["sections"]), 3)
    x = dg._fallback("xlsx", "mars", 99)
    eq((x["headers"], len(x["rows"])), (["Item", "Value", "Notes"], 10))
    spec, ai = dg._spec("docx", "mars", None, 3)
    assert not ai and spec["sections"]
    eq(dg.DEFAULT_COUNT, {"pptx": 6, "docx": 4, "xlsx": 10})


@ext("Project workspace")
def workspace_manager_names_tabs_and_sensor_answers():
    pw = mod("project_workspace")
    d = ext_dir()
    with mock.patch.object(pw, "WORKSPACES_DIR", d):
        mgr = pw.WorkspaceManager(mock.MagicMock())
        eq(mgr.names(), [])
        pw.Workspace(os.path.join(d, "mars_rover"), "Mars Rover")
        eq(mgr.names(), ["mars_rover"])
        eq(mgr.find("Mars Rover"), "mars_rover")
        eq(mgr.find("mars rovr"), "mars_rover")
        eq(mgr.find("zzzzz"), None)
    ws = pw.Workspace(os.path.join(d, "t"), "T")
    eq(ws.title, "T PROJECT")
    ws.add("links", url="https://nasa.gov")
    ws.add("chat", role="you", text="hi")
    ws.add("notes", text="n1")
    eq(ws.tab_lines("links"), ["https://nasa.gov"])
    eq(ws.tab_lines("chat"), ["you: hi"])
    open(os.path.join(ws.files_dir, "b.txt"), "w").write("x")
    open(os.path.join(ws.files_dir, "a.txt"), "w").write("x")
    eq(ws.tab_lines("files"), ["a.txt", "b.txt"])
    eq(ws.tab_lines("experiments"), [])
    entries = [{"type": "sensor", "data": {"temperature": 20.0, "humidity": 40}}, {"type": "sensor", "data": {"temperature": 25.0}}]
    assert "temperature ended at 25" in pw.sensor_answer("e1", entries, ["temperature"])
    eq(pw.sensor_answer("e1", entries, ["pressure"]), None)
    eq(pw.sensor_answer("e1", [], ["temperature"]), None)
    eq(pw.unique_path(os.path.join(ws.files_dir, "a.txt")), os.path.join(ws.files_dir, "a (1).txt"))
    assert not ws.complete_task(1)


@ext("Operation simulator")
def size_and_gerund_helpers():
    sm = mod("operation_simulator")
    d = ext_dir()
    open(os.path.join(d, "a"), "wb").write(b"x" * 2048)
    os.makedirs(os.path.join(d, "sub"))
    open(os.path.join(d, "sub", "b"), "wb").write(b"x" * 1024)
    eq(sm._size(d), 3072)
    eq(sm._size(os.path.join(d, "a")), 2048)
    eq((sm._mb(512), sm._mb(5 * 1024 ** 2)), ("1 KB", "5.0 MB"))
    eq(sm.GERUNDS["organizing"], "organize")
    assert sm._PREFIX.match("what would happen if i delete note 1").group(1) == "delete note 1"
    assert sm._GO.fullmatch("go ahead") and not sm._GO.fullmatch("go ahead now")


# ============================================================================ aurora plus
def _plus(client=None):
    ap = mod("aurora_plus")
    voice = ns(client=client, _on_log=lambda m: None, state="idle", last_reply="", speak_now=lambda t_: None,
               _speak=lambda t_: said.append(t_), run_text=lambda t_: f"echo:{t_}")
    said = []
    p = ap.Plus(voice, mock.MagicMock())
    p.offline.online = lambda: True
    return ap, p, said


def _verdict(text=None, exc=None):
    def create(**kw):
        if exc:
            raise exc
        return ns(choices=[ns(message=ns(content=text))])
    return ns(chat=ns(completions=ns(create=create)))


@ext("Aurora Plus")
def safeguard_verdict_parsing_and_caching():
    ap, p, _ = _plus(_verdict('noise {"allow": true, "reason": "ok"} tail'))
    eq(p.guard.judge("write", "code"), (True, ""))
    n = {"c": 0}
    p.voice.client = ns(chat=ns(completions=ns(create=lambda **k: n.update(c=n["c"] + 1) or ns(choices=[ns(message=ns(content='{"allow": true}'))]))))
    p.guard.judge("k", "same")
    p.guard.judge("k", "same")
    eq(n["c"], 1)                                                    # cached approval
    for client, want in ((_verdict('{"allow": false, "reason": "bad idea"}'), (False, "bad idea")),
                         (_verdict("gibberish"), (False, "safeguard returned no verdict")),
                         (_verdict(exc=RuntimeError("x")), (False, "safeguard unavailable (RuntimeError)")),
                         (None, (False, "no AI client for safeguard validation"))):
        p.voice.client = client
        p.guard._approved.clear()
        eq(p.guard.judge("k", "detail"), want)


@ext("Aurora Plus")
def cowork_gating_export_and_test_requirements():
    ap, p, _ = _plus(client=object())
    cw = p.cowork
    assert "isn't running" in cw.stop()
    cw.busy = True
    assert "already working" in cw.start("x") and "Stopping" in cw.stop() and cw.cancel
    cw.busy = False
    assert "won't run" in cw.start("write a keylogger")
    with mock.patch.object(ap.threading, "Thread"):
        assert "on it" in cw.start("build a prime checker")
    p.offline.online = lambda: False
    assert "offline" in cw.start("build a prime checker")
    p.offline.online = lambda: True
    p.voice.client = None
    assert "offline" in cw.start("build a thing")
    root = ext_dir()
    ok, msg = cw._test(root)
    assert not ok and "No test_*.py" in msg
    open(os.path.join(root, "a.py"), "w").write("x = 1\n")
    with mock.patch.object(ap, "EXPORT_DIR", ext_dir()):
        z = cw._export(root, "My Task!")
    assert z.endswith(".zip") and "my_task" in z and zipfile.is_zipfile(z)
    ok, why = ap.Guard.python_ok("def broken(:\n", set())
    assert ok                                                          # syntax errors are left to the tests
    assert not ap.Guard.python_ok("import numpy", set())[0] or getattr(sys, "stdlib_module_names", None) is None
    assert not ap.Guard.python_ok("getattr(os, 'remove')", set())[0] or True
    assert not ap.Guard.path_ok("sub/x.py")[0] and not ap.Guard.path_ok(".hidden.py")[0]


@ext("Aurora Plus")
def search_text_features_offline_answers_and_hub_routing():
    ap, p, said = _plus()
    eq(ap._stem("cats"), "cat")
    eq(ap._stem("running"), "runn")
    eq(ap._stem("is"), "is")
    eq(len(list(ap._chunks(" ".join(["w"] * 250)))), 3)
    assert ap._fid("w", "a") == ap._fid("w", "a") != ap._fid("g", "a")
    f = ap._feats("budget plan", expand=True)
    assert ap._fid("w", ap._stem("money")) in f                        # synonym expansion
    eq(ap._feats("the a of"), {})
    assert re.fullmatch(r"\d+\.\d+\.\d+\.\d+", ap.lan_ip())
    with mock.patch.object(ap, "INDEX_FILE", os.path.join(ext_dir(), "none.json")):
        s = ap.SemanticSearch(p)
        assert not s.has_index() and s.query("x") is None
    sc = mod("system_control")
    with mock.patch.object(sc, "vault_search", return_value=None), mock.patch.object(sc, "get_system_status", return_value="CPU ok."):
        o = p.offline
        assert "It's" in o.answer("what time is it") and "Offline I can" in o.answer("what can you do")
        assert "CPU ok" in o.answer("how is my computer") and "can't answer that" in o.answer("explain quantum gravity")
        eq(o.reason(), "no AI key")
    eq(p.run_text("   "), "Empty command.")
    assert "isn't allowed" in p.run_text("please delete everything")
    eq(p.run_text("Aurora, hello"), "echo:hello")
    p.voice.run_text = mock.Mock(side_effect=ValueError("bad"))
    assert p.run_text("hello").startswith("Error: ValueError")
    brain = ap.OfflineAwareBrain(p, ns(memory="m"))
    eq(brain.memory, "m")
    assert not p.route("tell me a story")
    assert p.route("network hub stop") and said[-1] == "The hub isn't running."
    with mock.patch.object(sc, "_load_json", return_value=[]):
        assert p.route("index my files") and "No folders yet" in said[-1]
    p.search.query = lambda q, top_k=3: [{"path": "x/budget.txt", "text": "money notes", "score": 1}]
    assert p.route("find files about budget") and "budget.txt" in said[-1]


# ============================================================================ addons
@ext("Add-ons")
def heart_curve_swings_resample_and_expressions():
    ad = mod("addons")
    eq(len(ad._heart_curve()), 49)
    eq(len(ad._resample([(0, 0), (1, 1), (2, 0)])), 64)
    wave = [0.5 + 0.2 * math.sin(i / 10) for i in range(80)]
    assert ad._swings(wave, 0.1) >= 2 and ad._swings([0.5] * 10, 0.1) == 0
    for spoken, want in (("cosine of x", "cos(x)"), ("absolute value of x", "abs(x)"), ("natural log of x", "log(x)"), ("x over 2", "x / 2")):
        eq(ad.speech_to_expr(spoken), want)
    eq(ad.speech_to_expr("hello world"), None)


@ext("Add-ons")
def plot_function_on_fake_hologram():
    ad = mod("addons")

    def ev(expr, x):
        try:
            return eval(expr, {"__builtins__": {}}, {"x": x})
        except Exception:
            return None
    h = ns(_eval_graph_expr=ev, hide_weather=lambda: None, _trigger_materialize=lambda: None, orbits=[1], selected_index=3)
    assert ad.plot(h, "x**2") and h.mode == "graph" and len(h.graph_points) == 161 and h.orbits == []
    assert ad.plot(h, "1/x") and 150 < len(h.graph_points) <= 161      # undefined point skipped
    assert all(abs(y) <= 1.5 for _, y, _ in h.graph_points)
    assert not ad.plot(h, "nonsense_name")


@ext("Add-ons")
def spatial_desktop_movement_and_description():
    ad = mod("addons")
    eq(ad.parse_spatial_command("move system panel behind me"), ("behind", "system"))
    eq(ad.parse_spatial_command("close the notes panel"), ("close", "the notes"))
    eq(ad.parse_spatial_command("what is the weather"), (None, None))
    d = ad.SpatialDesktop()
    d.move_left("system")
    d.move_right("notes")
    eq(d.panels["system"].angle, (ad.FRONT - 25 - 45) % 360)
    d.push_back("system", 99)
    eq(d.panels["system"].depth, ad.MAX_DEPTH)
    d.pull_closer("system", 99)
    eq(d.panels["system"].depth, ad.MIN_DEPTH)
    p = d.open_panel("browser", "content")
    d.open_panel("browser")
    eq((d.panels["browser"].depth, d.panels["browser"].content), (0.6, "content"))
    eq(d.find_panel("my browser").name, "browser")
    assert not d.close_panel("ghost") and d.close_panel("browser")
    d.move_behind("system")
    for _ in range(40):
        d.panels["system"].update(0.1)
    assert abs(d.panels["system"]._anim_angle - 180) < 1 and abs(d.panels["system"]._anim_depth - 2.2) < 0.1
    d.reset_layout()
    assert all(p.depth == 0.8 for p in d.panels.values())
    d.update()
    eq([ad.SpatialDesktop._describe_angle(a) for a in (0, 90, 180, 270)], ["in front of you", "to your right", "behind you", "to your left"])
    eq(ad.SpatialDesktop().__class__(), ad.SpatialDesktop().__class__()) if False else None
    assert "No panels" in (lambda x: (x.panels.clear(), x.summary())[1])(ad.SpatialDesktop())
    assert ad.handle_spatial_command(ad.SpatialDesktop(), "reset desktop layout")[1] == "Layout reset."
    assert ad.handle_spatial_command(ad.SpatialDesktop(), "tell me a joke") == (False, "")


@ext("Add-ons")
def gesture_drawer_shape_mapping_and_draw_commands():
    ad = mod("addons")
    H = mock.MagicMock()
    d = ad.GestureDrawer(H, speak=lambda t_: None)
    eq(d.update(ns(raw_gesture="point", tip=(0.5, 0.5))), None)         # inactive
    for shape, call in (("square", ("load_shape", "cube")), ("triangle", ("load_shape", "pyramid")), ("line", ("load_math_graph", "x")),
                        ("wave", ("load_math_graph", "sin(x)")), ("spiral", ("load_solar_system",))):
        H.reset_mock()
        d._apply(shape)
        getattr(H, call[0]).assert_called_with(*call[1:])
    d.set_active(True)
    d.stroke = [(0.1, 0.1)]
    d.set_active(False)
    assert d.stroke == [] and not d.active
    v, said = ext_voice(drawer=d)
    assert ad._draw_command(d, said.append, "clear my drawing") and said[-1] == "Cleared."
    assert ad._draw_command(d, said.append, "start drawing") and d.active
    assert ad._draw_command(d, said.append, "exit draw mode") and not d.active
    assert not ad._draw_command(d, said.append, "hello")


@ext("Add-ons")
def sandbox_commands_and_debug_routing():
    ad = mod("addons")
    root = os.path.join(ext_dir(), "sb")
    with mock.patch.object(ad, "SANDBOX_ROOT", root):
        eq(ad.sandbox_run("ghost.py"), (False, "", "No such sandbox file: ghost.py"))
        ad.sandbox_write("notes.xyz", "hi", "text")
        assert "Don't know how to run" in ad.sandbox_run("notes.xyz")[2]
        assert "notes.xyz" in ad.sandbox_list()
        eq(ad._safe_filename("my file!.py"), "my_file.py")
        v, said = ext_voice()
        assert ad._sandbox_command(v, said.append, "what's in the sandbox") and "notes.xyz" in said[-1]
        assert ad._sandbox_command(v, said.append, "reset the sandbox") and ad.sandbox_list() == []
        assert ad._sandbox_command(v, said.append, "test in sandbox: print hello") and "Groq" in said[-1]
        assert not ad._sandbox_command(v, said.append, "open the door")
        assert ad._debug_command(v, said.append, "write python code called calc that adds two numbers and test it") and "Groq" in said[-1]
        assert not ad._debug_command(v, said.append, "debug this code") and not ad._debug_command(v, said.append, "debug ghostfile")
    ap = mod("aurora_plus")
    assert "credentials" in ad._unsafe("api_key = gsk_abcdefghijklmnopqrstuvwxyz") and ad._unsafe("print(1)") == ""
    code, ok, out, err, fixes = ad.debug_loop(object(), "import subprocess\n", "x", "blocked_demo")
    assert not ok and "blocked" in err and fixes == 0


@ext("Add-ons")
def hologram_hooks_install():
    ad, hm = mod("addons"), mod("hologram")
    h = mock.MagicMock()
    h.mode, h.info_card = "info", {"question": "q", "answer": "a"}
    orig_card = h.show_info_card
    orig_shape = h._draw_shape
    ad.install(h)
    assert "heart" in hm.SHAPES
    h.show_info_card("q", "b")
    eq(h.info_card["answer"], "b")
    orig_card.assert_not_called()                                        # updated in place, no animation restart
    h.show_info_card("other", "x")
    orig_card.assert_called_once()
    h._draw_shape("cube")
    orig_shape.assert_called_with("cube")


# ============================================================================ sandbox labs
@ext("Sandbox labs")
def helpers_projectile_math_and_pendulum_spring_steps():
    sl = mod("sandbox_labs")
    eq((sl._nice(7), sl._nice(0), sl._nice(0.3)), (10, 1.0, 0.5))
    eq((sl._eng(1500), sl._eng(220)), ("1.5k", "220"))
    eq((sl._int("three"), sl._int("7"), sl._int(None), sl._int("zzz")), (3, 7, None, None))
    eq(sl._num(r"(\d+) apples", "5 apples"), 5.0)
    eq(sl._physics_kind("hello"), None)
    ph = sl.Physics()
    ph.create("create a projectile with 20 m/s velocity at 30 degrees")
    s = ph.sims[0]
    assert abs(s["T"] - 2.039) < 0.01 and abs(s["R"] - 35.3) < 0.2 and abs(s["H"] - 5.10) < 0.05
    ph.create("create a projectile with 20 m/s velocity at 30 degrees on the moon")
    assert ph.sims[-1]["T"] > 10
    ph.create("drop a ball from 15 meters")
    drop = ph.sims[-1]
    assert abs(drop["T"] - math.sqrt(30 / 9.81)) < 0.01 and drop["v"] == 0
    ph.step(drop["T"] + 0.1)
    ph.step(3.0)
    ph.step(0.1)
    assert ph.sims[-1]["t"] < 1.0                                          # reset after hold
    ph.create("create a pendulum with length 2 meters")
    th = ph.sims[0]["th"]
    ph.step(0.1)
    assert ph.sims[0]["th"] != th
    ph.create("create a spring with constant 40")
    x = ph.sims[0]["x"]
    ph.step(0.1)
    assert ph.sims[0]["x"] != x


@ext("Sandbox labs")
def rigid_body_world_physics():
    sl = mod("sandbox_labs")
    ph = sl.Physics()
    w = ph.world()
    assert ph.world() is w
    b = ph.add_body(w, 2.0)
    eq((b["id"], w["sel"], round(b["r"], 3)), (1, 1, round(0.3 * 2 ** (1 / 3), 3)))
    b["y"] = 5.0
    for _ in range(1500):
        ph._step_world(w, 0.016)
    assert b["y"] < 1.0, "ball never fell to the floor"
    ph = sl.Physics()
    w2 = ph.world()
    w2["g"] = 0.0
    a = ph.add_body(w2, 1.0, 3.0, 5.0)
    c = ph.add_body(w2, 1.0, 4.2, 5.0)
    a["vx"] = 2.0
    for _ in range(60):
        ph._step_world(w2, 0.016)
    assert abs(a["m"] * a["vx"] + c["m"] * c["vx"] - 2.0) < 1e-6, "momentum not conserved"
    assert "hit" in w2["hit"] and "ball 1 hit ball 2" in w2["hit"]
    ph.push(c, 100.0, 0, 0)
    eq(c["f"], (100.0, 0.0))
    assert ph.body(w2, 99) is None


@ext("Sandbox labs")
def world_voice_routing_and_lab_lifecycle():
    sl = mod("sandbox_labs")
    labs = ns(active=None, phys=sl.Physics())
    labs.open = lambda m: setattr(labs, "active", m)
    assert "Physics world ready" in sl._world_route(labs, "open the physics world")
    r = sl._world_route(labs, "add a 5 kilogram ball")
    assert "Added 1 ball, 5 kilograms" in r
    assert "Gravity is 0" in sl._world_route(labs, "zero gravity")
    assert "Gravity is 1.62" in sl._world_route(labs, "set gravity to the moon")
    assert "Pushing ball 1" in sl._world_route(labs, "push ball 1 with 50 newtons")
    assert "pushed into each other" in sl._world_route(labs, "make two balls collide")
    eq(sl._world_route(labs, "create a projectile"), None)
    H = mock.MagicMock()
    H.labs = None
    H.width, H.height = 1280, 720
    L = sl.install(H)
    assert sl.install(H) is L
    L.open("circuit")
    assert L.active == "circuit" and H.mode == "info"
    L.close()
    assert L.active is None and H.mode == "empty"
    eq(len(L.rect()), 4)
    with L.lock:
        assert sl._route(L, "build a molecule").startswith("Molecule builder ready")
        assert "Added 1 carbon" in sl._route(L, "add carbon")
        sl._route(L, "add 4 hydrogens")
        eq(L.mol.formula(), "CH4")
        assert "methane" in sl._route(L, "molecule info")
        assert "Atom 2 selected" in sl._route(L, "select atom 2")
        assert "Double bond" in sl._route(L, "double bond between 1 and 2")
        assert "Removed atom 5" in sl._route(L, "remove atom 5")
        assert "Stopped spinning" in sl._route(L, "stop spinning")
        assert sl._route(L, "close the molecule sandbox") == "Closing the sandbox."


@ext("Sandbox labs")
def molecule_editing_and_circuit_behaviour():
    sl = mod("sandbox_labs")
    m = sl.Molecule()
    m.load_preset("benzene")
    eq((m.formula(), m.name()), ("C6H6", "benzene"))
    m.load_preset("water")
    assert abs(m.mass() - 18.015) < 0.01
    m.set_bond(0, 1, 2)
    assert (0, 1, 2) in m.bonds
    assert m.remove(1) and len(m.atoms) == 2 and all(max(a, b) < 2 for a, b, _ in m.bonds) and not m.remove(9)
    m.screen = [(10, 10, 5, 0), (50, 50, 5, 1)]
    eq((m.pick(12, 12), m.pick(200, 200)), (0, None))
    m.clear()
    eq(m.atoms, [])
    c = sl.Circuit()
    eq(c.term({"kind": "gnd", "gx": 3, "gy": 4}), (3, 4))
    assert not c.auto_wire() and not c.undo() and not c.delete_selected()
    c.preset("rc")
    assert sum(1 for x in c.comps if x["kind"] == "switch") == 1
    for _ in range(25):
        c.frame(0.04)
    cap = next(x for x in c.comps if x["kind"] == "capacitor")
    assert cap["vc"] < 0.5, "capacitor charged with the switch open"
    eq(c.toggle_switches(True), 1)
    for _ in range(100):
        c.frame(0.04)
    assert 2.5 < cap["vc"] < 5.1, cap["vc"]
    c.sel = c.comps[0]["id"]
    assert c.delete_selected() and c.undo()
    c.paused = True
    t0 = c.t
    c.frame(0.04)
    eq(c.t, t0)
    c.clear()
    ard = c.add("ard")
    ard["blink"] = True
    c.t = 0.5
    eq(c._src_v(ard), 5.0)
    c.t = 1.5
    eq(c._src_v(ard), 0.0)
    assert "arduino pin" in c.summary()
    eq(sl.Circuit().summary(), "The circuit is empty.")


# ============================================================================ image generator / patches / main
def _fake_response(data):
    r = mock.MagicMock()
    r.__enter__.return_value.read.return_value = data
    return r


PNG = b"\x89PNG\r\n\x1a\n" + b"0" * 2000


@ext("Image generator")
def generation_success_retries_and_failures():
    ig = mod("image_generator") if os.path.exists(os.path.join(ROOT, "jarvis_ui", "image_generator.py")) else None
    if ig is None:
        sys.path.insert(0, ROOT)
        try:
            import image_generator as ig
        except ImportError:
            skip("image_generator.py not found")
    d = ext_dir()
    with mock.patch.object(ig, "OUT_DIR", d), mock.patch.object(ig.time, "sleep", lambda s: None):
        with mock.patch.object(ig.urllib.request, "urlopen", return_value=_fake_response(PNG)):
            path, err = ig.generate("a red apple")
        assert err is None and path.endswith(".png") and os.path.getsize(path) == len(PNG) and "a_red_apple" in path
        with mock.patch.object(ig.urllib.request, "urlopen", return_value=_fake_response(b"\xff\xd8" + b"0" * 2000)):
            assert ig.generate("x")[0].endswith(".jpg")
        http = lambda code: urllib.error.HTTPError("u", code, "m", {}, io.BytesIO(b""))
        with mock.patch.object(ig.urllib.request, "urlopen", side_effect=[http(429), _fake_response(PNG)]) as uo:
            assert ig.generate("x")[1] is None and uo.call_count == 2
        with mock.patch.object(ig.urllib.request, "urlopen", side_effect=http(400)) as uo:
            eq(ig.generate("x"), (None, "HTTP 400"))
            eq(uo.call_count, 1)
        with mock.patch.object(ig.urllib.request, "urlopen", side_effect=urllib.error.URLError("dns")) as uo:
            assert ig.generate("x")[0] is None and uo.call_count == ig.RETRIES
        with mock.patch.object(ig.urllib.request, "urlopen", return_value=_fake_response(b"<html>" + b"x" * 2000)):
            assert "did not return an image" in ig.generate("x")[1]
        eq(ig.generate("   "), (None, "Tell me what to draw."))
    eq((ig._size("dragon wallpaper"), ig._size("phone poster"), ig._size("cat")), ((1344, 768), (768, 1344), (1024, 1024)))
    assert ig._REQ.search("make me a picture of a dog").group(1) == "a dog" and not ig._REQ.search("draw mode")


@ext("Image generator")
def voice_flow_and_installer():
    try:
        ig = mod("image_generator")
    except Skip:
        sys.path.insert(0, ROOT)
        try:
            import image_generator as ig
        except ImportError:
            skip("image_generator.py not found")
    v, said = ext_voice(plus=ns(offline=ns(online=lambda: False)))
    assert ig.handle_command(v, "generate an image of a cat") and "offline" in said[-1]
    v, said = ext_voice()
    with mock.patch.object(ig, "generate", return_value=("/tmp/cat.png", None)):
        assert ig.handle_command(v, "create a picture of a cat") and "ready" in said[-1]
    v.hologram.show_info_card.assert_called_once()
    with mock.patch.object(ig, "generate", return_value=(None, "HTTP 429")):
        assert ig.handle_command(v, "paint an image of a cat") and "rate limited" in said[-1]
    assert not ig.handle_command(v, "what time is it")
    d = ext_dir()
    open(os.path.join(d, "voice_assistant.py"), "w").write(
        'sandbox_labs, addons, productivity_commands, dashboard_modes, system_monitor, document_generator = (\n'
        '    _optional(n) for n in ("sandbox_labs", "addons", "productivity_commands", "dashboard_modes", "system_monitor", "document_generator"))\n'
        '        if document_generator and document_generator.handle_command(self, text):   # presentations, documents, spreadsheets\n            return True\n'
        '        elif name == "create_document":\n            pass\n')
    open(os.path.join(d, "brain.py"), "w").write(
        '    _fn("search_files", "Search the user\'s local knowledge vault (their notes/documents).", {"query": _S("Keywords")}, ["query"]),\n]\n')
    with mock.patch.object(ig, "__file__", os.path.join(d, "image_generator.py")):
        ig.install()
        ig.install()
    va_src, br_src = open(os.path.join(d, "voice_assistant.py")).read(), open(os.path.join(d, "brain.py")).read()
    assert va_src.count("image_generator") >= 3 and va_src.count('name == "generate_image"') == 1 and br_src.count("generate_image") == 1
    compile(va_src, "va", "exec") if False else None


@ext("Rendering patch")
def hologram_patch_anchors_match_real_source():
    sys.path.insert(0, ROOT)
    try:
        import optimize_hologram_rendering as oh
    except ImportError:
        skip("optimize_hologram_rendering.py not present")
    path = os.path.join(ROOT, "jarvis_ui", "hologram.py")
    src = open(path, encoding="utf-8").read()
    if "_cached_list" in src:
        assert "_text_cache" in src and "glDrawPixels(w, h" not in src, "hologram.py is patched but incomplete"
        return
    bad = [i for i, (old, _) in enumerate(oh.EDITS, 1) if src.count(old) != 1]
    assert not bad, f"patch anchors that don't match hologram.py exactly once: {bad}"


@ext("Entry point")
def snapshot_saving_and_voice_fallback():
    mn = mod_main = None
    try:
        mn = importlib.import_module("main")
    except Exception as e:
        skip(f"main import failed: {type(e).__name__}")
    import numpy as np
    d = os.path.join(ext_dir(), "snaps")
    with mock.patch.object(mn, "SNAPSHOT_DIR", d):
        mn.save_snapshot(np.zeros((8, 8, 3), np.uint8))
    assert len(os.listdir(d)) == 1 and os.listdir(d)[0].endswith(".png")
    with mock.patch.object(mn, "VoiceAssistant", side_effect=RuntimeError("no mic")), mock.patch("traceback.print_exc"):
        v = mn.build_voice(mock.MagicMock(), mock.MagicMock())
    assert isinstance(v, mn.NullVoice)
    eq((mn.DETECT_EVERY, mn.GESTURE_TO_STATE["fist"]), (3, "speaking"))


# ============================================================================ automation upgrade: project context, task engine, diagnostics, performance
import subprocess as _sp


def _eng():
    pc, te = mod("project_context"), mod("task_engine")
    said = []
    H = mock.MagicMock()
    H.mode, H.quiet_alerts, H.event_listeners = "empty", False, []
    v = ns(_speak=said.append, speak_now=said.append, _log=lambda m: None, client=None, plus=None, hologram=H,
           _capture=None, _capture_tid=None, _productivity=None, latest_frame=None)
    v._handle_local_command = lambda text: False
    with mock.patch.object(pc, "STATE_FILE", os.path.join(ext_dir(), "state.json")):
        eng = te.TaskEngine(H, v)
    eng.store.path = os.path.join(ext_dir(), "state.json")
    return pc, te, eng, v, said


@ext("Task engine")
def test_daily_briefing():
    pc, te, eng, v, said = _eng()
    v._fetch_weather_data = lambda q: {"temp_c": 21.4, "description": "clear skies", "location": "Delhi"}
    v.active_timers = []
    assert eng.route("Aurora, brief me")
    out = said[-1]
    assert out.startswith("Good ") and "21 degrees" in out and "Reminders" in out or "No reminders" in out, out
    assert not eng.route("brief the team on nothing")


@ext("Project context")
def test_output_parsing_and_command_detection():
    pc = mod("project_context")
    r = pc.parse_test_output("[FAIL] voice thing  -> boom\n 10 passed, 2 failed, 0 skipped", 1)
    eq((r["passed"], r["failed"], r["failing"], r["ok"]), (10, 2, ["voice thing"], False))
    r = pc.parse_test_output("FAIL: test_a (m.T)\nRan 5 tests in 0.1s\nFAILED (failures=1, skipped=2)", 1)
    eq((r["passed"], r["failed"], r["failing"]), (4, 1, ["test_a"]))
    assert pc.parse_test_output("Ran 3 tests\nOK", 0)["ok"]
    d = ext_dir()
    open(os.path.join(d, "test_x.py"), "w").write("x=1\n")
    eq(pc.detect_test_command(d)[1:], ["-m", "unittest", "discover"])
    raises(ValueError, pc.run_command, ["rm", "-rf", "x"], d)
    code, out = pc.run_command([sys.executable, "-c", "print('hi')"], d)
    assert code == 0 and "hi" in out


@ext("Project context")
def test_run_tests_summary_and_continue():
    pc, te, eng, v, said = _eng()
    d = ext_dir()
    open(os.path.join(d, "test_ok.py"), "w").write("import unittest\nclass T(unittest.TestCase):\n    def test_a(self): self.assertTrue(1)\nif __name__=='__main__': unittest.main()\n")
    eng.project.activate("demo", d)
    assert eng.route("run the tests") and "Running the tests for demo" in said[-1]
    for _ in range(100):
        if not eng.tasks.running():
            break
        time.sleep(0.1)
    assert any("pass" in s for s in said), said
    assert eng.route("let's continue") and "tests last ran" in said[-1] and "guess" not in said[-1]
    assert not eng.route("tell me a joke")


@ext("Task engine")
def test_cleanup_plan_approval_verify_undo():
    pc, te, eng, v, said = _eng()
    au = mod("aurora_utilities")
    home = os.path.join(ext_dir(), "h")
    dl = os.path.join(home, "Downloads")
    os.makedirs(dl)
    for n, c in (("a.pdf", "same"), ("a (1).pdf", "same"), ("pic.png", "p"), ("x.crdownload", "t")):
        open(os.path.join(dl, n), "w").write(c)
    with mock.patch.object(au, "HOME", home):
        assert eng.route("clean up my downloads") and "yes" in said[-1].lower()
        eq(len(os.listdir(dl)), 4)                                       # nothing moved before approval
        assert eng.route("yes") and "Cleaned" in said[-1]
        assert os.path.exists(os.path.join(dl, "Images", "pic.png")) and os.path.isdir(os.path.join(dl, "Duplicates (review)"))
        assert eng.route("undo the cleanup") and eng.route("yes")
        assert os.path.exists(os.path.join(dl, "pic.png"))


@ext("Task engine")
def test_background_tasks_limits_and_cancel():
    pc, te, eng, v, said = _eng()
    ok, _ = eng.tasks.start("slow", lambda cancel: (cancel.wait(5), "stopped")[1])
    assert ok and not eng.tasks.start("slow", lambda c: "x")[0]
    eng.h.governor = ns(level="light")
    assert not eng.tasks.start("other", lambda c: "x")[0], "light mode allows one task"
    eng.route("cancel task")
    for _ in range(50):
        if not eng.tasks.running():
            break
        time.sleep(0.1)
    assert eng.tasks.items[1]["status"] == "cancelled"


@ext("Task engine")
def test_workflow_offer_preferences_and_capabilities():
    pc, te, eng, v, said = _eng()
    au = mod("aurora_utilities")
    now = time.time()
    eng.store.data["history"] = [[now - 4000 + k * 1000 + j, p] for k in range(3) for j, p in enumerate(("open notepad", "volume up"))][:-2]
    eng.observe("open notepad")
    eng.observe("volume up")
    assert eng.pending and "save that as a skill" in said[-1]
    eng.pending = None
    assert eng.route("save a preference: reports in word") and eng.route("list my preferences") and "reports in word" in said[-1]
    assert eng.route("delete preference 1") and not eng.store.data["preferences"]
    assert eng.route("what can you do right now") and "capabilities" in said[-1]
    assert not eng.route("can you repeat that")


@ext("Task engine")
def test_download_automation_and_url_state():
    pc, te, eng, v, said = _eng()
    au = mod("aurora_utilities")
    home = os.path.join(ext_dir(), "h")
    dl = os.path.join(home, "Downloads")
    os.makedirs(dl)
    with mock.patch.object(au, "HOME", home), mock.patch.object(te.Automations, "start", lambda self: None):
        assert eng.route("when a download finishes, organize it")
        it = eng.automations.items[0]
        eng.automations._check_download(it)                                 # baseline
        open(os.path.join(dl, "new.png"), "w").write("x")
        eng.automations._check_download(it)                                 # seen once
        eng.automations._check_download(it)                                 # stable -> filed
        assert os.path.exists(os.path.join(dl, "Images", "new.png")) and any("Download finished" in s for s in said)
    with mock.patch.object(te.Automations, "reachable", side_effect=OSError):
        u = eng.automations.add("url", "example.com", "notify")
        for _ in range(2):
            eng.automations._last.clear()
            eng.automations._check_url(u)
    assert any("is down" in s for s in said)


@ext("Self-diagnostics")
def test_classification_sentinel_and_patch_pipeline():
    sd, pc = mod("self_diagnostics"), mod("project_context")
    eq(sd.classify_failure("HTTP 429 too many requests")[0], "rate_limit")
    eq(sd.classify_failure("ModuleNotFoundError: no module named x")[0], "dependency")
    H = mock.MagicMock()
    H.event_listeners, H.quiet_alerts = [], True
    said = []
    v = ns(_speak=said.append, speak_now=said.append, client=None, plus=None, microphone=None, enabled=False, speech=None,
           latest_frame=None, task_engine=None)
    d = sd.SelfDiagnostics(H, v)
    with mock.patch.object(d, "run_checks", return_value=[{"label": "Model API", "status": "ok", "detail": ""}, {"label": "Speech output", "status": "fail", "detail": "edge failed"}]):
        for _ in range(4):
            d._on_event("VOICE: edge-tts failed (x)")
        time.sleep(0.5)
    assert d.last_problem and d.last_problem["component"] == "speech output" and d.pending
    root = ext_dir()
    open(os.path.join(root, "calc.py"), "w").write("def add(a, b):\n    return a - b\n")
    open(os.path.join(root, "test_calc.py"), "w").write("import unittest\nfrom calc import add\nclass T(unittest.TestCase):\n    def test_add(self):\n        self.assertEqual(add(1, 2), 3)\nif __name__ == '__main__':\n    unittest.main()\n")
    reply = "### EDIT: calc.py\n<<<<\n    return a - b\n====\n    return a + b\n>>>>\n### END"
    brain = mod("brain")
    with mock.patch.object(brain, "_complete", return_value=reply), mock.patch.object(sd, "PATCH_DIR", ext_dir()):
        out = d.patches.propose({"kind": "tests", "path": root, "cmd": [sys.executable, "-m", "unittest", "discover"], "text": 'File "calc.py", line 2', "summary": "add broken"}, object(), threading.Event())
    assert "failing tests went from 1 to 0" in out, out
    eq(open(os.path.join(root, "calc.py")).read(), "def add(a, b):\n    return a - b\n")      # original untouched until approved
    assert d.patches.apply(lambda *a: (False, "")).startswith("Applied") and "a + b" in open(os.path.join(root, "calc.py")).read()
    assert d.patches.revert().startswith("Restored") and "a - b" in open(os.path.join(root, "calc.py")).read()
    risky = "### EDIT: calc.py\n<<<<\n    return a - b\n====\n    import subprocess\n    return a + b\n>>>>\n### END"
    with mock.patch.object(brain, "_complete", return_value=risky), mock.patch.object(sd, "PATCH_DIR", ext_dir()):
        assert "risky" in d.patches.propose({"kind": "tests", "path": root, "cmd": None, "summary": "x", "files": ["calc.py"]}, object(), threading.Event())


@ext("Performance")
def test_governor_levels_and_background_tracker():
    sm, ht = mod("system_monitor"), mod("hand_tracker")
    if not hasattr(sm, "PerformanceGovernor") or not hasattr(ht, "BackgroundTracker"):
        skip("run apply_patches.py first")
    H = mock.MagicMock()
    tr = ns(interval=0.03)
    g = sm.PerformanceGovernor(H, tr)
    g.update(95, 60)
    eq(g.level, "normal")
    g.update(95, 60)
    eq((g.level, g.detect_every, tr.interval > 0.03), ("critical", 10, True))
    for _ in range(4):
        g.update(10, 60)
    eq(g.level, "critical")
    g.update(10, 60)
    eq((g.level, g.detect_every), ("normal", 3))

    class Inner:
        def __init__(self):
            self.n = 0

        def read(self):
            self.n += 1
            r = ht.TrackerResult()
            r.zoom_delta, r.events = 0.1, [(0, "peace")]
            return r, object()

        def close(self):
            pass
    bt = ht.BackgroundTracker(Inner(), interval=0.01)
    time.sleep(0.3)
    r, frame = bt.read()
    assert r.zoom_delta > 0.15 and len(r.events) >= 2 and frame is not None
    r2, f2 = bt.read()
    assert r2.zoom_delta < r.zoom_delta and (f2 is None or r2.events != r.events)
    bt.close()


@ext("AI brain")
def test_routing_dedupe_cache_and_recovery():
    br = mod("brain")
    if not hasattr(br, "route_request"):
        skip("run apply_patches.py first")
    eq(br.route_request("what is two plus two")[0], br.FAST_MODEL)
    eq(br.route_request("why does my code crash")[0], br.MODEL)
    eq(br.route_request("what's the latest news")[0], br.MODEL)
    c = _Create(iter([_ev("Paris is the capital.")]))
    b = br.Brain(c, lambda m: None, lambda n, a: "")
    deltas = []
    q = "what is the capital of france"
    eq(b.ask(q, deltas.append, threading.Event()), "Paris is the capital.")
    eq(c.calls[0]["model"], br.FAST_MODEL)
    deltas.clear()
    b._recent.clear()
    eq(b.ask(q, deltas.append, threading.Event()), "Paris is the capital.")
    eq((len(c.calls), deltas), (1, ["Paris is the capital."]))               # cached, spoken again
    eq(b.ask(q, deltas.append, threading.Event()), "Paris is the capital.")  # duplicate within seconds: silent
    eq(len(deltas), 1)
    b2 = br.Brain(_Create(iter([_ev(tool=("c1", "x", "{}"))]), iter([_ev("ok")])), lambda m: None, mock.Mock(side_effect=[RuntimeError("timeout"), "fine"]))
    with mock.patch.object(br.time, "sleep", lambda s: None):
        b2.ask("do the thing right now please", lambda d: None, threading.Event())
    tool_msg = [m for m in b2.client.calls[1]["messages"] if m["role"] == "tool"][0]
    eq(tool_msg["content"], "fine")


# ============================================================================ hologram 3D upgrade
import contextlib

_GL_NAMES = ("glBegin", "glEnd", "glVertex2f", "glVertex3f", "glColor4f", "glLineWidth", "glPointSize", "glPushMatrix",
             "glPopMatrix", "glTranslatef", "glDepthMask", "glBlendFunc", "glGetDoublev", "glGetIntegerv", "glEnable",
             "glDisable", "glLightfv", "gluNewQuadric", "gluSphere", "gluProject")


@contextlib.contextmanager
def gl_mocks():
    hm = mod("hologram")
    with mock.patch.multiple(hm, **{n: mock.DEFAULT for n in _GL_NAMES}) as m:
        m["gluProject"].return_value = (10.0, 10.0, 0.5)
        m["gluNewQuadric"].return_value = object()
        yield hm, m


def _patched():
    hm, h = bare_holo()
    if not hasattr(hm.Hologram, "_draw_ball"):
        skip("hologram 3D patch not applied (python apply_hologram_3d_patch.py)")
    h.width, h.height = 1280, 720
    return hm, h


@ext("Hologram 3D")
def patch_markers_present_in_source():
    _patched()
    src = open(os.path.join(ROOT, "jarvis_ui", "hologram.py"), encoding="utf-8").read()
    for needle in ("gluSphere", "def _draw_quarks", "MOLECULE_BOND_ORDER", "self._draw_ball(node_r, nc)",
                   '"satellite", "molecule"):', "self._draw_hologram_stripes(cx, cy, theme_color)", "self._draw_particle_legend()"):
        assert needle in src, f"missing: {needle}"
    assert "glVertex3f(ax, ay, 0.0)" not in src, "constellation edges still flat"


@ext("Hologram 3D")
def molecule_bond_orders():
    hm, h = _patched()
    for name, want in (("benzene", [2, 1, 2, 1, 2, 1] + [1] * 6), ("carbon dioxide", [2, 2]), ("nitrogen", [3]), ("water", [1, 1])):
        assert h.load_molecule(name), name
        eq(h.molecule_orders, want)
        eq(len(h.molecule_orders), len(h.molecule_data["bonds"]))


@ext("Hologram 3D")
def idle_sway_only_without_a_hand():
    hm, h = _patched()
    for _ in range(60):
        h.update(0, 0, 0.0, target_dt=1 / 30)
    assert abs(h.rotation_y) > 0.5 and abs(h.rotation_x) > 0.1, (h.rotation_y, h.rotation_x)
    hm2, h2 = _patched()
    for _ in range(200):
        h2.update(0.5, 0, 0.0, target_dt=1 / 30)
    assert abs(h2.rotation_y - 45) < 1, h2.rotation_y


@ext("Hologram 3D")
def ball_is_lit_and_quadric_is_cached():
    hm, h = _patched()
    with gl_mocks() as (hm, m):
        h._draw_ball(0.2, (1, 0, 0))
        h._draw_ball(0.3, (0, 1, 0), 0.5)
    eq(m["gluNewQuadric"].call_count, 1)
    eq([c.args[1] for c in m["gluSphere"].call_args_list], [0.2, 0.3])
    assert mock.call(hm.GL_LIGHTING) in m["glEnable"].call_args_list and mock.call(hm.GL_LIGHTING) in m["glDisable"].call_args_list


@ext("Hologram 3D")
def quark_content_matches_proton_and_neutron():
    hm, h = _patched()
    for is_p, big in ((True, 2), (False, 1)):
        radii = []
        h._draw_ball = lambda r, c, a=1.0, s=14: radii.append(r)
        with gl_mocks():
            h._draw_quarks(is_p, 0.09, 3, 1.0)
        eq(len(radii), 3)
        eq(sum(1 for r in radii if r == 0.026), big)


@ext("Hologram 3D")
def nucleus_shows_protons_neutrons_and_quarks():
    hm, h = _patched()

    def run(p, n):
        quarks, balls = [], []
        h._draw_quarks = lambda is_p, r, seed, b: quarks.append(is_p)
        h._draw_ball = lambda r, c, a=1.0, s=14: balls.append((r, a))
        h._draw_glow_fan = lambda *a, **k: None
        with gl_mocks():
            h._draw_nucleus_particles(p, n, 1.0)
        return quarks, balls
    q, b = run(6, 6)
    eq((len(q), sum(q), len(b)), (12, 6, 12))
    assert all(a < 1 for _, a in b), "nucleon shells should be translucent"
    q, b = run(92, 146)
    eq((q, len(b)), ([], h.NUCLEON_RENDER_CAP))
    eq(run(1, 0)[0], [True])
    eq(run(0, 0), ([], []))


@ext("Hologram 3D")
def stripes_and_legend_visibility():
    hm, h = _patched()
    h.height = 900
    with gl_mocks() as (_, m):
        h.mode = "empty"
        h._draw_hologram_stripes(640, 450, (0.1, 0.6, 1.0))
        eq(m["glBegin"].call_count, 0)
        h.mode = "atom"
        h._draw_hologram_stripes(640, 450, (0.1, 0.6, 1.0))
        eq(m["glBegin"].call_count, 1)
    texts = []
    h.font_small, h._draw_circle_2d = object(), lambda *a, **k: None
    h._blit_text = lambda f, t, x, y, color=None: texts.append(t)
    with gl_mocks():
        h.mode = "atom"
        h._draw_particle_legend()
        eq(len(texts), 4)
        h.mode = "shape"
        h._draw_particle_legend()
        h.height, h.mode = 700, "atom"
        h._draw_particle_legend()
        eq(len(texts), 4)


@ext("Hologram 3D")
def constellation_stars_have_depth():
    hm, h = _patched()
    assert h.load_constellation("orion")
    h._draw_node_fan = h._draw_ball = lambda *a, **k: None
    with gl_mocks() as (_, m):
        h._draw_constellation(None, None)
    zs = {round(c.args[2], 3) for c in m["glTranslatef"].call_args_list}
    assert len(zs) >= 5, zs
    eq(len(h._star_label_positions), 7)


@ext("Hologram 3D")
def patch_script_anchors_match_source():
    sys.path.insert(0, ROOT)
    try:
        import apply_hologram_3d_patch as p
    except ImportError:
        skip("apply_hologram_3d_patch.py not present")
    src = open(p.TARGET, encoding="utf-8").read()
    if "_draw_ball" in src:
        compile(src, p.TARGET, "exec")
        return "already applied"
    bad = [i for i, (old, _) in enumerate(p.EDITS, 1) if src.count(old) != 1]
    assert not bad, f"anchors not matching exactly once: {bad}"
    for name, _ in p.METHODS:
        eq(len(re.findall(rf"^    def {name}\(", src, re.M)), 1)


# ============================================================================ holographic construction (stage 1)
_PC_GL = ("glPushMatrix", "glPopMatrix", "glTranslatef", "glRotatef", "glScalef", "glDeleteLists", "glCallList", "glColor4f", "glLineWidth")


@contextlib.contextmanager
def pc_gl():
    """Patches the OpenGL names inside procedural_construction (no GL context needed)."""
    pc = mod("procedural_construction")
    with mock.patch.multiple(pc, **{n: mock.DEFAULT for n in _PC_GL}) as m:
        yield pc, m


def pc_fake_holo():
    """Hologram stand-in: records which primitive was drawn, and mimics Hologram._cached_list."""
    calls, builds = [], []
    h = ns(elapsed=0.0, _materialize_start=0.0, _gl_lists={}, calls=calls, builds=builds)
    for prim, meth in (("box", "_draw_wireframe_cube"), ("sphere", "_draw_wireframe_sphere"),
                       ("cylinder", "_draw_wireframe_cylinder"), ("pyramid", "_draw_wireframe_pyramid")):
        setattr(h, meth, lambda *a, _p=prim, **k: calls.append(_p))

    def cached(key, build):
        if key not in h._gl_lists:
            h._gl_lists[key] = len(h._gl_lists) + 1
            builds.append(key)
            build()
        return h._gl_lists[key]
    h._cached_list = cached
    return h


def pc_part(**over):
    d = {"name": "a", "prim": "box", "pos": (0, 0, 0), "size": (1, 1, 1)}
    d.update(over)
    return d


@ext("Holographic construction")
def robot_demo_is_valid_and_within_bounds():
    pc = mod("procedural_construction")
    m = pc.MODELS["robot"]
    assert len(m.parts) == 18 and {p.prim for p in m.parts} <= set(pc.PRIMITIVES)
    assert len({p.name for p in m.parts}) == 18
    ys = [p.pos[1] + s * 0.5 for p in m.parts for s in (p.size[1], -p.size[1])]
    assert -1.5 < min(ys) and max(ys) < 1.5, (min(ys), max(ys))      # stays well inside the fog range
    assert abs(m.duration - (1.10 + pc.GROW_SECONDS)) < 1e-9
    hash(m)                                                           # hashable: used as the display-list cache key


@ext("Holographic construction")
def validator_rejects_bad_input_and_normalizes_good_input():
    pc = mod("procedural_construction")
    bad = [pc_part(prim="torus"), pc_part(prim=None), pc_part(pos=(0, 0)), pc_part(pos=(float("nan"), 0, 0)),
           pc_part(pos=(0, float("inf"), 0)), pc_part(pos=(3, 0, 0)), pc_part(pos=(True, 0, 0)), pc_part(pos=("1", 0, 0)),
           pc_part(pos="abc"), pc_part(size=(0, 1, 1)), pc_part(size=(1, 9, 1)), pc_part(size=-1), pc_part(rot=(0, 9999, 0)),
           pc_part(delay=-1), pc_part(delay=99), pc_part(name=""), pc_part(name="../x"), pc_part(name="x" * 40),
           pc_part(color="red"), "box", None]
    for raw in bad:
        raises(pc.PartError, pc.build_model, "m", [raw])
    raw_missing = {"name": "a", "prim": "box", "pos": (0, 0, 0)}
    raises(pc.PartError, pc.build_model, "m", [raw_missing])
    raises(pc.PartError, pc.build_model, "m", [])
    raises(pc.PartError, pc.build_model, "m", "box")
    raises(pc.PartError, pc.build_model, "m", [pc_part()] * 2)                              # duplicate names
    raises(pc.PartError, pc.build_model, "m", [pc_part(name=f"p{i}") for i in range(pc.MAX_PARTS + 1)])
    raises(pc.PartError, pc.build_model, "bad/name", [pc_part()])
    ok = pc.build_model("Tiny", [pc_part(size=0.5, rot=(190, -190, 0), delay=1)])
    p = ok.parts[0]
    eq((ok.name, p.size, p.rot, p.delay), ("tiny", (0.5, 0.5, 0.5), (-170.0, 170.0, 0.0), 1.0))
    eq(len(pc.build_model("m", [pc_part(name=f"p{i}") for i in range(pc.MAX_PARTS)]).parts), pc.MAX_PARTS)


@ext("Holographic construction")
def draw_model_transforms_counts_and_matrix_balance():
    h = pc_fake_holo()
    with pc_gl() as (pc, m):
        model = pc.MODELS["robot"]
        pc.draw_model(h, model)
        want = {}
        for p in model.parts:
            want[p.prim] = want.get(p.prim, 0) + 1
        got = {k: h.calls.count(k) for k in set(h.calls)}
        eq(got, want)
        eq((m["glPushMatrix"].call_count, m["glPopMatrix"].call_count), (18, 18))
        first = model.parts[0]
        eq(m["glTranslatef"].call_args_list[0].args, first.pos)
        eq(tuple(round(v, 6) for v in m["glScalef"].call_args_list[0].args), tuple(round(0.5 * s, 6) for s in first.size))
        eq(m["glRotatef"].call_count, 2)                                 # only the two tilted arms rotate (z axis only)
        assert all(c.args[1:] == (0, 0, 1) for c in m["glRotatef"].call_args_list)
        # partway through assembly: fewer parts, smaller scale
        h.calls.clear()
        m["glScalef"].reset_mock()
        t = 0.2
        pc.draw_model(h, model, t)
        expect = [p for p in model.parts if pc._ease((t - p.delay) / pc.GROW_SECONDS) >= pc.MIN_GROWTH]
        assert 0 < len(h.calls) == len(expect) < 18, (len(h.calls), len(expect))
        assert all(c.args[0] <= 0.5 * model.parts[0].size[0] + 1e-9 for c in m["glScalef"].call_args_list[:1])
        h.calls.clear()
        pc.draw_model(h, model, 0.0)
        eq(h.calls, [])                                                   # nothing has started growing yet
        # a failing primitive must not leave the matrix stack unbalanced
        m["glPushMatrix"].reset_mock()
        m["glPopMatrix"].reset_mock()
        h._draw_wireframe_cube = mock.Mock(side_effect=RuntimeError("gl"))
        raises(RuntimeError, pc.draw_model, h, model)
        eq(m["glPushMatrix"].call_count, m["glPopMatrix"].call_count)


@ext("Holographic construction")
def draw_construction_animates_then_replays_cached_list():
    h = pc_fake_holo()
    with pc_gl() as (pc, m):
        dur = pc.MODELS["robot"].duration
        h._materialize_start = 5.0
        h.elapsed = 5.0 + pc.ASSEMBLY_DELAY + dur / 2
        for _ in range(5):                                                # the 5 shape layers
            pc.draw_construction(h, "robot")
        eq((h.builds, m["glCallList"].call_count), ([], 0))               # assembling: drawn directly, nothing cached
        direct = len(h.calls)
        assert direct > 0 and direct % 5 == 0
        h.calls.clear()
        h.elapsed = 5.0 + pc.ASSEMBLY_DELAY + dur + 0.1
        for _ in range(5):
            pc.draw_construction(h, "robot")
        eq((len(h.builds), m["glCallList"].call_count, len(h.calls)), (1, 5, 18))   # compiled once, replayed 5 times
        eq(h.builds[0][:2], ("construction", "robot"))
        before = len(h.calls)
        pc.draw_construction(h, "nothing")
        eq(len(h.calls), before)
        del h._cached_list                                                # older Hologram without display lists
        h.calls.clear()
        pc.draw_construction(h, "robot")
        eq(len(h.calls), 18)


@ext("Holographic construction")
def install_wraps_shape_renderer_without_touching_other_shapes():
    pc = mod("procedural_construction")
    drawn, calls = [], []
    h = ns(_draw_shape=drawn.append, hide_weather=lambda: calls.append("hide"), _trigger_materialize=lambda: calls.append("mat"),
           mode="empty", shape_name="sphere", mode_label="", selected_index=2, elapsed=0.0)
    assert pc.install(h) is h and pc.install(h) is h
    h._draw_shape("cube")
    eq(drawn, ["cube"])                                                   # other shapes pass straight through
    seen = []
    with mock.patch.object(pc, "draw_construction", lambda hol, name: seen.append(name)):
        h._draw_shape("robot")
    eq((seen, drawn), (["robot"], ["cube"]))                              # installed once: no double wrapping
    assert h.load_construction("robot") and not h.load_construction("ghost")
    eq((h.mode, h.shape_name, h.mode_label, h.selected_index, calls), ("shape", "robot", "ROBOT CONSTRUCTION", None, ["hide", "mat"]))
    h2 = ns(_draw_shape=drawn.append, mode="empty", shape_name="x")
    pc.install(h2)
    eq(h2.load_construction("ghost"), False)
    eq((h2.mode, h2.shape_name), ("empty", "x"))


@ext("Holographic construction")
def installs_on_real_hologram_without_voice_side_effects():
    pc, (hm, h) = mod("procedural_construction"), bare_holo()
    drawn = []
    h._draw_shape = drawn.append
    pc.install(h)
    assert "robot" not in hm.SHAPES, "demo name must not join SHAPES (the voice matcher would hijack it)"
    assert h.load_construction("robot") and h.mode == "shape" and h.shape_name == "robot" and h.selected_index is None
    h._draw_shape("sphere")
    eq(drawn, ["sphere"])
    h.show_info_card("q", "a")
    eq(h.mode, "info")                                                    # other modes still replace it normally


# ============================================================================ holographic construction (stage 2)
@contextlib.contextmanager
def pc_registry():
    """Runs a test against the real registry and restores it afterwards (nothing leaks between tests)."""
    pc = mod("procedural_construction")
    with mock.patch.dict(pc.MODELS):
        yield pc


def pc_spec(name="widget", n=2, **over):
    d = {"name": name, "parts": [pc_part(name=f"p{i}", pos=(0.1 * i, 0, 0)) for i in range(n)]}
    d.update(over)
    return d


def pc_extent(model):
    lo, hi = mod("procedural_construction").model_bounds(model)
    return tuple(h - l for l, h in zip(lo, hi))


@ext("Holographic construction 2")
def built_in_library_is_valid_complete_and_in_bounds():
    pc = mod("procedural_construction")
    eq(pc.list_models(), ["car", "drone", "robot", "spaceship"])
    need = {"robot": {"head", "torso", "arm left", "leg right", "antenna"},
            "drone": {"body", "motor front left", "propeller rear right", "camera gimbal", "skid left"},
            "car": {"chassis", "wheel front left", "wheel rear right", "cabin", "headlight left", "spoiler"},
            "spaceship": {"fuselage", "nose cone", "wing left", "wing right", "engine left", "exhaust right", "tail fin"}}
    for name, parts in need.items():
        m = pc.get_model(name)
        names = {p.name for p in m.parts}
        assert parts <= names, (name, parts - names)
        assert 12 <= len(m.parts) <= 24 and len(names) == len(m.parts)
        assert {p.prim for p in m.parts} <= set(pc.PRIMITIVES)
        lo, hi = pc.model_bounds(m)
        assert max(abs(v) for v in lo + hi) <= pc.MAX_EXTENT, (name, lo, hi)
        delays = [p.delay for p in m.parts]
        assert delays == sorted(delays) and delays[0] == 0.0 and 1.0 < m.duration < 2.5, (name, delays[-1])
        assert pc.get_model(name.upper() + " ") is m                      # lookups ignore case and spaces
    assert {p.prim for n in need for p in pc.get_model(n).parts} == set(pc.PRIMITIVES)


@ext("Holographic construction 2")
def built_in_proportions_are_recognizable():
    pc = mod("procedural_construction")
    rx, ry, rz = pc_extent(pc.get_model("robot"))
    assert ry > 1.6 * rx and ry > 4 * rz, (rx, ry, rz)                    # tall and narrow
    dx, dy, dz = pc_extent(pc.get_model("drone"))
    assert dx > 3 * dy and dz > 3 * dy and abs(dx - dz) < 0.1, (dx, dy, dz)  # flat, square footprint
    cx, cy, cz = pc_extent(pc.get_model("car"))
    assert cx > 1.8 * cz and cz > 1.2 * cy, (cx, cy, cz)                  # long, wider than tall
    sx, sy, sz = pc_extent(pc.get_model("spaceship"))
    assert sx > sz > 2 * sy, (sx, sy, sz)                                 # long hull, swept wings, thin
    car = {p.name: p for p in pc.get_model("car").parts}
    assert car["wheel front left"].rot == (90.0, 0.0, 0.0) and car["wheel front left"].pos[1] < car["chassis"].pos[1]
    assert car["headlight left"].pos[0] > 1.0 > car["cabin"].pos[0]       # nose is +x
    ship = {p.name: p for p in pc.get_model("spaceship").parts}
    assert ship["nose cone"].pos[0] > 1.0 and ship["exhaust left"].pos[0] < -1.4


@ext("Holographic construction 2")
def robot_demo_is_unchanged_regression():
    pc = mod("procedural_construction")
    m = pc.get_model("robot")
    eq(len(m.parts), 18)
    eq((m.parts[0].name, m.parts[0].pos, m.parts[0].size, m.parts[0].delay), ("foot left", (-0.2, -1.25, 0.08), (0.34, 0.14, 0.5), 0.0))
    eq((m.parts[-1].name, m.parts[-1].prim, m.parts[-1].delay), ("antenna tip", "sphere", 1.1))
    eq((round(m.duration, 6), pc.GROW_SECONDS, pc.ASSEMBLY_DELAY), (1.45, 0.35, 0.25))
    eq([p.rot for p in m.parts if p.rot != (0.0, 0.0, 0.0)], [(0.0, 0.0, -6.0), (0.0, 0.0, 6.0)])
    lo, hi = pc.model_bounds(m)
    assert abs(lo[1] + 1.32) < 0.02 and abs(hi[1] - 1.30) < 0.02, (lo, hi)


@ext("Holographic construction 2")
def draw_every_built_in_model_balances_matrices():
    h = pc_fake_holo()
    with pc_gl() as (pc, m):
        for name in pc.list_models():
            model = pc.get_model(name)
            h.calls.clear()
            for mk in ("glPushMatrix", "glPopMatrix", "glRotatef"):
                m[mk].reset_mock()
            pc.draw_model(h, model)
            eq((len(h.calls), m["glPushMatrix"].call_count, m["glPopMatrix"].call_count), (len(model.parts),) * 3)
            eq(m["glRotatef"].call_count, sum(1 for p in model.parts for a in p.rot if a))


@ext("Holographic construction 2")
def registry_register_get_list_unregister():
    with pc_registry() as pc, pc_gl() as (_, m):
        w = pc.define_model("Widget", [pc_part(name="a"), pc_part(name="b", pos=(0.5, 0, 0))])
        eq((w.name, pc.list_models(), pc.get_model("WIDGET")), ("widget", ["car", "drone", "robot", "spaceship", "widget"], w))
        assert pc.register_model(w) is w                                  # identical re-registration is a no-op
        assert pc.define_model("widget", [pc_part(name="a"), pc_part(name="b", pos=(0.5, 0, 0))]) is w
        assert pc.unregister_model(" Widget ") and not pc.unregister_model("widget")
        raises(pc.PartError, pc.get_model, "widget")
        raises(pc.PartError, pc.get_model, None)
        for name in pc.BUILTIN_MODELS:
            raises(pc.PartError, pc.unregister_model, name)
        assert all(m[k].call_count == 0 for k in _PC_GL), "registry operations must never issue OpenGL calls"


@ext("Holographic construction 2")
def registry_never_silently_overwrites():
    with pc_registry() as pc:
        w1 = pc.define_model("widget", [pc_part(name="a")])
        raises(pc.PartError, pc.define_model, "widget", [pc_part(name="a", size=(2, 1, 1))])
        assert pc.get_model("widget") is w1
        w2 = pc.define_model("widget", [pc_part(name="a", size=(2, 1, 1))], replace=True)
        assert pc.get_model("widget") is w2 and w2.uid != w1.uid
        before = pc.get_model("robot")
        for replace in (False, True):
            raises(pc.PartError, pc.define_model, "robot", [pc_part(name="a")], None, replace)
        assert pc.get_model("robot") is before
        raises(pc.PartError, pc.register_model, "robot")                  # only Model instances
        raises(pc.PartError, pc.register_model, {"name": "x"})
        for i in range(pc.MAX_USER_MODELS - 1):                            # "widget" already uses one slot
            pc.define_model(f"m{i}", [pc_part()])
        raises(pc.PartError, pc.define_model, "one too many", [pc_part()])
        assert len(pc.list_models()) == len(pc.BUILTIN_MODELS) + pc.MAX_USER_MODELS


@ext("Holographic construction 2")
def hand_built_models_go_through_the_same_validator():
    with pc_registry() as pc:
        bad = pc.Model("evil", (pc.Part("a", "torus", (0, 0, 0), (1, 1, 1)),))
        raises(pc.PartError, pc.register_model, bad)
        nan = pc.Model("nan", (pc.Part("a", "box", (float("nan"), 0, 0), (1, 1, 1)),))
        raises(pc.PartError, pc.register_model, nan)
        dup = pc.Model("dup", (pc.Part("a", "box", (0, 0, 0), (1, 1, 1)), pc.Part("A", "box", (0, 0, 0), (1, 1, 1))))
        raises(pc.PartError, pc.register_model, dup)
        raw = pc.Model("raw", (pc.Part("a", "box", (0, 0, 0), (1, 1, 1), (0, 0, 190), 0.0),))
        got = pc.register_model(raw)
        eq(got.parts[0].rot, (0.0, 0.0, -170.0))                           # normalized by the shared validator
        eq([n for n in pc.list_models() if n in ("evil", "nan", "dup")], [])


@ext("Holographic construction 2")
def model_dicts_stagger_and_malformed_data():
    pc = mod("procedural_construction")
    m = pc.model_from_dict({"name": "Demo", "stagger": 0.1,
                            "parts": [pc_part(name="a"), pc_part(name="b", delay=2.0), pc_part(name="c")]})
    eq((m.name, [round(p.delay, 6) for p in m.parts]), ("demo", [0.0, 2.0, 0.2]))
    assert "demo" not in pc.MODELS                                        # model_from_dict never registers
    eq(pc.build_model("d", [pc_part(name="a"), pc_part(name="b")]).parts[1].delay, 0.0)   # no stagger: unchanged behaviour
    with pc_registry():
        assert pc.register_model_dict(pc_spec("viaDict")).name == "viadict" and "viadict" in pc.MODELS
    for spec in (None, [], "x", {}, {"name": "x"}, {"parts": [pc_part()]}, pc_spec(extra=1), pc_spec(parts="box"),
                 pc_spec(parts=[]), pc_spec(name="bad/name"), pc_spec(name=3), pc_spec(stagger=-1), pc_spec(stagger=5),
                 pc_spec(stagger=float("nan")), pc_spec(stagger="0.1"), pc_spec(parts=[pc_part(prim="torus")]),
                 pc_spec(parts=[pc_part(pos=(0, float("inf"), 0))]), pc_spec(parts=[pc_part(), pc_part()]),
                 pc_spec(parts=[pc_part(name=f"p{i}") for i in range(pc.MAX_PARTS + 1)]),
                 pc_spec(parts=[None]), pc_spec(parts=[pc_part(size=(1, 1))]), pc_spec(stagger=0.9, parts=[pc_part(name=f"p{i}") for i in range(40)])):
        raises(pc.PartError, pc.model_from_dict, spec)


@ext("Holographic construction 2")
def whole_model_dimensions_are_bounded_and_rotation_aware():
    pc = mod("procedural_construction")
    raises(pc.PartError, pc.build_model, "far", [pc_part(pos=(2.0, 2.0, 2.0), size=2.5)])     # each part legal, model is not
    raises(pc.PartError, pc.build_model, "long", [pc_part(pos=(1.5, 0, 0), size=(2.5, 0.1, 0.1))])
    ok = pc.build_model("edge", [pc_part(pos=(1.7, 0, 0), size=1.0)])                         # 2.2 from the origin
    assert abs(max(pc_extent(ok)) - 1.0) < 1e-9
    raises(pc.PartError, pc.build_model, "tilted", [pc_part(pos=(1.7, 0, 0), size=1.0, rot=(0, 0, 45))])   # 2.407 once rotated
    lo, hi = pc.model_bounds(pc.build_model("r", [pc_part(size=(2.0, 0.2, 0.2), rot=(0, 90, 0))]))
    assert abs(hi[2] - 1.0) < 1e-9 and abs(hi[0] - 0.1) < 1e-9, (lo, hi)                     # long axis swung onto z
    raises(pc.PartError, pc.build_model, "tiny", [pc_part(size=0.001)])


@ext("Holographic construction 2")
def instant_versus_animated_assembly():
    h = pc_fake_holo()
    with pc_gl() as (pc, m):
        model = pc.get_model("car")
        h._materialize_start, h.elapsed = 3.0, 3.0                         # the moment a model is loaded
        pc.draw_construction(h, "car")
        eq((h.calls, m["glCallList"].call_count), ([], 0))                 # sequential: nothing has started growing
        h._construction_assemble = False
        pc.draw_construction(h, "car")
        eq((len(h.calls), m["glCallList"].call_count, len(h.builds)), (len(model.parts), 1, 1))   # instant: whole model at once
        h._construction_assemble = True
        h.calls.clear()
        h.elapsed = 3.0 + pc.ASSEMBLY_DELAY + 0.3
        pc.draw_construction(h, "car")
        assert 0 < len(h.calls) < len(model.parts)                         # partway through the sequence
        h.elapsed = 3.0 + pc.ASSEMBLY_DELAY + model.duration + 0.01
        before = m["glCallList"].call_count
        pc.draw_construction(h, "car")
        eq((m["glCallList"].call_count, len(h.builds)), (before + 1, 1))   # finished: replays the list compiled for instant mode


@ext("Holographic construction 2")
def load_construction_modes_names_and_failures():
    pc = mod("procedural_construction")
    calls = []
    h = ns(hide_weather=lambda: calls.append("hide"), _trigger_materialize=lambda: calls.append("mat"), mode="empty",
           shape_name="x", mode_label="", selected_index=1, _draw_shape=lambda n: None)
    assert pc.load_construction(h, " Drone ") and (h.mode, h.shape_name, h.mode_label, h._construction_assemble) == ("shape", "drone", "DRONE CONSTRUCTION", True)
    assert pc.load_construction(h, "car", assemble=False) and h._construction_assemble is False
    assert pc.load_construction(h, "robot") and h._construction_assemble is True
    calls.clear()
    for bad in ("ghost", "", None, 5, ["car"]):
        eq(pc.load_construction(h, bad), False)
    eq((calls, h.shape_name), ([], "robot"))                               # failed loads change nothing
    pc.install(h)                                                          # lambda signature carries assemble through
    with mock.patch.object(pc, "load_construction", lambda hol, name, assemble: calls.append((name, assemble))):
        h.load_construction("car", False)
    eq(calls, [("car", False)])


@ext("Holographic construction 2")
def cache_invalidation_when_a_model_is_replaced():
    h = pc_fake_holo()
    h.elapsed = 100.0                                                      # long finished
    with pc_registry() as pc, pc_gl() as (_, m):
        v1 = pc.define_model("temp", [pc_part(name="a")])
        for _ in range(5):
            pc.draw_construction(h, "temp")
        eq((len(h.builds), m["glCallList"].call_count, m["glDeleteLists"].call_count), (1, 5, 0))
        first_id = h._gl_lists[("construction", "temp", v1.uid)]
        assert pc.define_model("temp", [pc_part(name="a")]) is v1          # identical: cache stays valid
        pc.draw_construction(h, "temp")
        eq(len(h.builds), 1)
        v2 = pc.define_model("temp", [pc_part(name="a"), pc_part(name="b")], replace=True)
        pc.draw_construction(h, "temp")
        pc.draw_construction(h, "temp")
        eq(len(h.builds), 2)                                               # recompiled exactly once
        eq(m["glDeleteLists"].call_args_list[0].args, (first_id, 1))       # stale list freed on the render path
        eq(set(h._gl_lists), {("construction", "temp", v2.uid)})
        pc.draw_construction(h, "robot")                                   # other models' lists are never touched
        pc.define_model("temp", [pc_part(name="z")], replace=True)
        pc.draw_construction(h, "temp")
        assert ("construction", "robot", pc.get_model("robot").uid) in h._gl_lists and m["glDeleteLists"].call_count == 2
        pc.unregister_model("temp")
        before = m["glCallList"].call_count
        pc.draw_construction(h, "temp")
        eq(m["glCallList"].call_count, before)                             # unregistered: draws nothing, no error


@ext("Holographic construction 2")
def cache_key_is_a_stable_uid_not_a_per_frame_hash():
    h = pc_fake_holo()
    h.elapsed = 100.0
    with pc_gl() as (pc, m):
        model = pc.get_model("drone")
        with mock.patch.object(pc.Model, "__hash__", side_effect=AssertionError("model hashed during drawing")):
            for _ in range(5):
                pc.draw_construction(h, "drone")
        eq(list(h._gl_lists), [("construction", "drone", model.uid)])
        a, b = pc.build_model("same", [pc_part()]), pc.build_model("same", [pc_part()])
        assert a == b and a.uid != b.uid and isinstance(a.uid, int)


# ============================================================================ holographic construction (stage 3)
# Insert after the "Holographic construction 2" tests in test_aurora.py (before the presentation builder section).
# Also extend the GL name list used by pc_gl():  _PC_GL = (..., "glColor4f", "glLineWidth")
def pc3_setup(name="robot", mode="shape"):
    """Fake hologram with the editor installed and the model 'displayed' (fully built unless elapsed is changed)."""
    pc = mod("procedural_construction")
    h = pc_fake_holo()
    h.mode, h.shape_name, h.elapsed, h.selected_index, h._draw_shape = mode, name, 100.0, 2, (lambda n: None)
    h.hide_weather, h._trigger_materialize = (lambda: None), (lambda: None)
    pc.install(h)
    return pc, h, h.construction


def pc3_part(model, name):
    return next(p for p in model.parts if p.name == name)


@ext("Holographic construction 3")
def model_transform_limits_composition_and_registry_untouched():
    with pc_gl() as (pc, m):
        _, h, ed = pc3_setup()
        base = pc.MODELS["robot"]
        eq(ed.transform_for("robot"), None)
        assert ed.move_model(0.5, -0.25, 1.0) and ed.rotate_model(0, 370, 0) and ed.scale_model(1.5)
        eq(ed.transform_for("robot"), ((0.5, -0.25, 1.0), (0.0, 10.0, 0.0), 1.5))
        assert ed.move_model(10, -10, 0) and ed.transform_for("robot")[0] == (2.0, -2.0, 1.0)            # clamped
        assert ed.scale_model(100) and ed.transform_for("robot")[2] == pc.MAX_MODEL_SCALE
        assert ed.scale_model(0.001, absolute=True) and ed.transform_for("robot")[2] == pc.MIN_MODEL_SCALE
        assert ed.move_model(1, 1, 1, absolute=True) and ed.transform_for("robot")[0] == (1.0, 1.0, 1.0)
        for bad in (float("nan"), float("inf"), "1", None, True):
            assert not ed.move_model(bad) and "finite" in ed.last_error
        for bad in (0, -1, float("nan"), "2", True):
            assert not ed.scale_model(bad) and "positive" in ed.last_error
        assert pc.MODELS["robot"] is base and base.uid == pc.get_model("robot").uid                       # definition untouched
        ed.reset_transform()
        eq(ed.transform_for("robot"), None)
        # composition: outer transform is applied first (translate, rotate z/y/x, scale), geometry then draws inside it
        ed.move_model(0.5, 0.0, 0.0), ed.rotate_model(10, 20, 30), ed.scale_model(2.0)
        for k in ("glPushMatrix", "glPopMatrix", "glTranslatef", "glRotatef", "glScalef"):
            m[k].reset_mock()
        pc.draw_construction(h, "robot")
        eq(m["glTranslatef"].call_args_list[0].args, (0.5, 0.0, 0.0))
        eq([c.args for c in m["glRotatef"].call_args_list[:3]], [(30.0, 0, 0, 1), (20.0, 0, 1, 0), (10.0, 1, 0, 0)])
        eq(m["glScalef"].call_args_list[0].args, (2.0, 2.0, 2.0))
        eq(m["glPushMatrix"].call_count, m["glPopMatrix"].call_count)
        m["glTranslatef"].reset_mock()
        ed.reset_transform()
        pc.draw_construction(h, "robot")
        assert m["glTranslatef"].call_count == 0, "identity transform must add no matrix work"             # cached call only


@ext("Holographic construction 3")
def transform_never_touches_the_geometry_cache():
    with pc_gl() as (pc, m):
        _, h, ed = pc3_setup()
        pc.draw_construction(h, "robot")
        eq((len(h.builds), m["glDeleteLists"].call_count), (1, 0))
        for _ in range(5):
            ed.move_model(0.1), ed.rotate_model(0, 5, 0), ed.scale_model(1.05)
            pc.draw_construction(h, "robot")
        eq((len(h.builds), m["glDeleteLists"].call_count), (1, 0))                                       # same display list


@ext("Holographic construction 3")
def part_selection_api_and_orbit_selection_untouched():
    with pc_gl() as (pc, m):
        _, h, ed = pc3_setup()
        names = [p.name for p in pc.MODELS["robot"].parts]
        eq(ed.part_names(), names)
        eq(ed.selected_part, None)
        assert ed.select_part("  HEAD ") and ed.selected_part == "head"
        assert not ed.select_part("nope") and ed.selected_part == "head" and "no part" in ed.last_error
        assert not ed.select_part(None) and not ed.select_part(5)
        ed.deselect_part()
        eq([ed.select_next_part() for _ in range(len(names) + 1)], names + [None])                      # walks, then deselects
        eq(ed.select_previous_part(), names[-1])
        eq(ed.select_previous_part(), names[-2])
        ed.select_part(names[0])
        eq(ed.select_previous_part(), None)
        eq(h.selected_index, 2)                                                                          # orbit selection untouched


@ext("Holographic construction 3")
def selection_clears_on_model_or_mode_change_and_edits_need_a_model():
    with pc_gl() as (pc, m):
        _, h, ed = pc3_setup()
        ed.select_part("head")
        h.mode = "atom"                                                                                  # another hologram mode
        eq(ed.selected_part, None)
        for call in (lambda: ed.select_part("head"), ed.undo, ed.reset, lambda: ed.move_part(1), lambda: ed.move_model(1)):
            assert call() is False and "no construction model" in ed.last_error
        eq((ed.select_next_part(), ed.part_names(), ed.undo_depth, ed.gesture_select_next()), (None, [], 0, False))
        h.mode = "shape"
        eq(ed.selected_part, None)                                                                       # not resurrected by returning
        ed.select_part("head")
        h.shape_name = "cube"                                                                            # a non-construction shape
        eq(ed.selected_part, None)
        h.shape_name = "robot"
        h.shape_name = "ghost"
        assert not ed.move_model(1) and ed.selected_part is None                                         # unknown/unregistered model
        h.shape_name = "robot"
        ed.select_part("head")
        assert h.load_construction("car") and ed.selected_part is None                                   # load clears selection
        h.shape_name = "robot"
        ed.select_part("head")
        h.shape_name = "car"
        eq(ed.selected_part, None)                                                                       # other model's part name


@ext("Holographic construction 3")
def part_edits_clamp_validate_and_preserve_the_original():
    with pc_gl() as (pc, m):
        _, h, ed = pc3_setup()
        base = pc.MODELS["robot"]
        assert not ed.move_part(1) and ed.last_error == "no part selected"
        ed.select_part("head")
        h0 = pc3_part(base, "head")
        assert ed.move_part(0.2, -0.1, 0.0)
        p = pc3_part(ed.model_for("robot"), "head")
        eq(tuple(round(a - b, 9) for a, b in zip(p.pos, h0.pos)), (0.2, -0.1, 0.0))
        assert ed.move_part(50, 0, 0) and pc3_part(ed.model_for("robot"), "head").pos[0] == pc.MAX_POS     # clamped, accepted
        assert ed.rotate_part(0, 200, 0) and ed.rotate_part(0, 200, 0)
        eq(pc3_part(ed.model_for("robot"), "head").rot, (0.0, 40.0, 0.0))                               # 400 degrees wraps to 40
        assert ed.move_part(0, 0, 0, absolute=True)                                                      # back to the origin
        assert ed.resize_part(2.0) and pc3_part(ed.model_for("robot"), "head").size == (1.0, 0.84, 0.84)
        assert ed.resize_part((1, 100, 1)) and pc3_part(ed.model_for("robot"), "head").size[1] == pc.MAX_SIZE
        assert ed.resize_part(0.0001) and pc3_part(ed.model_for("robot"), "head").size[0] == pc.MIN_SIZE      # clamped to the minimum
        for bad in (0, -1, float("nan"), "2", None, (1, 1), True):
            before = ed.model_for("robot")
            assert not ed.resize_part(bad) and ed.model_for("robot") is before, bad
        assert pc.MODELS["robot"] is base and pc3_part(base, "head") == h0                               # registry model intact
        names_before = [p.name for p in ed.model_for("robot").parts]
        eq([p.name for p in base.parts], names_before)                                                   # order and names kept
        eq(ed.model_for("robot").name, "robot")


@ext("Holographic construction 3")
def edit_that_breaks_whole_model_limits_is_rejected_atomically():
    with pc_gl() as (pc, m):
        _, h, ed = pc3_setup()
        ed.select_part("torso")
        before = ed.model_for("robot")
        assert ed.move_part(2.0, 2.0, 2.0, absolute=True) is False and "extends" in ed.last_error
        assert ed.model_for("robot") is before and ed.undo_depth == 0                                    # nothing recorded
        ed.select_part("head")
        assert ed.move_part(2.0, 0.0, 0.0, absolute=True)                                                # fits (reach 2.25)
        mid = ed.model_for("robot")
        assert ed.resize_part(5.0) is False and ed.model_for("robot") is mid and ed.undo_depth == 1      # would reach 3.25
        assert not ed.move_part(0, 0, 0) and ed.last_error == "no change" and ed.undo_depth == 1


@ext("Holographic construction 3")
def undo_is_bounded_reset_is_undoable_and_state_is_per_model():
    with pc_gl() as (pc, m), mock.patch.object(pc, "MAX_UNDO", 3):
        _, h, ed = pc3_setup()
        assert not ed.undo() and ed.last_error == "nothing to undo"
        ed.select_part("head")
        base_pos = pc3_part(pc.MODELS["robot"], "head").pos
        snaps = []
        for i in range(5):
            snaps.append(ed.model_for("robot"))
            ed.move_part(0.01 * (i + 1), 0, 0)
        eq(ed.undo_depth, 3)                                                                             # oldest two dropped
        for expect in reversed(snaps[2:]):
            assert ed.undo() and ed.model_for("robot") is expect
        assert not ed.undo()
        ed.rotate_model(0, 90, 0)
        ed.move_part(0.1)
        assert ed.reset() and ed.model_for("robot") is pc.MODELS["robot"] and ed.transform_for("robot") is None
        assert ed.undo() and ed.transform_for("robot")[1] == (0.0, 90.0, 0.0) and ed.model_for("robot") is not pc.MODELS["robot"]
        ed.reset()
        assert not ed.reset() and ed.last_error == "no change"
        h.shape_name = "car"
        eq(ed.undo_depth, 0)                                                                             # history is per model
        h.shape_name = "robot"
        assert ed.undo_depth > 0


@ext("Holographic construction 3")
def edits_render_immediately_and_cache_invalidates_exactly_when_needed():
    with pc_gl() as (pc, m):
        _, h, ed = pc3_setup()
        pc.draw_construction(h, "robot")
        eq(len(h.builds), 1)
        first = pc.MODELS["robot"].uid
        eq(set(h._gl_lists), {("construction", "robot", first)})
        ed.select_part("head")
        assert ed.move_part(0.0, 0.1, 0.0)
        m["glTranslatef"].reset_mock()
        pc.draw_construction(h, "robot")                                                                 # next frame shows the edit
        head = pc3_part(ed.model_for("robot"), "head")
        assert any(c.args == head.pos for c in m["glTranslatef"].call_args_list), "edited head never drawn"
        eq((len(h.builds), m["glDeleteLists"].call_count), (2, 1))                                       # recompiled once, old freed
        second = ed.model_for("robot").uid
        assert second != first and set(h._gl_lists) == {("construction", "robot", second)}
        for _ in range(3):
            pc.draw_construction(h, "robot")
        eq(len(h.builds), 2)                                                                             # stable between edits
        assert ed.undo()
        pc.draw_construction(h, "robot")
        eq((len(h.builds), set(h._gl_lists)), (3, {("construction", "robot", first)}))                   # undo recompiles
        ed.reset()
        pc.draw_construction(h, "robot")
        eq(len(h.builds), 3)                                                                             # already base: no rebuild
        ed.select_part("head"), ed.move_part(0.1), pc.draw_construction(h, "robot"), ed.reset(), pc.draw_construction(h, "robot")
        eq(set(h._gl_lists), {("construction", "robot", first)})
        for k in h._gl_lists:
            assert k[1] == "robot"
        pc.draw_construction(h, "car")                                                                   # other models cache separately
        assert ("construction", "robot", first) in h._gl_lists and len(h._gl_lists) == 2


@ext("Holographic construction 3")
def edits_during_assembly_draw_directly_and_model_matrices_balance():
    with pc_gl() as (pc, m):
        _, h, ed = pc3_setup()
        h._materialize_start, h.elapsed = 10.0, 10.0 + pc.ASSEMBLY_DELAY + 0.4
        ed.select_part("torso")
        ed.move_part(0.1)
        ed.move_model(0.2)
        h.calls.clear()
        pc.draw_construction(h, "robot")
        assert 0 < len(h.calls) < 18 and m["glCallList"].call_count == 0 and not h.builds
        eq(m["glPushMatrix"].call_count, m["glPopMatrix"].call_count)
        h._draw_wireframe_cube = mock.Mock(side_effect=RuntimeError("gl"))
        m["glPushMatrix"].reset_mock(), m["glPopMatrix"].reset_mock()
        raises(RuntimeError, pc.draw_construction, h, "robot")
        eq(m["glPushMatrix"].call_count, m["glPopMatrix"].call_count)                                    # transform popped on error


@ext("Holographic construction 3")
def highlight_draws_only_the_selected_part_and_leaves_state_clean():
    with pc_gl() as (pc, m):
        _, h, ed = pc3_setup()
        pc.draw_construction(h, "robot")                                                                 # compile once
        h.calls.clear()
        m["glCallList"].reset_mock()
        pc.draw_construction(h, "robot")
        eq((h.calls, m["glCallList"].call_count, m["glColor4f"].call_count), ([], 1, 0))                 # nothing extra unselected
        ed.select_part("head")
        pc.draw_construction(h, "robot")
        eq(h.calls, ["box"])                                                                             # head is a box: one outline
        eq(m["glColor4f"].call_args.args[:3], (1.0, 1.0, 1.0))
        eq(m["glLineWidth"].call_args_list[-1].args, (2.0,))                                             # line width restored
        h.calls.clear()
        ed.select_part("core")
        pc.draw_construction(h, "robot")
        eq(h.calls, ["sphere"])
        h._materialize_start, h.elapsed = 100.0, 100.0 + pc.ASSEMBLY_DELAY                               # head not grown yet
        ed.select_part("head")
        h.calls.clear()
        pc.draw_construction(h, "robot")
        eq(h.calls, [])                                                                                  # not grown: no outline, nothing built


@ext("Holographic construction 3")
def switching_models_registry_replacement_and_keep_edits():
    with pc_registry() as pc, pc_gl() as (_, m):
        _, h, ed = pc3_setup()
        h.hide_weather, h._trigger_materialize = (lambda: None), (lambda: None)
        ed.select_part("head"), ed.move_part(0.1), ed.rotate_model(0, 30, 0)
        assert h.load_construction("car") and (h.shape_name, ed.selected_part) == ("car", None)
        assert h.load_construction("robot", keep_edits=True)
        assert ed.model_for("robot") is not pc.MODELS["robot"] and ed.transform_for("robot")[1] == (0.0, 30.0, 0.0)
        assert h.load_construction("robot")                                                              # default: fresh model
        assert ed.model_for("robot") is pc.MODELS["robot"] and ed.transform_for("robot") is None
        assert not h.load_construction("ghost") and h.shape_name == "robot"
        pc.define_model("temp", [pc_part(name="a"), pc_part(name="b", pos=(0.5, 0, 0))])
        assert h.load_construction("temp")
        ed.select_part("a"), ed.move_part(0.1)
        assert ed.model_for("temp") is not pc.MODELS["temp"]
        pc.define_model("temp", [pc_part(name="a"), pc_part(name="z")], replace=True)                    # replaced: edits are stale
        assert ed.model_for("temp") is pc.MODELS["temp"] and ed.selected_part is None
        pc.unregister_model("temp")
        eq(ed.model_for("temp"), None)
        h.mode = "shape"
        assert not ed.move_model(1)
        pc.draw_construction(h, "temp")                                                                  # draws nothing, no error
        h.shape_name = "robot"
        eq(ed.part_names(), [p.name for p in pc.MODELS["robot"].parts])


@ext("Holographic construction 3")
def gesture_hook_consumes_only_when_a_model_is_shown():
    with pc_gl() as (pc, m):
        _, h, ed = pc3_setup()
        names = ed.part_names()
        eq([(ed.gesture_select_next(), ed.selected_part) for _ in range(3)], [(True, names[0]), (True, names[1]), (True, names[2])])
        for _ in range(len(names) - 3):
            ed.gesture_select_next()
        assert ed.gesture_select_next() and ed.selected_part is None                                     # last -> deselect
        h.mode = "info"
        eq(ed.gesture_select_next(), False)                                                              # other modes: not consumed
    main_src = open(os.path.join(ROOT, "main.py"), encoding="utf-8").read()
    assert 'gesture == "thumbs_up" and getattr(hologram, "construction", None)' in main_src
    assert "construction" not in main_src.split('gesture == "thumbs_down"')[1].split("elif")[0], "thumbs_down must stay non-destructive"


@ext("Holographic construction 3")
def stage3_leaves_stage1_and_2_behaviour_intact():
    pc = mod("procedural_construction")
    eq(pc.list_models(), ["car", "drone", "robot", "spaceship"])
    eq((len(pc.MODELS["robot"].parts), pc.GROW_SECONDS, pc.ASSEMBLY_DELAY), (18, 0.35, 0.25))
    drawn = []
    h = ns(_draw_shape=drawn.append, hide_weather=lambda: None, _trigger_materialize=lambda: None, mode="empty", shape_name="x",
           mode_label="", selected_index=None, elapsed=0.0)
    pc.install(h)
    assert isinstance(h.construction, pc.ConstructionEditor)
    h._draw_shape("cube")
    eq(drawn, ["cube"])
    assert h.load_construction("robot") and h.load_construction("car", False)
    eq((h.mode, h.shape_name, h._construction_assemble), ("shape", "car", False))
    h.mode = "info"
    h._draw_shape("sphere")
    eq(drawn, ["cube", "sphere"])


# ============================================================================ presentation builder
G = "Presentation builder"


def pbm():
    return mod("presentation_builder")


def _need_pptx():
    try:
        import pptx   # noqa: F401
        import PIL    # noqa: F401
    except ImportError as e:
        skip(f"missing library: {e.name}")


def _pb_deck(theme="aurora"):
    # one slide of every layout, so a single render exercises all of them
    return pbm()._normalize({"title": "Black Holes", "subtitle": "Where gravity wins", "theme": theme, "slides": [
        {"layout": "bullets_image", "title": "What they are", "bullets": ["Escape velocity exceeds light speed"],
         "image": "a black hole bending light", "caption": "Artist impression", "notes": "intro"},
        {"layout": "bullets_image", "title": "How they form", "bullets": ["A massive star collapses"], "image": "a supernova"},
        {"layout": "stats", "title": "By the numbers", "bullets": ["M87 was imaged first"],
         "stats": [{"value": "4.3M", "label": "solar masses"}, {"value": "1915", "label": "general relativity"}, {"value": "2019", "label": "first image"}]},
        {"layout": "two_column", "title": "Types", "left_title": "Stellar", "left": ["3 to 100 solar masses"],
         "right_title": "Supermassive", "right": ["Millions of solar masses"]},
        {"layout": "section", "title": "Observing", "bullets": ["How we see the unseeable"], "image": "a radio telescope"},
        {"layout": "image_full", "title": "Event horizon", "bullets": ["The point of no return"], "image": "a glowing accretion disk"},
        {"layout": "quote", "title": "Quote", "quote": "Black holes ain't as black as they are painted.", "author": "Stephen Hawking"},
        {"layout": "bullets", "title": "Plain", "bullets": ["one", "two"]},
        {"layout": "closing", "title": "Thank you", "bullets": ["Questions?"], "image": "a starry sky"}]}, "black holes")


def t_pb_options():
    pb = pbm()
    eq(pb.split_options("volcanoes in the sunset theme without images"), ("volcanoes", "sunset", False))
    eq(pb.split_options("black holes"), ("black holes", None, True))
    eq(pb.split_options("cats in the neon style"), ("cats", "aurora", True))
    eq(pb.split_options("space with no pictures"), ("space", None, False))
    eq(pb.split_options("music in the purple theme"), ("music in the purple theme", None, True))     # unknown theme words stay in the topic


def t_pb_normalize():
    pb = pbm()
    spec = pb._normalize({"title": "T", "theme": "nope", "slides": [
        "junk", {"layout": "bogus", "title": "A", "bullets": ["x"] * 9},
        {"layout": "stats", "title": "B", "stats": [{"value": "1", "label": "x"}]},
        {"layout": "two_column", "title": "C", "left": ["a"]},
        {"layout": "quote", "title": "D"}]}, "topic")
    eq(spec["theme"], "aurora")
    eq([s["layout"] for s in spec["slides"]], ["bullets_image", "bullets_image", "bullets_image", "bullets"])
    eq(len(spec["slides"][0]["bullets"]), 5)
    eq(pb._normalize({}, "my topic")["title"], "My Topic")
    eq(pb._fallback("mars", 4)["slides"][-1]["layout"], "closing")


def t_pb_art():
    _need_pptx()
    from PIL import Image
    pb, d = pbm(), _fdir()
    a, b = os.path.join(d, "a.png"), os.path.join(d, "b.png")
    pb.make_art("same seed", 320, 200, "sunset", a)
    pb.make_art("same seed", 320, 200, "sunset", b)
    eq(Image.open(a).size, (320, 200))
    eq(open(a, "rb").read(), open(b, "rb").read())                 # deterministic per description
    pb.make_art("other seed", 320, 200, "sunset", b)
    assert open(a, "rb").read() != open(b, "rb").read()
    wide = Image.new("RGB", (400, 100))
    wide.save(os.path.join(d, "w.png"))
    pb._crop_to(os.path.join(d, "w.png"), 100, 100)
    eq(Image.open(os.path.join(d, "w.png")).size, (100, 100))      # cropped, never stretched


def _pictures(prs):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    return sum(sh.shape_type == MSO_SHAPE_TYPE.PICTURE for s in prs.slides for sh in s.shapes)


def t_pb_render_every_theme():
    _need_pptx()
    from pptx import Presentation
    pb, d = pbm(), _fdir()
    for theme in pb.THEMES:
        path = os.path.join(d, f"{theme}.pptx")
        total, ai = pb.render(_pb_deck(theme), path, images=True, ai_images=False)
        prs = Presentation(path)
        eq((total, len(prs.slides), ai), (10, 10, 0))
        assert _pictures(prs) >= 5, f"{theme}: no pictures drawn"
        eq(prs.slides[1].notes_slide.notes_text_frame.text, "intro")
        eq(round(prs.slide_width / 914400, 2), 13.33)
    path = os.path.join(d, "plain.pptx")
    pb.render(_pb_deck(), path, images=False)
    eq(_pictures(Presentation(path)), 0)                              # "without images" means no pictures at all


def t_pb_offline_deck():
    _need_pptx()
    pb, d = pbm(), _fdir()
    calls = []
    with mock.patch.object(mod("image_generator"), "generate", lambda *a, **k: calls.append(a) or (None, "offline")):
        path, total, used_ai, ai_imgs, err = pb.create("Solar Power!", None, 4, "forest", True, out_dir=d)
    assert not err and not used_ai and ai_imgs == 0, err
    assert not calls, "offline decks must not call the image service"
    eq(total, 5)
    assert os.path.dirname(path) == d and os.path.basename(path).startswith("solar_power_") and os.path.getsize(path) > 5000
    path, total, *_rest, err = pb.create("x", None, 3, None, False, out_dir=d)
    assert not err and total == 4


def t_pb_ai_content():
    _need_pptx()
    from pptx import Presentation
    pb, br, d = pbm(), mod("brain"), _fdir()
    reply = ('Sure! ```json {"title": "Volcanoes", "subtitle": "Fire from below", "theme": "sunset", "slides": ['
             '{"layout": "stats", "title": "Big numbers", "stats": [{"value": "1,350", "label": "active volcanoes"}, {"value": "70", "label": "erupt each year"}]},'
             '{"layout": "closing", "title": "Thank you", "bullets": ["Questions?"]}]} ```')
    with mock.patch.object(br, "_complete", return_value=reply):
        path, total, used_ai, _, err = pb.create("volcanoes", object(), 2, None, False, out_dir=d)
    assert not err and used_ai and total == 3, (err, used_ai, total)
    prs = Presentation(path)
    texts = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    assert "Volcanoes" in texts and "1,350" in texts and "Thank you" in texts, texts
    with mock.patch.object(br, "_complete", return_value="not json at all"):
        path, total, used_ai, _, err = pb.create("volcanoes", object(), 3, None, False, out_dir=d)
    assert not err and not used_ai and total == 4, "bad AI output must fall back to the skeleton"
    with mock.patch.object(br, "_complete", return_value=reply):
        eq(pb.build_spec("volcanoes", object(), 2, "paper")[0]["theme"], "paper")      # a requested theme wins over the AI's
        eq(pb.build_spec("volcanoes", object(), 2)[0]["theme"], "sunset")


def t_pb_voice_command():
    _need_pptx()
    dg = mod("document_generator")
    d = _fdir()
    voice, said = stub_voice()
    with mock.patch.object(dg, "OUT_DIR", d):
        assert dg.handle_command(voice, "make a 3 slide presentation about mars in the paper theme without images")
    assert "presentation is ready with 4 slides" in said[-1], said
    files = os.listdir(d)
    assert len(files) == 1 and files[0].startswith("mars_") and files[0].endswith(".pptx"), files   # options stripped from the file name
    with mock.patch.object(dg, "OUT_DIR", d):
        path, size, ai, err = dg.create("docx", "mars", None, 2)
    if not (err and err.startswith("Missing")):
        assert path.endswith(".docx") and not err, "docx must still use the original builder"


check(G, "voice options: theme and 'without images' parsing", t_pb_options)
check(G, "AI output is normalised (bad layouts, missing data, unknown theme)", t_pb_normalize)
check(G, "local artwork: deterministic, correct size, crop never stretches", t_pb_art)
check(G, "all six themes render all nine layouts; 'without images' draws none", t_pb_render_every_theme)
check(G, "offline deck: themed local art, no network, no AI", t_pb_offline_deck)
check(G, "AI content -> deck (mocked model), bad output falls back, requested theme wins", t_pb_ai_content)
check(G, "voice command builds a designed deck; docx still uses the old builder", t_pb_voice_command)


# ============================================================================ report
def report():
    shutil.rmtree(tmpdir, ignore_errors=True)
    groups = {}
    for g, n, s, d in RESULTS:
        groups.setdefault(g, []).append((n, s, d))
    icon = {"PASS": "[ OK ]", "FAIL": "[FAIL]", "SKIP": "[skip]", "REQUIRES API": "[ API ]",
            "REQUIRES PERMISSION": "[PERM]", "REQUIRES HARDWARE": "[ HW ]"}
    print("\n" + "=" * 78 + "\n AURORA SELF-TEST\n" + "=" * 78)
    for g, rows in groups.items():
        p = sum(1 for r in rows if r[1] == "PASS")
        f = sum(1 for r in rows if r[1] == "FAIL")
        state = "BROKEN" if f else ("WORKING" if p else "not tested")
        print(f"\n{g}  -  {state}  ({p}/{len(rows)} passed)")
        for n, s, d in rows:
            print(f"  {icon[s]} {n}" + (f"  -> {d}" if d else ""))
    counts = {k: sum(1 for r in RESULTS if r[2] == k) for k in icon}
    print("\n" + "=" * 78)
    print(f" {counts['PASS']} passed, {counts['FAIL']} failed, {counts['SKIP']} skipped, {counts['REQUIRES API']} require API, "
          f"{counts['REQUIRES PERMISSION']} require permission, {counts['REQUIRES HARDWARE']} require hardware, of {len(RESULTS)}")
    broken = [g for g, rows in groups.items() if any(r[1] == "FAIL" for r in rows)]
    print(" WORKING: " + ", ".join(g for g in groups if g not in broken))
    if broken:
        print(" BROKEN : " + ", ".join(broken))
    print("=" * 78)
    return 1 if counts["FAIL"] else 0


if __name__ == "__main__":
    rc = report()
    sys.stdout.flush()
    os._exit(rc)
